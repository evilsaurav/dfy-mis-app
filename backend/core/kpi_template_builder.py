# -*- coding: utf-8 -*-
"""
backend/core/kpi_template_builder.py
Dynamic KPI Workbook Skeleton Builder for District KPI Reports (Backlog #88 Option B).
Generates production-grade 33-tab KPI templates dynamically scaled to any staff roster.
"""
import io
import hashlib
from typing import List, Optional

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from backend.core.styles import (
    EXCEL_HEADER_BORDER,
    EXCEL_THIN_BORDER,
    EXCEL_TOTAL_ROW_BORDER,
    EXCEL_CLUSTER_DIVIDER,
)
from backend.core.kpi_constants import EXCEL_KPI_CATEGORIES, get_kpi_tab_name

# Typography & Color Palettes matching production templates
FONT_TITLE = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
FONT_HEADER = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
FONT_BOLD = Font(name="Calibri", size=10, bold=True, color="000000")
FONT_REGULAR = Font(name="Calibri", size=10, bold=False, color="000000")

FILL_NAVY = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
FILL_INDIGO = PatternFill(start_color="4338CA", end_color="4338CA", fill_type="solid")
FILL_GRAY_HEADER = PatternFill(start_color="374151", end_color="374151", fill_type="solid")
FILL_GREEN_HEADER = PatternFill(start_color="047857", end_color="047857", fill_type="solid")
FILL_GOLD_TOTAL = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")

BLOCK_BOTTOM_BORDER = Border(
    left=Side(style='thin', color='CBD5E1'),
    right=Side(style='thin', color='CBD5E1'),
    top=Side(style='thin', color='CBD5E1'),
    bottom=Side(style='medium', color='475569')
)

ALIGN_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")


def compute_roster_fingerprint(staff_ids: List[int]) -> str:
    """
    Computes a deterministic SHA-256 fingerprint from a list of staff IDs.
    Sorted to ensure order-independence across roster queries.
    """
    sorted_ids = sorted(staff_ids)
    raw_str = ",".join(str(sid) for sid in sorted_ids)
    return hashlib.sha256(raw_str.encode("utf-8")).hexdigest()


def build_district_skeleton_bytes(district: str, staff_list: List[str]) -> bytes:
    """
    Dynamically generates a 33-tab KPI Excel workbook skeleton sized exactly
    to the provided staff_list for the given district.

    Returns the raw Excel file bytes (in-memory, no disk writes).
    """
    wb = openpyxl.Workbook()
    wb.properties.creator = "Saurav Kumar"
    wb.properties.lastModifiedBy = "Saurav Kumar"
    wb.properties.company = "Doctors For You"
    wb.properties.title = f"District KPI Performance Sheet - {district.upper()}"
    default_sheet = wb.active

    # -------------------------------------------------------------
    # 1. TAB 1: 'Performance sheet'
    # -------------------------------------------------------------
    ws_perf = wb.create_sheet(title="Performance sheet")
    ws_perf.views.sheetView[0].showGridLines = True

    perf_headers = [
        "Employee Name", "DESIG.", "Target", "NOTIFICATION", "% Achieved"
    ] + [kpi[0] for kpi in EXCEL_KPI_CATEGORIES[1:]]  # 21 columns (Cols A to U)

    # Title Banner (Rows 1-2 merged)
    banner_end_col = get_column_letter(len(perf_headers))
    ws_perf.merge_cells(f"A1:{banner_end_col}2")
    ws_perf["A1"] = f"DOCTORS FOR YOU (DFY) -- DISTRICT KPI PERFORMANCE SHEET ({district.upper()})"
    ws_perf["A1"].font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    ws_perf["A1"].fill = FILL_NAVY
    ws_perf["A1"].alignment = ALIGN_CENTER
    for r in range(1, 3):
        for c in range(1, len(perf_headers) + 1):
            ws_perf.cell(row=r, column=c).border = EXCEL_HEADER_BORDER

    # Subtitle (Row 3)
    ws_perf["A3"] = "Monthly Evaluation & Notification Target Achievement Matrix | Generated on DFY MIS by Saurav Kumar"
    ws_perf["A3"].font = Font(name="Calibri", size=10, italic=True, color="6B7280")

    # Column Headers (Row 4)
    for c_idx, h in enumerate(perf_headers, start=1):
        cell = ws_perf.cell(row=4, column=c_idx, value=h)
        cell.font = FONT_HEADER
        cell.fill = FILL_INDIGO if c_idx <= 5 else FILL_GRAY_HEADER
        cell.alignment = ALIGN_CENTER
        cell.border = EXCEL_HEADER_BORDER

    # Staff Rows (Row 5 onwards)
    start_r = 5
    for idx, staff_name in enumerate(staff_list):
        r = start_r + idx
        cell_name = ws_perf.cell(row=r, column=1, value=staff_name)
        cell_name.font = FONT_BOLD
        cell_name.alignment = ALIGN_LEFT

        cell_desig = ws_perf.cell(row=r, column=2, value="Field Officer")
        cell_desig.font = FONT_REGULAR
        cell_desig.alignment = ALIGN_CENTER

        cell_tgt = ws_perf.cell(row=r, column=3, value=50)  # Default target
        cell_tgt.font = FONT_BOLD
        cell_tgt.alignment = ALIGN_CENTER

        cell_notif = ws_perf.cell(row=r, column=4, value=0)
        cell_notif.font = FONT_BOLD
        cell_notif.alignment = ALIGN_CENTER

        cell_pct = ws_perf.cell(row=r, column=5, value=f"=IF(C{r}>0, D{r}/C{r}, 0)")
        cell_pct.font = FONT_BOLD
        cell_pct.alignment = ALIGN_CENTER
        cell_pct.number_format = "0.0%"

        for k_idx in range(1, len(EXCEL_KPI_CATEGORIES)):
            col_target = 5 + k_idx
            c = ws_perf.cell(row=r, column=col_target, value=0)
            c.font = FONT_REGULAR
            c.alignment = ALIGN_CENTER
            c.border = EXCEL_THIN_BORDER

        for c in range(1, len(perf_headers) + 1):
            ws_perf.cell(row=r, column=c).border = EXCEL_THIN_BORDER

    # Grand Total Row
    total_r = start_r + len(staff_list)
    cell_gt = ws_perf.cell(row=total_r, column=1, value="GRAND TOTAL")
    cell_gt.font = FONT_TITLE
    cell_gt.fill = FILL_NAVY
    cell_gt.alignment = ALIGN_LEFT

    ws_perf.cell(row=total_r, column=2, value="").fill = FILL_NAVY

    # Target Total Formula
    c_tgt = ws_perf.cell(row=total_r, column=3, value=f"=SUM(C{start_r}:C{total_r-1})")
    c_tgt.font = FONT_TITLE
    c_tgt.fill = FILL_NAVY
    c_tgt.alignment = ALIGN_CENTER

    # Notification Total Formula
    c_notif = ws_perf.cell(row=total_r, column=4, value=f"=SUM(D{start_r}:D{total_r-1})")
    c_notif.font = FONT_TITLE
    c_notif.fill = FILL_NAVY
    c_notif.alignment = ALIGN_CENTER

    # Overall % Achieved Total Formula
    cell_tot_pct = ws_perf.cell(row=total_r, column=5, value=f"=IF(C{total_r}>0, D{total_r}/C{total_r}, 0)")
    cell_tot_pct.font = FONT_TITLE
    cell_tot_pct.fill = FILL_NAVY
    cell_tot_pct.alignment = ALIGN_CENTER
    cell_tot_pct.number_format = "0.0%"

    # KPI Columns 6..21 Total Formulas
    for k_idx in range(1, len(EXCEL_KPI_CATEGORIES)):
        col_target = 5 + k_idx
        col_letter = get_column_letter(col_target)
        c = ws_perf.cell(row=total_r, column=col_target, value=f"=SUM({col_letter}{start_r}:{col_letter}{total_r-1})")
        c.font = FONT_TITLE
        c.fill = FILL_NAVY
        c.alignment = ALIGN_CENTER
        c.border = EXCEL_TOTAL_ROW_BORDER

    for c in range(1, len(perf_headers) + 1):
        ws_perf.cell(row=total_r, column=c).border = EXCEL_TOTAL_ROW_BORDER

    ws_perf.column_dimensions['A'].width = 24
    ws_perf.column_dimensions['B'].width = 14
    ws_perf.column_dimensions['C'].width = 12
    ws_perf.column_dimensions['D'].width = 15
    ws_perf.column_dimensions['E'].width = 14
    for c in range(6, len(perf_headers) + 1):
        ws_perf.column_dimensions[get_column_letter(c)].width = 16

    # -------------------------------------------------------------
    # 2. TAB 2: 'CONSOLIDATED SHEET'
    # -------------------------------------------------------------
    ws_cons = wb.create_sheet(title="CONSOLIDATED SHEET")
    ws_cons.views.sheetView[0].showGridLines = True

    # Wing 1: Left Side (District Master Rollup & Master Log) -- Cols 1 to 48 (16 clusters x 3 cols)
    left_clusters = [kpi[0] for kpi in EXCEL_KPI_CATEGORIES if kpi[0] != "Kit Consumption"]

    for c_idx, cluster_name in enumerate(left_clusters):
        start_c = 1 + (c_idx * 3)
        end_c = start_c + 2

        ws_cons.merge_cells(start_row=1, start_column=start_c, end_row=1, end_column=end_c)
        r1_cell = ws_cons.cell(row=1, column=start_c, value=cluster_name)
        r1_cell.font = FONT_TITLE
        r1_cell.fill = FILL_NAVY
        r1_cell.alignment = ALIGN_CENTER

        ws_cons.merge_cells(start_row=2, start_column=start_c, end_row=2, end_column=end_c)
        r2_cell = ws_cons.cell(row=2, column=start_c, value=0)
        r2_cell.font = FONT_TITLE
        r2_cell.fill = FILL_INDIGO
        r2_cell.alignment = ALIGN_CENTER

        subheaders = ["Patient ID", "Date", "Reported by"]
        for s_idx, sh in enumerate(subheaders):
            sub_cell = ws_cons.cell(row=3, column=start_c + s_idx, value=sh)
            sub_cell.font = FONT_HEADER
            sub_cell.fill = FILL_GRAY_HEADER
            sub_cell.alignment = ALIGN_CENTER

        for r in range(1, 4):
            for c in range(start_c, end_c + 1):
                if c == end_c:
                    ws_cons.cell(row=r, column=c).border = EXCEL_CLUSTER_DIVIDER
                else:
                    ws_cons.cell(row=r, column=c).border = EXCEL_THIN_BORDER

        # Pre-border initial 50 rows of data
        for r in range(4, 54):
            for c in range(start_c, end_c + 1):
                if c == end_c:
                    ws_cons.cell(row=r, column=c).border = EXCEL_CLUSTER_DIVIDER
                else:
                    ws_cons.cell(row=r, column=c).border = EXCEL_THIN_BORDER

        ws_cons.column_dimensions[get_column_letter(start_c)].width = 14
        ws_cons.column_dimensions[get_column_letter(start_c + 1)].width = 12
        ws_cons.column_dimensions[get_column_letter(start_c + 2)].width = 18

    # Wing 2: Right Side (Staff-Wise Performance & Indicator Wing) -- Col 49 onwards
    for s_idx, staff_name in enumerate(staff_list):
        staff_start_c = 49 + (s_idx * 17)
        staff_end_c = staff_start_c + 16

        ws_cons.merge_cells(start_row=1, start_column=staff_start_c, end_row=1, end_column=staff_end_c)
        r1_staff = ws_cons.cell(row=1, column=staff_start_c, value=staff_name)
        r1_staff.font = FONT_TITLE
        r1_staff.fill = FILL_GREEN_HEADER
        r1_staff.alignment = ALIGN_CENTER

        for k_idx, (kpi_name, _, _) in enumerate(EXCEL_KPI_CATEGORIES):
            c_num = staff_start_c + k_idx

            h_cell = ws_cons.cell(row=2, column=c_num, value=kpi_name)
            h_cell.font = FONT_HEADER
            h_cell.fill = FILL_GRAY_HEADER
            h_cell.alignment = ALIGN_CENTER

            r3_cell = ws_cons.cell(row=3, column=c_num, value=0)
            r3_cell.font = FONT_BOLD
            r3_cell.fill = FILL_GOLD_TOTAL
            r3_cell.alignment = ALIGN_CENTER

            ws_cons.column_dimensions[get_column_letter(c_num)].width = 15

        for r in range(1, 4):
            for c in range(staff_start_c, staff_end_c + 1):
                if c == staff_end_c:
                    ws_cons.cell(row=r, column=c).border = EXCEL_CLUSTER_DIVIDER
                else:
                    ws_cons.cell(row=r, column=c).border = EXCEL_THIN_BORDER

        for r in range(4, 54):
            for c in range(staff_start_c, staff_end_c + 1):
                if c == staff_end_c:
                    ws_cons.cell(row=r, column=c).border = EXCEL_CLUSTER_DIVIDER
                else:
                    ws_cons.cell(row=r, column=c).border = EXCEL_THIN_BORDER

    # -------------------------------------------------------------
    # 3. TABS 3 to 33: Daily Tabs ('1ST' to '31st')
    # -------------------------------------------------------------
    daily_headers = ["NAME", "DESIGNATION"] + [kpi[0] for kpi in EXCEL_KPI_CATEGORIES]

    for day_int in range(1, 32):
        tab_title = get_kpi_tab_name(day_int)
        ws_day = wb.create_sheet(title=tab_title)
        ws_day.views.sheetView[0].showGridLines = True

        for c_idx, h in enumerate(daily_headers, start=1):
            cell = ws_day.cell(row=1, column=c_idx, value=h)
            cell.font = FONT_HEADER
            cell.fill = FILL_NAVY if c_idx <= 2 else FILL_GRAY_HEADER
            cell.alignment = ALIGN_CENTER
            cell.border = EXCEL_HEADER_BORDER

        for s_idx, staff_name in enumerate(staff_list):
            fo_start_r = 2 + (s_idx * 40)
            fo_end_r = fo_start_r + 39

            cell_name = ws_day.cell(row=fo_start_r, column=1, value=staff_name)
            cell_name.font = FONT_BOLD
            cell_name.alignment = ALIGN_LEFT

            cell_des = ws_day.cell(row=fo_start_r, column=2, value="Field Officer")
            cell_des.font = FONT_REGULAR
            cell_des.alignment = ALIGN_CENTER

            for r in range(fo_start_r, fo_end_r + 1):
                is_bottom_row = (r == fo_end_r)
                for c in range(1, len(daily_headers) + 1):
                    if is_bottom_row:
                        ws_day.cell(row=r, column=c).border = BLOCK_BOTTOM_BORDER
                    else:
                        ws_day.cell(row=r, column=c).border = EXCEL_THIN_BORDER
                    ws_day.cell(row=r, column=c).alignment = ALIGN_CENTER if c > 2 else (ALIGN_LEFT if c == 1 else ALIGN_CENTER)

        ws_day.column_dimensions['A'].width = 22
        ws_day.column_dimensions['B'].width = 14
        for c in range(3, len(daily_headers) + 1):
            ws_day.column_dimensions[get_column_letter(c)].width = 15

    # Remove default initial sheet
    if default_sheet.title in wb.sheetnames and len(wb.sheetnames) > 33:
        wb.remove(default_sheet)
    elif "Sheet" in wb.sheetnames:
        wb.remove(wb["Sheet"])

    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()
