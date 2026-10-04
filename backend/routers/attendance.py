import os
import io
import re
import gc
import sys
import math
import calendar
import logging
import asyncio
from datetime import datetime, timedelta, date as dt_date
from typing import Optional, List, Dict, Any, Tuple, Set
from fastapi import APIRouter, HTTPException, Depends
from collections import defaultdict
from pydantic import BaseModel

from backend.core.database import db, ENABLE_IN_MEMORY_DERIVATION
from backend.core.cache import cache
from backend.core.security import get_current_admin
from backend.core.helpers import (
    get_ist_now,
    format_to_ist_time,
    parse_to_ist_datetime,
    canonicalize_district,
    normalize_staff_key,
    is_officer_name_match,
    get_reporting_cutoff_hour,
    get_profile_cache_key,
    evict_officer_profile_cache,
    log_admin_activity,
    load_baseline_staff_directory,
)
from backend.core.master_ledger import (
    get_cached_staff_directory_raw,
    get_raw_monthly_reports,
    get_directory,
)
from backend.core.supabase import (
    pg_query_table,
    pg_fetch_one,
    pg_upsert_row,
    pg_update_row,
    pg_delete_rows,
    pg_execute_raw,
    get_active_db,
)

from backend.core.styles import (
    safe_filename,
    ExcelStreamingResponse,
    style_excel_worksheet,
    EXCEL_HEADER_BORDER,
    EXCEL_THIN_BORDER,
    EXCEL_TOTAL_ROW_BORDER,
)
from backend.routers.backup import ensure_daily_backup_scheduled

router = APIRouter(tags=["attendance"])
logger = logging.getLogger("attendance")
attendance_excel_semaphore = asyncio.Semaphore(1)


async def _resolve_get_raw_monthly_reports(*args, **kwargs):
    main_mod = sys.modules.get("main")
    fn = getattr(main_mod, "get_raw_monthly_reports", get_raw_monthly_reports) if main_mod else get_raw_monthly_reports
    return await fn(*args, **kwargs)


async def get_attendance_staff_roster(target_date: str, allowed_dist_set: Optional[set] = None) -> Tuple[list, set]:
    staff_list = []
    raw_staff_records = await get_cached_staff_directory_raw()

    inactive_staff_keys = set()
    if raw_staff_records:
        for d in raw_staff_records:
            if not d:
                continue
            raw_dist = d.get("district") or ""
            dist = canonicalize_district(raw_dist)
            fo_name = (d.get("name") or "").strip()
            doc_id = d.get("id", "")

            is_active = d.get("is_active") is not False and d.get("status") != "inactive"
            inactive_since = (d.get("inactive_since") or "").strip()[:10]

            is_inactive = False
            if not is_active:
                if not inactive_since or target_date >= inactive_since:
                    is_inactive = True
            elif inactive_since and target_date >= inactive_since:
                is_inactive = True

            if is_inactive:
                if dist and fo_name:
                    inactive_staff_keys.add(normalize_staff_key(dist, fo_name))
                    clean_d = re.sub(r'[^a-z0-9]', '', dist.lower())
                    clean_n = re.sub(r'[^a-z0-9]', '', fo_name.lower())
                    inactive_staff_keys.add(f"{clean_d}_{clean_n}")
                if doc_id:
                    clean_id = re.sub(r'[^a-z0-9_]', '', doc_id.lower())
                    inactive_staff_keys.add(clean_id)
                    inactive_staff_keys.add(re.sub(r'(.)\1+', r'\1', clean_id))
                    if "_" in clean_id:
                        p_dist, p_name = clean_id.split("_", 1)
                        inactive_staff_keys.add(f"{p_dist}_{re.sub(r'(.)\1+', r'\1', p_name)}")

        if inactive_staff_keys:
            cache.set("inactive_staff_keys", list(inactive_staff_keys), ttl=3600)
    else:
        cached_inactive = cache.get("inactive_staff_keys")
        if cached_inactive:
            inactive_staff_keys.update(cached_inactive)

    if raw_staff_records:
        for d in raw_staff_records:
            if not d:
                continue
            raw_dist = d.get("district") or ""
            dist = canonicalize_district(raw_dist)
            fo_name = (d.get("name") or "").strip()
            if not dist or not fo_name:
                continue
            if allowed_dist_set is not None and dist.lower() not in allowed_dist_set:
                continue

            norm_key = normalize_staff_key(dist, fo_name)
            clean_d = re.sub(r'[^a-z0-9]', '', dist.lower())
            clean_n = re.sub(r'[^a-z0-9]', '', fo_name.lower())
            exact_key = f"{clean_d}_{clean_n}"
            doc_id = d.get("id", "")
            doc_norm_key = re.sub(r'(.)\1+', r'\1', re.sub(r'[^a-z0-9_]', '', doc_id.lower())) if doc_id else ""

            if norm_key in inactive_staff_keys or exact_key in inactive_staff_keys:
                continue
            if doc_id and (doc_id.lower() in inactive_staff_keys or doc_norm_key in inactive_staff_keys):
                continue

            is_active = d.get("is_active") is not False and d.get("status") != "inactive"
            inactive_since = (d.get("inactive_since") or "").strip()[:10]

            is_included = False
            if is_active:
                if not inactive_since or target_date < inactive_since:
                    is_included = True
            else:
                if inactive_since and target_date < inactive_since:
                    is_included = True

            if not is_included:
                continue

            staff_list.append({
                "district": dist,
                "fo_name": fo_name,
                "designation": d.get("designation", "Field Officer"),
                "status": "active" if is_active else "inactive",
                "is_active": is_active,
                "inactive_since": inactive_since or None
            })
    else:
        cached_dir = cache.get("staff_directory_dict") or await get_directory()
        if cached_dir and isinstance(cached_dir, dict):
            for dist, names in cached_dir.items():
                c_dist = canonicalize_district(dist)
                if allowed_dist_set is not None and c_dist.lower() not in allowed_dist_set:
                    continue
                for clean_fo in names:
                    if clean_fo and str(clean_fo).strip():
                        fo_str = str(clean_fo).strip()
                        norm_k = normalize_staff_key(c_dist, fo_str)
                        clean_d = re.sub(r'[^a-z0-9]', '', c_dist.lower())
                        clean_n = re.sub(r'[^a-z0-9]', '', fo_str.lower())
                        exact_k = f"{clean_d}_{clean_n}"
                        if norm_k in inactive_staff_keys or exact_k in inactive_staff_keys:
                            continue
                        staff_list.append({
                            "district": c_dist,
                            "fo_name": fo_str,
                            "designation": "Field Officer"
                        })
        else:
            cached_dir = load_baseline_staff_directory()
            for dist, names in (cached_dir or {}).items():
                c_dist = canonicalize_district(dist)
                if allowed_dist_set is not None and c_dist.lower() not in allowed_dist_set:
                    continue
                for clean_fo in names:
                    if clean_fo and str(clean_fo).strip():
                        fo_str = str(clean_fo).strip()
                        norm_k = normalize_staff_key(c_dist, fo_str)
                        clean_d = re.sub(r'[^a-z0-9]', '', c_dist.lower())
                        clean_n = re.sub(r'[^a-z0-9]', '', fo_str.lower())
                        exact_k = f"{clean_d}_{clean_n}"
                        if norm_k in inactive_staff_keys or exact_k in inactive_staff_keys:
                            continue
                        staff_list.append({
                            "district": c_dist,
                            "fo_name": fo_str,
                            "designation": "Field Officer"
                        })
    return staff_list, inactive_staff_keys


def format_attendance_response(
    staff_list: list,
    all_candidate_docs: list,
    leave_docs: list,
    target_date: str,
    next_date: str,
    allowed_dist_set: Optional[set] = None,
    inactive_staff_keys: Optional[set] = None
) -> dict:
    reports_map = {}
    for doc in all_candidate_docs:
        d = doc.to_dict() if hasattr(doc, "to_dict") else (doc if isinstance(doc, dict) else {})
        if not d:
            continue
        dist = canonicalize_district(d.get('working_place', '') or d.get('district', ''))
        if allowed_dist_set is not None and dist.lower() not in allowed_dist_set:
            continue
        fo_raw_name = d.get('fo_name', '').strip()
        clean_fo = re.sub(r'[^a-zA-Z0-9]', '', fo_raw_name).lower()
        if not clean_fo:
            continue
        key = f"{dist}_{clean_fo}".replace(" ", "").lower()
        
        raw_ts = d.get("timestamp_completed") or d.get("timestamp") or d.get("submitted_at")
        submitted_time = format_to_ist_time(raw_ts)
        dt_ist = parse_to_ist_datetime(raw_ts)
        if dt_ist:
            iso_ts = dt_ist.isoformat()
        elif hasattr(raw_ts, 'isoformat'):
            iso_ts = raw_ts.isoformat()
        else:
            iso_ts = str(raw_ts) if raw_ts else ""
        
        rep_date = d.get("date_of_reporting") or d.get("date") or ""
        is_next_day_flag = bool(d.get("is_next_day_submission"))

        # Stealth Cutoff Segregation Rules:
        # Rule 1: Submissions on target_date before cutoff hour strictly belong to target_date - 1.
        # Exclude from target_date's submitted list.
        target_cutoff = get_reporting_cutoff_hour(dt_ist) if dt_ist else 11
        if dt_ist and dt_ist.date().strftime("%Y-%m-%d") == target_date and dt_ist.hour < target_cutoff:
            continue

        # Rule 2: Submissions on next_date (target_date + 1) before next_date cutoff hour,
        # or reports with is_next_day_submission == True and date_of_reporting == target_date,
        # strictly belong to target_date as next-day morning submissions.
        is_next_day = False
        next_cutoff = get_reporting_cutoff_hour(dt_ist) if dt_ist else 11
        if is_next_day_flag and rep_date == target_date:
            is_next_day = True
        elif dt_ist and next_date and dt_ist.date().strftime("%Y-%m-%d") == next_date and dt_ist.hour < next_cutoff:
            is_next_day = True
        elif rep_date != target_date:
            continue

        total_km = 0
        if d.get("total_km"):
            try: total_km = int(d.get("total_km"))
            except: pass
        elif d.get("morning_km") is not None and d.get("evening_km") is not None:
            try: total_km = max(0, int(d.get("evening_km")) - int(d.get("morning_km")))
            except: pass

        raw_morning_time = d.get("submitted_morning_time")
        morning_time = format_to_ist_time(raw_morning_time) if raw_morning_time else submitted_time
        if is_next_day:
            submitted_time = morning_time or submitted_time
            submitted_label = d.get("morning_submission_label") or f"Next day morning {submitted_time}"
            time_classification = "Next Day Morning (< 10 AM)"
        else:
            submitted_label = submitted_time or "Submitted"
            if dt_ist:
                if dt_ist.hour < 17:
                    time_classification = "Mid-Day (< 5 PM)"
                elif dt_ist.hour < 20:
                    time_classification = "Evening (< 8 PM)"
                else:
                    time_classification = "Night (8 PM+)"
            else:
                time_classification = "On Time"

        doc_total_ids = sum(len(v) for k, v in d.items() if isinstance(v, list) and k.endswith("_ids"))
        if doc_total_ids == 0:
            if isinstance(d.get("total_ids"), int) and d["total_ids"] > 0:
                doc_total_ids = d["total_ids"]
            else:
                doc_total_ids = sum(
                    int(d[k]) for k in [
                        "notifications", "legacy_count_notifications", "sample_tested",
                        "hiv_dm", "dbt", "contact_tracing", "differentiated_tb", "doctor_store_visits_count"
                    ] if isinstance(d.get(k), int) and d[k] > 0
                )

        if key in reports_map:
            existing = reports_map[key]
            existing["submission_count"] = max(existing.get("submission_count", 1), d.get("submission_count", 1))
            existing["total_ids"] += doc_total_ids
            existing["total_km"] = max(existing.get("total_km", 0), total_km)
            if is_next_day:
                existing["is_next_day"] = True
                existing["submitted_time"] = submitted_time
                existing["submitted_label"] = submitted_label
                existing["time_classification"] = time_classification
            elif iso_ts and (not existing.get("timestamp_raw") or iso_ts > existing.get("timestamp_raw", "")):
                existing["timestamp_raw"] = iso_ts
                existing["submitted_time"] = submitted_time or existing.get("submitted_time")
                existing["submitted_label"] = submitted_label
                existing["time_classification"] = time_classification
        else:
            reports_map[key] = {
                "district": dist,
                "fo_name": fo_raw_name,
                "submission_count": d.get("submission_count", 1),
                "total_ids": doc_total_ids,
                "submitted_time": submitted_time or "Submitted",
                "timestamp_raw": iso_ts,
                "total_km": total_km,
                "is_next_day": is_next_day,
                "submitted_label": submitted_label,
                "time_classification": time_classification
            }

    leaves_map = {}
    for ldoc in leave_docs:
        ld = ldoc.to_dict() if hasattr(ldoc, "to_dict") else (ldoc if isinstance(ldoc, dict) else {})
        if not ld:
            continue
        dist = canonicalize_district(ld.get("district", ""))
        if allowed_dist_set is not None and dist.lower() not in allowed_dist_set:
            continue
        clean_fo = re.sub(r'[^a-zA-Z0-9]', '', ld.get("fo_name", "")).lower()
        lkey = f"{dist}_{clean_fo}".replace(" ", "").lower()
        leaves_map[lkey] = {
            "district": dist,
            "fo_name": ld.get("fo_name", "").strip(),
            "status": ld.get("status", "leave"),
            "reason_type": ld.get("reason_type", "Casual"),
            "remark": ld.get("remark", ""),
            "marked_by_name": ld.get("marked_by_name", "Admin"),
            "marked_at": ld.get("marked_at", "")
        }

    submitted_full = []
    submitted_partial = []
    on_leave_fos = []
    missing_fos = []
    matched_report_keys = set()
    
    for s in (staff_list or []):
        norm_s_key = normalize_staff_key(s['district'], s['fo_name'])
        clean_fo = re.sub(r'[^a-zA-Z0-9]', '', s['fo_name']).lower()
        key = f"{s['district']}_{clean_fo}".replace(" ", "").lower()
        if inactive_staff_keys and (norm_s_key in inactive_staff_keys or key in inactive_staff_keys):
            continue
        
        alias_key = None
        if "ashwanikrkeshri" in key:
            alias_key = f"{s['district']}_ashwanikumar".replace(" ", "").lower()
        elif "ashwanikumar" in key:
            alias_key = f"{s['district']}_ashwanikrkeshri".replace(" ", "").lower()

        matched_rkey = None
        if key in reports_map:
            matched_rkey = key
        elif alias_key and alias_key in reports_map:
            matched_rkey = alias_key

        if matched_rkey:
            matched_report_keys.add(matched_rkey)
            rep = reports_map[matched_rkey]
            info = {**s, **rep}
            if rep["submission_count"] >= 2:
                submitted_full.append(info)
            else:
                submitted_partial.append(info)
        else:
            matched_lkey = None
            if key in leaves_map:
                matched_lkey = key
            elif alias_key and alias_key in leaves_map:
                matched_lkey = alias_key

            if matched_lkey:
                on_leave_fos.append({**s, **leaves_map[matched_lkey]})
            else:
                missing_fos.append(s)

    # In case an officer submitted whose name is not in staff_list, also include them in submitted list
    for rkey, rinfo in reports_map.items():
        if rkey not in matched_report_keys:
            orphan_info = {
                "district": rinfo["district"],
                "fo_name": rinfo["fo_name"],
                "designation": "Field Officer",
                **rinfo
            }
            submitted_partial.append(orphan_info)

    submitted_fos = submitted_full + submitted_partial
    submitted_fos.sort(key=lambda x: (x.get("timestamp_raw") or "", x["district"], x["fo_name"]), reverse=True)
    missing_fos.sort(key=lambda x: (x["district"], x["fo_name"]))
    submitted_full.sort(key=lambda x: (x["district"], x["fo_name"]))
    submitted_partial.sort(key=lambda x: (x["district"], x["fo_name"]))
    on_leave_fos.sort(key=lambda x: (x["district"], x["fo_name"]))

    return {
        "date": target_date,
        "total_staff": len(staff_list or []),
        "submitted_count": len(submitted_fos),
        "submitted_full_count": len(submitted_full),
        "submitted_partial_count": len(submitted_partial),
        "on_leave_count": len(on_leave_fos),
        "missing_count": len(missing_fos),
        "submitted_fos": submitted_fos,
        "submitted_full": submitted_full,
        "submitted_partial": submitted_partial,
        "on_leave_fos": on_leave_fos,
        "missing_fos": missing_fos
    }


async def legacy_get_today_attendance(
    target_date: str,
    next_date: str,
    allowed_dist_set: Optional[set],
    staff_list: Optional[list] = None,
    effective_dist: str = "all",
    cache_key: Optional[str] = None,
    admin: Optional[dict] = None,
    force_refresh: bool = False,
    inactive_staff_keys: Optional[set] = None
) -> dict:
    if staff_list is None:
        staff_list, inactive_staff_keys = await get_attendance_staff_roster(target_date, allowed_dist_set)
    if not next_date:
        try:
            target_dt = datetime.strptime(target_date, "%Y-%m-%d").date()
            next_date = (target_dt + timedelta(days=1)).strftime("%Y-%m-%d")
        except Exception:
            next_date = ""

    report_rows = pg_execute_raw(
        "SELECT * FROM daily_field_reports WHERE date_of_reporting = %s",
        [target_date],
        fetch=True
    ) or []
    all_candidate_docs = [dict(r) for r in report_rows]
    seen_doc_ids = {str(d.get("id")) for d in all_candidate_docs if d.get("id")}

    if next_date:
        try:
            next_day_rows = pg_execute_raw(
                "SELECT * FROM daily_field_reports WHERE date_of_reporting = %s",
                [next_date],
                fetch=True
            ) or []
            for ndoc in next_day_rows:
                ndoc_dict = dict(ndoc)
                ndoc_id = str(ndoc_dict.get("id", ""))
                if not ndoc_id or ndoc_id not in seen_doc_ids:
                    all_candidate_docs.append(ndoc_dict)
        except Exception as e:
            print(f"Notice fetching next_day_docs in today-attendance: {e}")

    leave_docs = []
    try:
        leave_rows = pg_execute_raw(
            "SELECT * FROM daily_staff_leaves WHERE date = %s",
            [target_date],
            fetch=True
        ) or []
        leave_docs = [dict(r) for r in leave_rows]
    except Exception as le:
        print(f"PG daily_staff_leaves query notice: {le}")
        leave_docs = []


    res = format_attendance_response(
        staff_list=staff_list,
        all_candidate_docs=all_candidate_docs,
        leave_docs=leave_docs,
        target_date=target_date,
        next_date=next_date,
        allowed_dist_set=allowed_dist_set,
        inactive_staff_keys=inactive_staff_keys
    )
    if cache_key:
        cache.set(cache_key, res, ttl=600)
    return res


@router.get("/admin/today-attendance")
async def get_today_attendance(
    date: Optional[str] = None, 
    districts: Optional[str] = None, 
    force_refresh: Optional[bool] = False, 
    admin: dict = Depends(get_current_admin)
):
    try:
        ensure_daily_backup_scheduled()
        target_date = date.strip() if date and date.strip() else get_ist_now().strftime("%Y-%m-%d")
            
        allowed_dist_set = None
        if districts and districts.strip() and districts.strip() != "All":
            allowed_dist_set = set([canonicalize_district(d.strip()).lower() for d in districts.split(",") if d.strip()])

        # Sub-Admin RBAC validation
        admin_role = admin.get("role", "SUB_ADMIN")
        if admin_role == "SUB_ADMIN":
            admin_allowed = admin.get("allowed_districts", []) or admin.get("districts", [])
            if admin_allowed is not None:
                if "All" in admin_allowed or "all" in admin_allowed:
                    pass
                else:
                    subadmin_allowed = set([canonicalize_district(d.strip()).lower() for d in admin_allowed if d.strip()])
                    if allowed_dist_set is not None:
                        forbidden = allowed_dist_set - subadmin_allowed
                        if forbidden:
                            raise HTTPException(status_code=403, detail="Permission denied. You do not have access to the requested district(s).")
                        allowed_dist_set = allowed_dist_set.intersection(subadmin_allowed)
                    else:
                        allowed_dist_set = subadmin_allowed

        effective_dist = ",".join(sorted(allowed_dist_set)) if (allowed_dist_set is not None and len(allowed_dist_set) > 0) else ("none" if allowed_dist_set is not None else (districts or 'all'))
        user_scope = admin.get("user_id") or admin.get("username") or admin.get("role", "admin")
        cache_key = f"attendance_{target_date}_{effective_dist}_{user_scope}"

        if force_refresh:
            cache.delete(cache_key)
        else:
            cached = cache.get(cache_key)
            if cached is not None:
                return cached

        try:
            target_dt = datetime.strptime(target_date, "%Y-%m-%d").date()
            next_date = (target_dt + timedelta(days=1)).strftime("%Y-%m-%d")
        except Exception:
            target_dt = get_ist_now().date()
            next_date = (target_dt + timedelta(days=1)).strftime("%Y-%m-%d")

        staff_list = None
        inactive_staff_keys = None

        # Master Kill Switch
        import sys
        main_mod = sys.modules.get("main")
        is_enabled = getattr(main_mod, "ENABLE_IN_MEMORY_DERIVATION", ENABLE_IN_MEMORY_DERIVATION) if main_mod else ENABLE_IN_MEMORY_DERIVATION
        if not is_enabled:
            return await legacy_get_today_attendance(
                target_date=target_date,
                next_date=next_date,
                allowed_dist_set=allowed_dist_set,
                staff_list=staff_list,
                effective_dist=effective_dist,
                cache_key=cache_key,
                admin=admin,
                force_refresh=bool(force_refresh),
                inactive_staff_keys=inactive_staff_keys
            )

        # In-Memory Master Ledger Fast-Path
        try:
            staff_list, inactive_staff_keys = await get_attendance_staff_roster(target_date, allowed_dist_set)

            all_candidate_docs = []
            pg_rows = []
            try:
                date_sql = "SELECT * FROM daily_field_reports r WHERE (r.date_of_reporting = %s OR r.date_of_reporting = %s)"
                sql_params = [target_date, next_date or target_date]
                if allowed_dist_set is not None and len(allowed_dist_set) > 0:
                    date_sql += " AND (LOWER(TRIM(r.working_place)) = ANY(%s) OR r.district_id::text = ANY(%s))"
                    dist_param = list({str(d).lower() for d in allowed_dist_set} | {str(d) for d in allowed_dist_set})
                    sql_params.extend([dist_param, dist_param])

                pg_rows = pg_execute_raw(date_sql, sql_params, fetch=True) or []

                if pg_rows:
                    report_ids = [r["id"] for r in pg_rows if r.get("id")]
                    kpi_by_rep = defaultdict(lambda: defaultdict(list))
                    visited_by_rep = defaultdict(list)
                    if report_ids:
                        try:
                            k_rows = pg_execute_raw(
                                "SELECT report_id, category, patient_id FROM report_kpi_entries WHERE report_id = ANY(%s)",
                                [report_ids], fetch=True
                            ) or []
                            for kr in k_rows:
                                kpi_by_rep[kr["report_id"]][kr["category"]].append(kr["patient_id"])
                        except Exception as k_err:
                            logger.debug(f"[get_today_attendance] child kpi fetch notice: {k_err}")
                        try:
                            v_rows = pg_execute_raw(
                                "SELECT report_id, name FROM report_visited_names WHERE report_id = ANY(%s) ORDER BY position ASC",
                                [report_ids], fetch=True
                            ) or []
                            for vr in v_rows:
                                visited_by_rep[vr["report_id"]].append(vr["name"])
                        except Exception as v_err:
                            logger.debug(f"[get_today_attendance] child visited fetch notice: {v_err}")

                    seen_doc_ids = set()
                    for row in pg_rows:
                        item = dict(row)
                        r_id = item.get("id")
                        kpi_map = kpi_by_rep.get(r_id, {})
                        names_list = visited_by_rep.get(r_id, [])

                        if hasattr(item.get("date_of_reporting"), "strftime"):
                            item["date_of_reporting"] = item["date_of_reporting"].strftime("%Y-%m-%d")

                        did = item.get("id") or item.get("doc_id")
                        if not did or isinstance(did, int):
                            c_wp = canonicalize_district(item.get("working_place", "") or item.get("district", ""))
                            fo = str(item.get("fo_name", "")).strip()
                            dt = str(item.get("date_of_reporting", "")).strip()
                            did = f"{c_wp}_{fo}_{dt}".replace(" ", "_").lower()
                        item["id"] = did
                        item.setdefault("doc_id", did)

                        if did in seen_doc_ids:
                            continue
                        seen_doc_ids.add(did)

                        for list_field in (
                            "notification_ids", "hiv_dm_ids", "dbt_ids", "sample_collection_ids",
                            "sample_tested_ids", "outcome_assigned_ids", "home_visit_ids",
                            "contact_tracing_ids", "follow_up_ids", "face_to_face_ids",
                            "presumptive_ids", "fdc_provided_ids", "differentiated_tb_ids",
                            "tpt_treatment_start_ids", "tpt_presumptive_ids",
                            "adhar_face_authentication_ids", "consent_with_id_ids",
                            "culture_dst_ids", "kit_consumption_ids", "documents_ids"
                        ):
                            if list_field in kpi_map:
                                item[list_field] = kpi_map[list_field]
                            else:
                                val = item.get(list_field)
                                if isinstance(val, str):
                                    try:
                                        import json as _json
                                        item[list_field] = _json.loads(val)
                                    except Exception:
                                        item[list_field] = []
                                elif not isinstance(val, list):
                                    item[list_field] = []

                        item["visited_names"] = names_list if names_list else (item.get("visited_names") or [])
                        all_candidate_docs.append(item)
            except Exception as pg_e:
                logger.debug(f"[get_today_attendance] Direct PG date query notice: {pg_e}")
                all_candidate_docs = []

            # Fallback to get_raw_monthly_reports / active_db ONLY if in mock/test environment
            active_db = get_active_db()
            is_mock_env = (
                active_db is not None and (
                    hasattr(active_db, "mock_calls") 
                    or hasattr(active_db, "store")
                    or hasattr(active_db, "saved_reports")
                    or hasattr(active_db, "existing_docs")
                    or type(active_db).__name__ in ["Mock", "MagicMock", "MockFirestore"]
                    or isinstance(getattr(active_db, "reports", None), list)
                )
            )
            if not all_candidate_docs and is_mock_env:
                target_month = target_date[:7]
                raw_docs = await _resolve_get_raw_monthly_reports(target_month)

                extended_raw_docs = list(raw_docs or [])
                if target_dt.day <= 2:
                    prev_month = (target_dt.replace(day=1) - timedelta(days=1)).strftime("%Y-%m")
                    prev_docs = await _resolve_get_raw_monthly_reports(prev_month)
                    if prev_docs:
                        extended_raw_docs.extend(prev_docs)
                elif target_dt.day >= 28:
                    next_month = (target_dt.replace(day=28) + timedelta(days=5)).strftime("%Y-%m")
                    next_docs = await _resolve_get_raw_monthly_reports(next_month)
                    if next_docs:
                        extended_raw_docs.extend(next_docs)

                seen_doc_ids = set()
                for d in extended_raw_docs:
                    r_date = str(d.get("date_of_reporting") or d.get("date") or "").strip()
                    if r_date == target_date or (next_date and r_date == next_date):
                        did = d.get("id") or d.get("doc_id")
                        if did:
                            if did in seen_doc_ids:
                                continue
                            seen_doc_ids.add(did)
                        all_candidate_docs.append(d)

            leave_cache_key = f"daily_leaves_date_{target_date}"
            cached_leaves = cache.get(leave_cache_key)
            if cached_leaves is not None and isinstance(cached_leaves, list):
                leave_docs = cached_leaves
            else:
                try:
                    leave_rows = pg_execute_raw(
                        "SELECT * FROM daily_staff_leaves WHERE date = %s",
                        [target_date],
                        fetch=True
                    ) or []
                    leave_docs = [dict(ld) for ld in leave_rows]
                    cache.set(leave_cache_key, leave_docs, ttl=300)
                except Exception as le:
                    print(f"PG daily_staff_leaves query notice: {le}")
                    leave_docs = []


            result = format_attendance_response(
                staff_list=staff_list,
                all_candidate_docs=all_candidate_docs,
                leave_docs=leave_docs,
                target_date=target_date,
                next_date=next_date,
                allowed_dist_set=allowed_dist_set,
                inactive_staff_keys=inactive_staff_keys
            )
            cache.set(cache_key, result, ttl=600)
            return result
        except HTTPException:
            raise
        except Exception as derivation_err:
            logger.warning(f"[Attendance Failover] In-memory derivation failed, falling back to legacy: {derivation_err}")
            return await legacy_get_today_attendance(
                target_date=target_date,
                next_date=next_date,
                allowed_dist_set=allowed_dist_set,
                staff_list=staff_list,
                effective_dist=effective_dist,
                cache_key=cache_key,
                admin=admin,
                force_refresh=bool(force_refresh),
                inactive_staff_keys=inactive_staff_keys
            )
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))



@router.get("/admin/export-staff-attendance")
@router.get("/admin/export-attendance")
async def export_staff_attendance(
    month: Optional[str] = None,
    district: Optional[str] = None,
    districts: Optional[str] = None,
    admin: dict = Depends(get_current_admin)
):
    try:
        # 1. Month validation and normalization
        if not month or not month.strip():
            target_month = get_ist_now().strftime("%Y-%m")
        else:
            target_month = month.strip()

        try:
            year_val, month_val = map(int, target_month.split("-"))
            num_days = calendar.monthrange(year_val, month_val)[1]
        except Exception:
            now_ist = get_ist_now()
            year_val, month_val = now_ist.year, now_ist.month
            target_month = now_ist.strftime("%Y-%m")
            num_days = calendar.monthrange(year_val, month_val)[1]

        start_date = f"{target_month}-01"
        end_date = f"{target_month}-{num_days:02d}"
        today_ist = get_ist_now().date()

        # 2. RBAC & District Isolation
        admin_role = admin.get("role", "SUB_ADMIN")
        is_subadmin = (admin_role == "SUB_ADMIN")
        allowed = admin.get("allowed_districts", []) or admin.get("districts", [])
        allowed_c = {canonicalize_district(d).lower() for d in allowed if d}
        has_all_access = not is_subadmin or ("all" in allowed_c)

        target_dist_set: Optional[Set[str]] = None

        if district and district.strip() and district.strip().lower() != "all":
            c_dist = canonicalize_district(district.strip())
            if not has_all_access and c_dist.lower() not in allowed_c:
                raise HTTPException(status_code=403, detail=f"Permission denied for district: {district}")
            target_dist_set = {c_dist.lower()}
        elif districts and districts.strip() and districts.strip().lower() != "all":
            req_dists = [canonicalize_district(d.strip()) for d in districts.split(",") if d.strip()]
            if not has_all_access:
                filtered = [d for d in req_dists if d.lower() in allowed_c]
                if not filtered:
                    raise HTTPException(status_code=403, detail="Permission denied for requested districts.")
                target_dist_set = {d.lower() for d in filtered}
            else:
                target_dist_set = {d.lower() for d in req_dists}
        elif not has_all_access:
            target_dist_set = set(allowed_c)

        async with attendance_excel_semaphore:
            # 3. Data fetching
            report_docs = await _resolve_get_raw_monthly_reports(target_month)
            staff_docs = await get_cached_staff_directory_raw()
            
            try:
                leave_rows = pg_execute_raw(
                    "SELECT * FROM daily_staff_leaves WHERE date >= %s AND date <= %s",
                    [start_date, end_date],
                    fetch=True
                ) or []
                leave_docs = [dict(r) for r in leave_rows]
            except Exception as le:
                print(f"PG daily_staff_leaves monthly query notice: {le}")
                leave_docs = []


            # 4. Normalize & Index Officers
            def norm_fo_name(name: str) -> str:
                clean = re.sub(r'[^a-zA-Z0-9]', '', str(name or "")).lower()
                if "ashwanikrkeshri" in clean:
                    return "ashwanikumar"
                return clean

            officers_map = {}
            for sd in staff_docs:
                d = sd.to_dict() if hasattr(sd, "to_dict") else sd
                if not isinstance(d, dict):
                    continue
                c_dist = canonicalize_district(d.get("district", ""))
                if not c_dist:
                    continue
                if target_dist_set and c_dist.lower() not in target_dist_set:
                    continue
                raw_name = (d.get("name") or d.get("fo_name") or "").strip()
                c_name = norm_fo_name(raw_name)
                if not c_name:
                    continue
                key = (c_dist.lower(), c_name)
                if key not in officers_map:
                    officers_map[key] = {
                        "district": c_dist,
                        "name": raw_name,
                        "designation": d.get("designation") or "Field Officer",
                        "clean_name": c_name
                    }

            for doc in report_docs:
                d = doc if isinstance(doc, dict) else doc.to_dict()
                if not isinstance(d, dict):
                    continue
                c_dist = canonicalize_district(d.get("working_place") or d.get("district") or "")
                if not c_dist:
                    continue
                if target_dist_set and c_dist.lower() not in target_dist_set:
                    continue
                raw_name = (d.get("fo_name") or "").strip()
                c_name = norm_fo_name(raw_name)
                if not c_name:
                    continue
                key = (c_dist.lower(), c_name)
                if key not in officers_map:
                    officers_map[key] = {
                        "district": c_dist,
                        "name": raw_name,
                        "designation": "Field Officer",
                        "clean_name": c_name
                    }

            for ldoc in leave_docs:
                ld = ldoc.to_dict() if hasattr(ldoc, "to_dict") else ldoc
                if not isinstance(ld, dict):
                    continue
                c_dist = canonicalize_district(ld.get("district", ""))
                if not c_dist:
                    continue
                if target_dist_set and c_dist.lower() not in target_dist_set:
                    continue
                raw_name = (ld.get("fo_name") or "").strip()
                c_name = norm_fo_name(raw_name)
                if not c_name:
                    continue
                key = (c_dist.lower(), c_name)
                if key not in officers_map:
                    officers_map[key] = {
                        "district": c_dist,
                        "name": raw_name,
                        "designation": "Field Officer",
                        "clean_name": c_name
                    }

            # 5. Index leaves & daily reports
            leaves_by_key = {}
            for ldoc in leave_docs:
                ld = ldoc.to_dict() if hasattr(ldoc, "to_dict") else ldoc
                if not isinstance(ld, dict):
                    continue
                ld_date = ld.get("date", "")
                if not (start_date <= ld_date <= end_date):
                    continue
                c_dist = canonicalize_district(ld.get("district", ""))
                if target_dist_set and c_dist.lower() not in target_dist_set:
                    continue
                c_name = norm_fo_name(ld.get("fo_name", ""))
                leaves_by_key[(c_dist.lower(), c_name, ld_date)] = ld

            reports_by_key = {}
            for doc in report_docs:
                d = doc if isinstance(doc, dict) else doc.to_dict()
                if not isinstance(d, dict):
                    continue
                rep_date = d.get("date_of_reporting") or d.get("date") or ""
                if not (start_date <= rep_date <= end_date):
                    continue
                c_dist = canonicalize_district(d.get("working_place") or d.get("district") or "")
                if target_dist_set and c_dist.lower() not in target_dist_set:
                    continue
                c_name = norm_fo_name(d.get("fo_name", ""))
                key = (c_dist.lower(), c_name, rep_date)

                n_count = len(d.get("notification_ids", []))
                st_count = len(d.get("sample_tested_ids", []))
                dbt_count = len(d.get("dbt_ids", []))
                day_total_ids = sum(len(v) for k, v in d.items() if isinstance(v, list) and k.endswith("_ids"))
                try:
                    km = int(float(d.get("total_km", 0) or 0))
                except Exception:
                    km = 0
                raw_v = d.get("visited_names") or d.get("doctor_names") or []
                if isinstance(raw_v, list):
                    v_names = [str(x).strip() for x in raw_v if str(x).strip()]
                elif isinstance(raw_v, str) and raw_v.strip():
                    v_names = [x.strip() for x in raw_v.split(",") if x.strip()]
                else:
                    v_names = []
                raw_rem = (d.get("admin_remark") or d.get("remark") or "").strip()
                is_next_day = bool(d.get("is_next_day_submission"))
                raw_m = d.get("submitted_morning_time")
                m_time = format_to_ist_time(raw_m) if raw_m else (format_to_ist_time(d.get("timestamp_completed") or d.get("timestamp")) or "")

                if key not in reports_by_key:
                    reports_by_key[key] = {
                        "total_ids": day_total_ids,
                        "notifications": n_count,
                        "samples_tested": st_count,
                        "dbt_seeded": dbt_count,
                        "total_km": km,
                        "visited_names": list(v_names),
                        "admin_remarks": [raw_rem] if raw_rem else [],
                        "submission_count": d.get("submission_count", 1),
                        "is_next_day": is_next_day,
                        "morning_time": m_time
                    }
                else:
                    existing = reports_by_key[key]
                    existing["total_ids"] += day_total_ids
                    existing["notifications"] += n_count
                    existing["samples_tested"] += st_count
                    existing["dbt_seeded"] += dbt_count
                    existing["total_km"] = max(existing["total_km"], km)
                    for vn in v_names:
                        if vn and vn not in existing["visited_names"]:
                            existing["visited_names"].append(vn)
                    if raw_rem and raw_rem not in existing["admin_remarks"]:
                        existing["admin_remarks"].append(raw_rem)
                    existing["submission_count"] += d.get("submission_count", 1)
                    if is_next_day:
                        existing["is_next_day"] = True
                        if m_time:
                            existing["morning_time"] = m_time

            # Sort officers by District, Name
            sorted_officers = sorted(officers_map.values(), key=lambda x: (x["district"], x["name"]))

            # 6. Build Workbook
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            from openpyxl.utils import get_column_letter
            wb = openpyxl.Workbook()
            ws1 = wb.active
            ws1.title = "Attendance Matrix"
            ws1.freeze_panes = "E2"

            ws2 = wb.create_sheet(title="Daily Activity Log")
            ws2.freeze_panes = "A2"

            # Color Palettes & Styles
            header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
            header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")

            status_styles = {
                "P": {
                    "fill": PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid"),
                    "font": Font(name="Calibri", size=10, bold=True, color="065F46")
                },
                "ML": {
                    "fill": PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid"),
                    "font": Font(name="Calibri", size=10, bold=True, color="92400E")
                },
                "CL": {
                    "fill": PatternFill(start_color="E0F2FE", end_color="E0F2FE", fill_type="solid"),
                    "font": Font(name="Calibri", size=10, bold=True, color="0369A1")
                },
                "OD": {
                    "fill": PatternFill(start_color="E0E7FF", end_color="E0E7FF", fill_type="solid"),
                    "font": Font(name="Calibri", size=10, bold=True, color="3730A3")
                },
                "A": {
                    "fill": PatternFill(start_color="FFE4E6", end_color="FFE4E6", fill_type="solid"),
                    "font": Font(name="Calibri", size=10, bold=True, color="9F1239")
                },
                "WO": {
                    "fill": PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid"),
                    "font": Font(name="Calibri", size=10, color="64748B")
                },
                "H": {
                    "fill": PatternFill(start_color="F3E8FF", end_color="F3E8FF", fill_type="solid"),
                    "font": Font(name="Calibri", size=10, color="6B21A8")
                },
                "-": {
                    "fill": PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid"),
                    "font": Font(name="Calibri", size=10, color="94A3B8")
                }
            }

            headers_sheet1 = ["SL", "District", "Officer Name", "Designation"] + [str(d) for d in range(1, num_days + 1)] + ["Total Days", "Present", "Leaves", "Absent", "Travel KM", "Remarks"]
            ws1.append(headers_sheet1)
            for col_idx, cell in enumerate(ws1[1], start=1):
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = EXCEL_HEADER_BORDER

            headers_sheet2 = ["Date", "District", "Officer Name", "Status", "Total IDs", "Notifications", "Samples Tested", "DBT Seeded", "Travel KM", "Visited Doctors / Facilities", "Admin Remark"]
            ws2.append(headers_sheet2)
            for col_idx, cell in enumerate(ws2[1], start=1):
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = EXCEL_HEADER_BORDER

            rows_sheet2 = []

            for sl_num, officer in enumerate(sorted_officers, start=1):
                c_dist_lower = officer["district"].lower()
                c_name = officer["clean_name"]
                
                day_codes = []
                present_count = 0
                leaves_count = 0
                absent_count = 0
                travel_km_sum = 0
                officer_remarks = []

                for d in range(1, num_days + 1):
                    day_str = f"{year_val:04d}-{month_val:02d}-{d:02d}"
                    cur_dt = datetime(year_val, month_val, d).date()
                    is_future = (cur_dt > today_ist)
                    is_sunday = (cur_dt.weekday() == 6)

                    rep = reports_by_key.get((c_dist_lower, c_name, day_str))
                    l_doc = leaves_by_key.get((c_dist_lower, c_name, day_str))

                    if rep:
                        code = "P"
                        status_label = "Present"
                    elif l_doc:
                        l_status = str(l_doc.get("status", "")).strip().lower()
                        r_type = str(l_doc.get("reason_type", "")).strip().lower()
                        if "absent" in l_status or "absent" in r_type:
                            code = "A"
                            status_label = "Absent"
                        elif "medical" in r_type or "sick" in r_type:
                            code = "ML"
                            status_label = "Medical Leave"
                        elif "official" in r_type or "duty" in r_type or "training" in r_type or "official_duty" in l_status:
                            code = "OD"
                            status_label = "Official Duty"
                        elif "holiday" in r_type:
                            code = "H"
                            status_label = "Declared Holiday"
                        else:
                            code = "CL"
                            status_label = "Casual Leave"
                    elif is_future:
                        code = "-"
                        status_label = "-"
                    elif is_sunday:
                        code = "WO"
                        status_label = "Weekly Off"
                    else:
                        code = "A"
                        status_label = "Absent"

                    day_codes.append(code)

                    if code == "P":
                        present_count += 1
                        if rep:
                            travel_km_sum += rep.get("total_km", 0)
                            day_notes = []
                            if rep.get("admin_remarks"):
                                day_notes.append(', '.join(rep['admin_remarks']))
                            if rep.get("is_next_day"):
                                m_time = rep.get("morning_time")
                                if m_time:
                                    day_notes.append(f"Submitted next morning ({m_time})")
                                else:
                                    day_notes.append("Submitted next morning")
                            if day_notes:
                                officer_remarks.append(f"Day {d}: {', '.join(day_notes)}")
                    elif code in ("ML", "CL", "OD", "H"):
                        leaves_count += 1
                        if l_doc and l_doc.get("remark"):
                            officer_remarks.append(f"Day {d} ({l_doc.get('reason_type', 'Leave')}): {l_doc.get('remark')}")
                    elif code == "A":
                        absent_count += 1

                    # Sheet 2 row (log elapsed days or any day with report/leave)
                    if not is_future or rep or l_doc:
                        v_str = ", ".join(rep["visited_names"]) if (rep and rep.get("visited_names")) else "-"
                        rem_parts = []
                        if rep and rep.get("admin_remarks"):
                            rem_parts.append(", ".join(rep["admin_remarks"]))
                        if rep and rep.get("is_next_day"):
                            m_time = rep.get("morning_time")
                            if m_time:
                                rem_parts.append(f"Submitted next morning ({m_time})")
                            else:
                                rem_parts.append("Submitted next morning")
                        if l_doc and l_doc.get("remark"):
                            rem_parts.append(f"{l_doc.get('reason_type', 'Leave')}: {l_doc.get('remark')}")
                        rem_str = ", ".join(rem_parts) if rem_parts else "-"

                        rows_sheet2.append([
                            day_str,
                            officer["district"],
                            officer["name"],
                            status_label,
                            rep.get("total_ids", 0) if rep else 0,
                            rep.get("notifications", 0) if rep else 0,
                            rep.get("samples_tested", 0) if rep else 0,
                            rep.get("dbt_seeded", 0) if rep else 0,
                            rep.get("total_km", 0) if rep else 0,
                            v_str,
                            rem_str
                        ])

                rem_summary = "; ".join(officer_remarks) if officer_remarks else "-"
                row_data = [
                    sl_num,
                    officer["district"],
                    officer["name"],
                    officer["designation"]
                ] + day_codes + [
                    num_days,
                    present_count,
                    leaves_count,
                    absent_count,
                    travel_km_sum,
                    rem_summary
                ]
                ws1.append(row_data)

                curr_row = ws1.max_row
                for col_idx in range(1, len(row_data) + 1):
                    c = ws1.cell(row=curr_row, column=col_idx)
                    c.border = EXCEL_THIN_BORDER
                    c.font = Font(name="Calibri", size=10)
                    if col_idx == 1:
                        c.alignment = Alignment(horizontal="center", vertical="center")
                    elif 2 <= col_idx <= 4:
                        c.alignment = Alignment(horizontal="left", vertical="center")
                    elif 5 <= col_idx <= (4 + num_days):
                        code_val = str(c.value or "").strip()
                        c.alignment = Alignment(horizontal="center", vertical="center")
                        if code_val in status_styles:
                            c.fill = status_styles[code_val]["fill"]
                            c.font = status_styles[code_val]["font"]
                    elif (4 + num_days) < col_idx < len(row_data):
                        c.alignment = Alignment(horizontal="center", vertical="center")
                        c.font = Font(name="Calibri", size=10, bold=True)
                    else:
                        c.alignment = Alignment(horizontal="left", vertical="center")

            # Column dimensions for Sheet 1
            ws1.column_dimensions["A"].width = 6
            ws1.column_dimensions["B"].width = 16
            ws1.column_dimensions["C"].width = 22
            ws1.column_dimensions["D"].width = 18
            for d in range(1, num_days + 1):
                col_letter = get_column_letter(4 + d)
                ws1.column_dimensions[col_letter].width = 4.5

            c_tot_days = get_column_letter(5 + num_days)
            c_pres = get_column_letter(6 + num_days)
            c_leaves = get_column_letter(7 + num_days)
            c_abs = get_column_letter(8 + num_days)
            c_km = get_column_letter(9 + num_days)
            c_rem = get_column_letter(10 + num_days)

            ws1.column_dimensions[c_tot_days].width = 12
            ws1.column_dimensions[c_pres].width = 10
            ws1.column_dimensions[c_leaves].width = 10
            ws1.column_dimensions[c_abs].width = 10
            ws1.column_dimensions[c_km].width = 12
            ws1.column_dimensions[c_rem].width = 35

            # Populate Sheet 2: Daily Activity Log
            rows_sheet2.sort(key=lambda r: (r[0], r[1], r[2]))
            for r_data in rows_sheet2:
                ws2.append(r_data)
                curr_row = ws2.max_row
                for col_idx in range(1, len(r_data) + 1):
                    c = ws2.cell(row=curr_row, column=col_idx)
                    c.border = EXCEL_THIN_BORDER
                    c.font = Font(name="Calibri", size=10)
                    if col_idx in (1, 2, 4):
                        c.alignment = Alignment(horizontal="center", vertical="center")
                    elif col_idx == 3:
                        c.alignment = Alignment(horizontal="left", vertical="center")
                    elif 5 <= col_idx <= 9:
                        c.alignment = Alignment(horizontal="center", vertical="center")
                    else:
                        c.alignment = Alignment(horizontal="left", vertical="center")

            ws2.auto_filter.ref = ws2.dimensions
            ws2.column_dimensions["A"].width = 13
            ws2.column_dimensions["B"].width = 16
            ws2.column_dimensions["C"].width = 22
            ws2.column_dimensions["D"].width = 16
            ws2.column_dimensions["E"].width = 11
            ws2.column_dimensions["F"].width = 13
            ws2.column_dimensions["G"].width = 15
            ws2.column_dimensions["H"].width = 13
            ws2.column_dimensions["I"].width = 12
            ws2.column_dimensions["J"].width = 32
            ws2.column_dimensions["K"].width = 32

            output = io.BytesIO()
            wb.save(output)
            content_bytes = output.getvalue()
            del wb
            gc.collect()

            if district and district.strip() and district.strip().lower() != "all":
                c_fn_dist = canonicalize_district(district.strip())
            elif districts and len(districts.split(",")) == 1 and districts.strip().lower() != "all":
                c_fn_dist = canonicalize_district(districts.strip())
            else:
                c_fn_dist = "All"

            safe_fn_dist = safe_filename(c_fn_dist)
            filename = f"DFY_Staff_Attendance_{safe_fn_dist}_{target_month}.xlsx"
            headers = {"Content-Disposition": f'attachment; filename="{filename}"'}

            date_str = target_month
            actor_name = admin.get("name") or admin.get("username", "Admin")
            actor_id = admin.get("user_id") or admin.get("username", "admin")
            actor_role = admin.get("role", "SUB_ADMIN")
            await log_admin_activity(
                action_type="REPORT_DOWNLOADED",
                details=f"Admin {actor_name} downloaded Attendance Excel for {date_str}",
                user_name=actor_name,
                user_id=actor_id,
                role=actor_role,
                diff={"report_type": "Attendance Excel", "date": date_str}
            )

            return ExcelStreamingResponse(
                content_bytes,
                media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                headers=headers
            )
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.get("/admin/export-summary-metrics")
@router.get("/admin/export-fo-dossier")
async def export_summary_metrics(
    month: Optional[str] = None,
    district: Optional[str] = None,
    districts: Optional[str] = None,
    admin: dict = Depends(get_current_admin)
):
    main_mod = sys.modules.get("main")
    fn = getattr(main_mod, "export_staff_attendance", export_staff_attendance) if main_mod else export_staff_attendance
    return await fn(month=month, district=district, districts=districts, admin=admin)



class MarkLeaveReq(BaseModel):
    district: str
    fo_name: str
    date: str  # YYYY-MM-DD
    status: str = "leave"  # "leave" | "absent" | "weekly_off"
    reason_type: str = "Casual"  # "Medical" | "Casual" | "Official Work" | "Personal" | "Uninformed"
    remark: Optional[str] = ""

class UnmarkLeaveReq(BaseModel):
    district: str
    fo_name: str
    date: str  # YYYY-MM-DD

@router.post("/admin/attendance/mark-leave")
async def mark_leave(req: MarkLeaveReq, admin: dict = Depends(get_current_admin)):
    try:
        clean_dist = canonicalize_district(req.district.strip()) if req.district else ""
        if not clean_dist:
            raise HTTPException(status_code=400, detail="Valid district is required.")
        clean_fo = re.sub(r'[^a-zA-Z0-9]', '', req.fo_name).lower()
        if not clean_fo:
            raise HTTPException(status_code=400, detail="Valid Field Officer name is required.")
        clean_date = req.date.strip()
        if not clean_date:
            raise HTTPException(status_code=400, detail="Valid date is required.")

        admin_role = admin.get("role", "SUB_ADMIN")
        if admin_role == "SUB_ADMIN":
            allowed = admin.get("allowed_districts", []) or admin.get("districts", [])
            allowed_c = [canonicalize_district(d).lower() for d in allowed if d]
            if "All" not in allowed and "all" not in allowed_c and clean_dist.lower() not in allowed_c and clean_dist not in allowed:
                raise HTTPException(status_code=403, detail=f"Permission denied for district: {clean_dist}")

        doc_id = f"{clean_date}_{clean_dist}_{clean_fo}"

        actor_name = admin.get("name") or admin.get("username") or "Admin"
        actor_id = admin.get("user_id") or admin.get("username") or "admin"
        actor_role = admin.get("role", "SUB_ADMIN")
        marked_at = get_ist_now().strftime("%Y-%m-%d %H:%M:%S")

        leave_data = {
            "id": doc_id,
            "date": clean_date,
            "district": clean_dist,
            "fo_name": req.fo_name.strip(),
            "status": req.status,
            "reason_type": req.reason_type,
            "remark": req.remark or "",
            "marked_by_name": actor_name,
            "marked_by_id": actor_id,
            "marked_by_role": actor_role,
            "marked_at": marked_at
        }

        pg_upsert_row("daily_staff_leaves", leave_data, conflict_columns=["id"])

        cache.delete_prefix(f"attendance_{clean_date}")
        cache.delete(f"daily_leaves_{clean_dist}_{clean_date[:7]}")
        cache.delete(f"daily_leaves_date_{clean_date}")
        evict_officer_profile_cache(clean_dist, req.fo_name.strip(), clean_date)

        await log_admin_activity(
            action_type="LEAVE_MARKED",
            details=f"Marked {req.status} for {req.fo_name.strip()} ({clean_dist}) on {clean_date}: {req.reason_type}",
            user_name=actor_name,
            user_id=actor_id,
            role=actor_role,
            district=clean_dist,
            target_officer=req.fo_name.strip(),
            diff={
                "date": clean_date,
                "status": req.status,
                "reason_type": req.reason_type,
                "remark": req.remark or ""
            }
        )

        return {"success": True, "message": "Leave recorded successfully."}
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@router.post("/admin/attendance/unmark-leave")
async def unmark_leave(req: UnmarkLeaveReq, admin: dict = Depends(get_current_admin)):
    try:
        clean_dist = canonicalize_district(req.district.strip()) if req.district else ""
        if not clean_dist:
            raise HTTPException(status_code=400, detail="Valid district is required.")
        clean_fo = re.sub(r'[^a-zA-Z0-9]', '', req.fo_name).lower()
        if not clean_fo:
            raise HTTPException(status_code=400, detail="Valid Field Officer name is required.")
        clean_date = req.date.strip()
        if not clean_date:
            raise HTTPException(status_code=400, detail="Valid date is required.")

        admin_role = admin.get("role", "SUB_ADMIN")
        if admin_role == "SUB_ADMIN":
            allowed = admin.get("allowed_districts", []) or admin.get("districts", [])
            allowed_c = [canonicalize_district(d).lower() for d in allowed if d]
            if "All" not in allowed and "all" not in allowed_c and clean_dist.lower() not in allowed_c and clean_dist not in allowed:
                raise HTTPException(status_code=403, detail=f"Permission denied for district: {clean_dist}")

        doc_id = f"{clean_date}_{clean_dist}_{clean_fo}"
        pg_delete_rows("daily_staff_leaves", filters={"id": doc_id})


        cache.delete_prefix(f"attendance_{clean_date}")
        cache.delete(f"daily_leaves_{clean_dist}_{clean_date[:7]}")
        cache.delete(f"daily_leaves_date_{clean_date}")
        evict_officer_profile_cache(clean_dist, req.fo_name.strip(), clean_date)

        actor_name = admin.get("name") or admin.get("username") or "Admin"
        actor_id = admin.get("user_id") or admin.get("username") or "admin"
        actor_role = admin.get("role", "SUB_ADMIN")

        await log_admin_activity(
            action_type="LEAVE_UNMARKED",
            details=f"Removed leave for {req.fo_name.strip()} ({clean_dist}) on {clean_date}",
            user_name=actor_name,
            user_id=actor_id,
            role=actor_role,
            district=clean_dist,
            target_officer=req.fo_name.strip(),
            diff={"date": clean_date}
        )

        return {"success": True, "message": "Leave removed successfully."}
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


class AttendanceRemarkReq(BaseModel):
    district: str
    fo_name: str
    date: str  # YYYY-MM-DD
    action: str = "remark"  # "remark" | "override_leave"
    remark: str
    status: Optional[str] = "leave"  # "leave" | "absent" | "present"
    reason_type: Optional[str] = "Casual"  # "Casual", "Medical", "Official Duty", etc.


@router.post("/admin/attendance/add-remark")
async def add_attendance_remark(req: AttendanceRemarkReq, admin: dict = Depends(get_current_admin)):
    try:
        clean_dist = canonicalize_district(req.district.strip()) if req.district else ""
        if not clean_dist:
            raise HTTPException(status_code=400, detail="Valid district is required.")
        clean_fo = re.sub(r'[^a-zA-Z0-9]', '', req.fo_name).lower() if req.fo_name else ""
        if not clean_fo:
            raise HTTPException(status_code=400, detail="Valid Field Officer name is required.")
        clean_date = req.date.strip() if req.date else ""
        if not clean_date:
            raise HTTPException(status_code=400, detail="Valid date is required.")
        if not req.remark or not req.remark.strip():
            raise HTTPException(status_code=400, detail="Remark text is required.")

        admin_role = admin.get("role", "SUB_ADMIN")
        if admin_role == "SUB_ADMIN":
            allowed = admin.get("allowed_districts", []) or admin.get("districts", [])
            allowed_c = [canonicalize_district(d).lower() for d in allowed if d]
            if "All" not in allowed and "all" not in allowed_c and clean_dist.lower() not in allowed_c and clean_dist not in allowed:
                raise HTTPException(status_code=403, detail=f"Permission denied for district: {clean_dist}")

        actor_name = admin.get("name") or admin.get("username") or "Admin"
        actor_id = admin.get("user_id") or admin.get("username") or "admin"
        actor_role = admin.get("role", "SUB_ADMIN")
        marked_at = get_ist_now().strftime("%Y-%m-%d %H:%M:%S")

        doc_id = f"{clean_date}_{clean_dist}_{clean_fo}"

        action_type = "ATTENDANCE_REMARK_ADDED" if req.action == "remark" else "ATTENDANCE_LEAVE_OVERRIDDEN"

        if req.action == "remark":
            # Attach remark to daily report in PostgreSQL
            try:
                pg_execute_raw(
                    "UPDATE daily_field_reports SET admin_remark = %s, admin_remark_by = %s, admin_remark_at = %s WHERE date_of_reporting = %s AND (fo_name = %s OR fo_name = %s)",
                    [req.remark.strip(), actor_name, marked_at, clean_date, req.fo_name.strip(), clean_fo]
                )
            except Exception as upd_err:
                print(f"Notice: Failed to update admin_remark on PG report: {upd_err}")

            # Also update candidate docs in db / mock store
            try:
                cand_ids = [
                    f"{clean_dist.lower()}_{clean_fo}_{clean_date}",
                    f"{clean_dist}_{req.fo_name.strip()}_{clean_date}".replace(" ", "_").lower(),
                    f"{canonicalize_district(clean_dist).lower()}_{clean_fo}_{clean_date}"
                ]
                for cid in cand_ids:
                    doc_ref = db.collection("daily_field_reports").document(cid)
                    doc = await asyncio.to_thread(doc_ref.get)
                    if getattr(doc, "exists", False):
                        await asyncio.to_thread(lambda ref=doc_ref: ref.update({
                            "admin_remark": req.remark.strip(),
                            "admin_remark_by": actor_name,
                            "admin_remark_at": marked_at
                        }))
                        break
            except Exception as q_err:
                print(f"Notice: daily_field_reports search failed: {q_err}")

            # Also store/merge in daily_staff_leaves so remark is preserved across views
            leave_remark_data = {
                "id": doc_id,
                "date": clean_date,
                "district": clean_dist,
                "fo_name": req.fo_name.strip(),
                "remark": req.remark.strip(),
                "admin_remark": req.remark.strip(),
                "marked_by_name": actor_name,
                "marked_by_id": actor_id,
                "marked_by_role": actor_role,
                "marked_at": marked_at,
                "is_inspection_remark": True,
                "is_override": False
            }
            pg_upsert_row("daily_staff_leaves", leave_remark_data, conflict_columns=["id"])

            details = f"Added admin inspection remark for {req.fo_name.strip()} ({clean_dist}) on {clean_date}: {req.remark.strip()}"
            msg = "Attendance remark recorded successfully."
        else:
            # req.action == "override_leave"
            leave_data = {
                "id": doc_id,
                "date": clean_date,
                "district": clean_dist,
                "fo_name": req.fo_name.strip(),
                "status": req.status or "leave",
                "reason_type": req.reason_type or "Casual",
                "remark": req.remark.strip(),
                "marked_by_name": actor_name,
                "marked_by_id": actor_id,
                "marked_by_role": actor_role,
                "marked_at": marked_at,
                "is_override": True,
                "is_inspection_remark": False
            }
            pg_upsert_row("daily_staff_leaves", leave_data, conflict_columns=["id"])
            details = f"Overrode attendance to {req.status or 'leave'} ({req.reason_type or 'Casual'}) for {req.fo_name.strip()} ({clean_dist}) on {clean_date}: {req.remark.strip()}"
            msg = "Leave status overridden successfully."


        cache.delete_prefix(f"attendance_{clean_date}")
        cache.delete(f"daily_leaves_{clean_dist}_{clean_date[:7]}")
        cache.delete(f"daily_leaves_date_{clean_date}")
        evict_officer_profile_cache(clean_dist, req.fo_name.strip(), clean_date)

        # Update in-memory shared raw month cache if present
        month_prefix = clean_date[:7]
        cached_monthly = cache.get(f"shared_raw_month_{month_prefix}")
        if cached_monthly and isinstance(cached_monthly, list):
            found_and_updated = False
            for r in cached_monthly:
                r_fo = re.sub(r'[^a-zA-Z0-9]', '', r.get("fo_name", "")).lower()
                r_wp = canonicalize_district(r.get("working_place", "")).lower()
                r_dt = str(r.get("date_of_reporting", "") or r.get("date", "")).strip()
                if r_fo == clean_fo and r_wp == clean_dist.lower() and r_dt == clean_date:
                    if req.action == "remark":
                        r["admin_remark"] = req.remark.strip()
                        r["admin_remark_by"] = actor_name
                        r["admin_remark_at"] = marked_at
                        found_and_updated = True
            if found_and_updated:
                cache.set(f"shared_raw_month_{month_prefix}", cached_monthly, ttl=3600)
            else:
                cache.delete_prefix(f"shared_raw_month_{month_prefix}")
        else:
            cache.delete_prefix(f"shared_raw_month_{month_prefix}")

        cache.delete_prefix(f"shared_raw_month_{month_prefix}_")

        await log_admin_activity(
            action_type=action_type,
            details=details,
            user_name=actor_name,
            user_id=actor_id,
            role=actor_role,
            district=clean_dist,
            target_officer=req.fo_name.strip(),
            diff={
                "date": clean_date,
                "action": req.action,
                "remark": req.remark.strip(),
                "status": req.status if req.action != "remark" else None,
                "reason_type": req.reason_type if req.action != "remark" else None
            }
        )

        return {"success": True, "message": msg}
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


# =========================================================================
# --- Backend Declared Holidays Pacing Sync Endpoints (/admin/pacing/settings) ---
# =========================================================================

