import calendar
import gc
import io
import json
import logging
import re
from datetime import datetime, timedelta
from typing import Optional, List, Dict, Any, Union

from fastapi import APIRouter, HTTPException, Depends, Header, Query
from pydantic import BaseModel, Field
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from backend.core.styles import (
    EXCEL_NAVY_HEADER_FILL,
    EXCEL_WHITE_BOLD_FONT,
    EXCEL_HEADER_BORDER,
    EXCEL_THIN_BORDER,
    EXCEL_TOTAL_ROW_BORDER,
    ExcelStreamingResponse,
    safe_filename
)
from backend.core.cache import cache
from backend.core import security
from backend.core.security import get_current_admin
from backend.core.helpers import (
    canonicalize_district,
    is_officer_name_match
)
from backend.core.master_ledger import (
    get_cached_staff_directory_raw,
    get_directory
)
from backend.core.supabase import (
    pg_execute_raw,
    get_active_db,
    get_db_connection
)

logger = logging.getLogger("travel_allowance")
router = APIRouter(tags=["Travel Allowance"])


def get_current_user(
    authorization: Optional[str] = Header(None),
    token: Optional[str] = Query(None)
) -> dict:
    """Dependency delegating to security.get_current_user."""
    return security.get_current_user(authorization=authorization, token=token)


def normalize_staff_slug(name: str) -> str:
    """Creates a URL/ID safe slug for staff names."""
    if not name:
        return ""
    return re.sub(r'[^a-z0-9]+', '_', name.strip().lower()).strip('_')


def clean_alphanumeric(s: str) -> str:
    """Normalizes string to alphanumeric characters only."""
    return re.sub(r'[^a-z0-9]', '', (s or "").lower())


def check_district_access(admin: dict, district: str) -> str:
    """
    Validates Sub-Admin access to a specific district.
    Super Admins, Admins, and Main Incharges have statewide access.
    """
    canon_dist = canonicalize_district(district)
    role = admin.get("role", "")
    if role in ("SUPER_ADMIN", "MAIN_INCHARGE", "ADMIN"):
        return canon_dist

    allowed = admin.get("allowed_districts", [])
    if "All" in allowed:
        return canon_dist

    canon_allowed = [canonicalize_district(d) for d in allowed]
    if canon_dist not in canon_allowed:
        raise HTTPException(
            status_code=403,
            detail="Access denied: You do not have permission for this district."
        )
    return canon_dist


# --- Pydantic Request Models ---

class TaRateUpdateReq(BaseModel):
    rate_per_km: float


class TaPrefillAccessReq(BaseModel):
    admin_id: Optional[int] = None
    username: Optional[str] = None


class TaDayReading(BaseModel):
    day: int
    date: str
    morning_km: float = 0.0
    evening_km: float = 0.0
    total_km: float = 0.0
    visited_names: str = ""
    purpose: str = ""
    is_manual_override: bool = False
    admin_remarks: str = ""


class TaLogSaveReq(BaseModel):
    month: str
    district: str
    staff_name: str
    staff_key: Optional[str] = ""
    designation: Optional[str] = "Field Officer"
    deduction_amount: Optional[float] = 0.0
    deduction_reason: Optional[str] = ""
    admin_remarks: Optional[str] = ""
    days: List[Dict[str, Any]] = []


class TaPrefillReq(BaseModel):
    month: str
    district: str
    staff_name: Optional[str] = None
    staff_key: Optional[str] = None


class TaStaffActionReq(BaseModel):
    staff_key: str
    action: Optional[str] = None
    revert_reason: Optional[str] = None


class TaSubmitRosterReq(BaseModel):
    month: str
    district: str
    staff_keys: Optional[List[str]] = None


class TaPassStaffReq(BaseModel):
    month: str
    district: str
    staff_key: str
    staff_keys: Optional[List[str]] = None


class TaRevertStaffReq(BaseModel):
    month: str
    district: str
    staff_key: str
    revert_reason: str


class TaUnlockStaffReq(BaseModel):
    month: str
    district: str
    staff_key: str


class TaDisputeReq(BaseModel):
    month: str
    district: str
    fo_name: str
    dispute_reason: Optional[str] = None
    reason: Optional[str] = None


class TaResolveDisputeReq(BaseModel):
    month: str
    district: str
    staff_key: str
    action: str
    resolution_remarks: Optional[str] = ""


# --- Financial & Normalization Helpers ---

def calculate_log_totals(
    days: list,
    rate_per_km: float = 4.0,
    deduction_amount: float = 0.0
) -> dict:
    """
    Computes total distance, gross allowance, deduction, and final payable amount.
    Guarantees payable amount never drops below 0.0.
    """
    total_km = 0.0
    for d in days:
        if isinstance(d, dict):
            m = float(d.get("morning_km") or 0.0)
            e = float(d.get("evening_km") or 0.0)
            calc_diff = max(0.0, e - m)
            is_manual = bool(d.get("is_manual_override") or d.get("is_override"))
            manual_t = d.get("manual_total_km")
            t = d.get("total_km")

            if is_manual and manual_t is not None and str(manual_t).strip() != "":
                try:
                    total_km += max(0.0, float(manual_t))
                except (ValueError, TypeError):
                    total_km += calc_diff
            elif t is not None and float(t or 0.0) > 0:
                try:
                    total_km += float(t)
                except (ValueError, TypeError):
                    total_km += calc_diff
            else:
                total_km += calc_diff
        elif hasattr(d, "total_km"):
            m = float(getattr(d, "morning_km", 0.0) or 0.0)
            e = float(getattr(d, "evening_km", 0.0) or 0.0)
            calc_diff = max(0.0, e - m)
            is_manual = bool(getattr(d, "is_manual_override", False) or getattr(d, "is_override", False))
            manual_t = getattr(d, "manual_total_km", None)
            t = getattr(d, "total_km", None)

            if is_manual and manual_t is not None and str(manual_t).strip() != "":
                try:
                    total_km += max(0.0, float(manual_t))
                except (ValueError, TypeError):
                    total_km += calc_diff
            elif t is not None and float(t or 0.0) > 0:
                try:
                    total_km += float(t)
                except (ValueError, TypeError):
                    total_km += calc_diff
            else:
                total_km += calc_diff
        else:
            m = float(getattr(d, "morning_km", 0.0) or 0.0)
            e = float(getattr(d, "evening_km", 0.0) or 0.0)
            total_km += max(0.0, e - m)

    total_km = round(total_km, 2)
    rate = float(rate_per_km if rate_per_km is not None else 4.0)
    try:
        ded = max(0.0, float(deduction_amount if deduction_amount is not None else 0.0))
    except (ValueError, TypeError):
        ded = 0.0
    gross_amount = round(total_km * rate, 2)
    final_payable_amount = round(max(0.0, gross_amount - ded), 2)

    return {
        "total_km": total_km,
        "gross_amount": gross_amount,
        "deduction_amount": ded,
        "final_payable_amount": final_payable_amount
    }


def normalize_days_to_list(raw_days_or_logs: Any, month: str = "") -> List[Dict[str, Any]]:
    """
    Standardizes daily logs into a uniform 28-31 day array matching month length.
    """
    year = 2026
    month_num = 10
    if month and isinstance(month, str) and "-" in month:
        try:
            parts = month.split("-")
            year = int(parts[0])
            month_num = int(parts[1])
        except Exception as e:
            logger.debug(f"month parsing fallback: {e}")
    _, num_days = calendar.monthrange(year, month_num)

    day_map: Dict[int, Dict[str, Any]] = {}
    if isinstance(raw_days_or_logs, dict):
        for k, v in raw_days_or_logs.items():
            if not isinstance(v, dict):
                continue
            day_num = None
            if "-" in str(k):
                try:
                    day_num = int(str(k).split("-")[2])
                except Exception as e:
                    logger.debug(f"day key parsing fallback: {e}")
            elif str(k).isdigit():
                day_num = int(k)
            else:
                day_num = v.get("day")
            if day_num:
                day_map[int(day_num)] = v
    elif isinstance(raw_days_or_logs, list):
        for idx, d in enumerate(raw_days_or_logs, start=1):
            if isinstance(d, dict):
                day_num = d.get("day") or idx
                try:
                    day_map[int(day_num)] = d
                except Exception as e:
                    logger.debug(f"day list map fallback: {e}")

    normalized: List[Dict[str, Any]] = []
    for day in range(1, num_days + 1):
        date_str = f"{year:04d}-{month_num:02d}-{day:02d}"
        d = day_map.get(day, {})
        m_km = float(d.get("morning_km") if d.get("morning_km") is not None else (d.get("initial_reading") or 0.0))
        e_km = float(d.get("evening_km") if d.get("evening_km") is not None else (d.get("final_reading") or 0.0))
        calc_diff = max(0.0, e_km - m_km)
        override = bool(d.get("is_manual_override") or d.get("is_override") or False)
        manual_t = d.get("manual_total_km")

        if override and manual_t is not None and str(manual_t).strip() != "":
            try:
                t_km = max(0.0, float(manual_t))
            except (ValueError, TypeError):
                t_km = calc_diff
        elif d.get("total_km") is not None and float(d.get("total_km") or 0.0) > 0:
            try:
                t_km = float(d.get("total_km"))
            except (ValueError, TypeError):
                t_km = calc_diff
        else:
            t_km = calc_diff
        visited = str(d.get("visited_names") or d.get("to_location") or d.get("places_visited") or "")
        purpose = str(d.get("purpose") or d.get("remarks") or "")
        admin_remarks = str(d.get("admin_remarks") or "")

        normalized.append({
            "day": day,
            "date": str(d.get("date") or date_str),
            "morning_km": round(m_km, 2),
            "evening_km": round(e_km, 2),
            "total_km": round(t_km, 2),
            "visited_names": visited,
            "purpose": purpose,
            "is_manual_override": override,
            "admin_remarks": admin_remarks
        })

    return normalized


def apply_staff_status_transition(
    record: dict,
    action: str,
    actor_role: str,
    actor_name: str,
    reason: str = ""
) -> dict:
    """
    State machine transitions for staff TA records:
    - PASS: SUBMITTED -> APPROVED, is_locked = True
    - REVERT: [SUBMITTED, APPROVED] -> REVERTED, is_locked = False
    - UNLOCK: Any -> REVERTED, is_locked = False
    - SUBMIT: [DRAFT, REVERTED] -> SUBMITTED
    """
    iso_now = datetime.utcnow().isoformat()
    rec = dict(record)
    act = (action or "").strip().upper()
    current_status = rec.get("status") or "DRAFT"

    if act == "PASS":
        if current_status != "SUBMITTED":
            raise HTTPException(
                status_code=400,
                detail=f"Cannot pass record in '{current_status}' status. Record must be SUBMITTED."
            )
        if actor_role not in ("MAIN_INCHARGE", "SUPER_ADMIN", "ADMIN"):
            raise HTTPException(
                status_code=403,
                detail="Only Incharge or Super Admin can pass staff records."
            )
        rec["status"] = "APPROVED"
        rec["is_locked"] = True
        rec["approved_at"] = iso_now
        rec["approved_by"] = actor_name
        return rec

    elif act == "REVERT":
        if current_status not in ("SUBMITTED", "APPROVED"):
            raise HTTPException(
                status_code=400,
                detail=f"Cannot revert record in '{current_status}' status. Record must be SUBMITTED or APPROVED."
            )
        if actor_role not in ("MAIN_INCHARGE", "SUPER_ADMIN", "ADMIN"):
            raise HTTPException(
                status_code=403,
                detail="Only Incharge or Super Admin can revert staff records."
            )
        rec["status"] = "REVERTED"
        rec["is_locked"] = False
        rec["reverted_at"] = iso_now
        rec["reverted_by"] = actor_name
        rec["revert_reason"] = reason or ""
        return rec

    elif act == "UNLOCK":
        if actor_role not in ("MAIN_INCHARGE", "SUPER_ADMIN", "ADMIN"):
            raise HTTPException(
                status_code=403,
                detail="Only Incharge or Super Admin can unlock staff records."
            )
        rec["is_locked"] = False
        rec["status"] = "REVERTED"
        rec["unlocked_at"] = iso_now
        rec["unlocked_by"] = actor_name
        return rec

    elif act == "SUBMIT":
        if current_status not in ("DRAFT", "REVERTED", None, ""):
            raise HTTPException(
                status_code=400,
                detail=f"Cannot submit record in '{current_status}' status. Only DRAFT or REVERTED records can be submitted."
            )
        if actor_role not in ("SUB_ADMIN", "SUPER_ADMIN", "ADMIN"):
            raise HTTPException(
                status_code=403,
                detail="Only Sub-Admin or Super Admin can submit staff records."
            )
        rec["status"] = "SUBMITTED"
        rec["is_locked"] = True
        rec["submitted_at"] = iso_now
        rec["submitted_by"] = actor_name
        return rec

    else:
        raise HTTPException(
            status_code=400,
            detail=f"Unknown transition action '{action}'."
        )


def validate_edit_permission(record: dict, user_role: str) -> tuple[bool, str]:
    """
    Enforces RBAC and record lock constraints for editing TA readings:
    - Super Admin: global override allowed.
    - Locked record: cannot be edited unless unlocked.
    - Main Incharge: read-only inspection.
    - Sub-Admin: cannot edit if record is locked or in SUBMITTED/APPROVED state.
    """
    if user_role == "SUPER_ADMIN":
        return True, ""

    if bool(record.get("is_locked", False)):
        return False, "Record is locked by Incharge. Request Incharge to unlock."

    if user_role == "MAIN_INCHARGE":
        return False, "Incharge role is read-only for inspection and approval. Edits must be made by Sub-Admin."

    status = record.get("status", "DRAFT")
    if user_role == "SUB_ADMIN" and status in ("APPROVED", "SUBMITTED"):
        return False, f"Cannot edit record in {status} state. Request Incharge to unlock."

    return True, ""


def is_dispute_window_open(approved_at: Optional[str], now_dt: Optional[datetime] = None) -> bool:
    """Checks if an approved record is within the 24-hour dispute window."""
    if not approved_at or not str(approved_at).strip():
        return False

    try:
        clean_str = str(approved_at).strip()
        if clean_str.endswith("Z"):
            clean_str = clean_str[:-1]
        approved_dt = datetime.fromisoformat(clean_str)

        if now_dt is None:
            now_dt = datetime.utcnow()

        if approved_dt.tzinfo is not None and now_dt.tzinfo is None:
            approved_dt = approved_dt.replace(tzinfo=None)
        elif approved_dt.tzinfo is None and now_dt.tzinfo is not None:
            now_dt = now_dt.replace(tzinfo=None)

        deadline = approved_dt + timedelta(hours=24)
        return now_dt <= deadline
    except Exception as e:
        logger.warning(f"Error parsing approved_at datetime '{approved_at}': {e}")
        return False


def get_current_ta_rate_value() -> float:
    """Reads current dynamic rate per KM from PostgreSQL travel_allowance_settings or cache."""
    cached = cache.get("ta_global_rate")
    if cached is not None and isinstance(cached, dict) and "rate_per_km" in cached:
        return float(cached["rate_per_km"])

    try:
        rows = pg_execute_raw(
            "SELECT rate_per_km FROM travel_allowance_settings WHERE id = 'global' LIMIT 1",
            fetch=True
        )
        if rows and rows[0].get("rate_per_km") is not None:
            rate = float(rows[0]["rate_per_km"])
            cache.set("ta_global_rate", {"rate_per_km": rate}, ttl=300)
            return rate
    except Exception as e:
        logger.warning(f"Notice reading TA rate from postgres: {e}")

    # Fallback to active_db mock if present
    active_db = get_active_db()
    if active_db and hasattr(active_db, "collection"):
        try:
            doc = active_db.collection("app_settings").document("travel_allowance").get()
            if hasattr(doc, "exists") and doc.exists:
                d = doc.to_dict() if callable(doc.to_dict) else dict(doc)
                return float(d.get("rate_per_km", 4.0))
        except Exception as e:
            logger.debug(f"app_settings mock lookup skipped: {e}")

    return 4.0


# --- Super-Admin Prefill Permission Gate Helper ---

_IN_MEMORY_PREFILL_PERMS: Dict[str, dict] = {}

def check_prefill_permission(current_user: dict) -> bool:
    """
    Enforces Super Admin gate for the prefill action:
    - SUPER_ADMIN: Always permitted.
    - MAIN_INCHARGE: Never permitted (read-only inspection role).
    - SUB_ADMIN / ADMIN: Must have an explicit grant in travel_allowance_permissions table.
    """
    role = (current_user.get("role") or "").upper()
    if role == "SUPER_ADMIN":
        return True

    if role == "MAIN_INCHARGE":
        raise HTTPException(
            status_code=403,
            detail="Main Incharge role is read-only and cannot prefill reports."
        )

    admin_id = current_user.get("id") or current_user.get("user_id") or current_user.get("uid")
    username = current_user.get("username") or ""

    has_permission = False
    try:
        rows = pg_execute_raw(
            """SELECT p.can_prefill 
               FROM travel_allowance_permissions p
               JOIN admin_users u ON u.id = p.admin_id
               WHERE p.admin_id = %s::bigint 
                  OR LOWER(u.username) = LOWER(%s)
                  OR LOWER(u.user_id) = LOWER(%s)
               LIMIT 1""",
            [int(admin_id) if str(admin_id).isdigit() else -1, str(username), str(admin_id)],
            fetch=True
        )
        if rows and rows[0].get("can_prefill"):
            has_permission = True
    except Exception as e:
        logger.warning(f"Error querying travel_allowance_permissions: {e}")

    # In-memory mock support for unit tests
    if not has_permission and _IN_MEMORY_PREFILL_PERMS:
        p = _IN_MEMORY_PREFILL_PERMS.get(str(admin_id)) or _IN_MEMORY_PREFILL_PERMS.get(str(username).lower())
        if p and p.get("can_prefill"):
            has_permission = True

    if not has_permission:
        raise HTTPException(
            status_code=403,
            detail="Prefill access denied. Please contact Super Admin for prefill authorization."
        )

    return True


# --- Endpoints ---

@router.get("/admin/ta/rate")
def get_travel_allowance_rate(current_user: dict = Depends(get_current_user)):
    """Fetches the current global dynamic travel allowance rate per KM."""
    rate = get_current_ta_rate_value()
    return {
        "status": "success",
        "rate_per_km": rate,
        "default_rate": rate
    }


@router.post("/admin/ta/rate")
def update_travel_allowance_rate(
    req: TaRateUpdateReq,
    current_user: dict = Depends(get_current_user)
):
    """Updates the global rate per KM. Restricted to SUPER_ADMIN and MAIN_INCHARGE."""
    role = current_user.get("role", "")
    if role not in ("SUPER_ADMIN", "MAIN_INCHARGE"):
        raise HTTPException(
            status_code=403,
            detail="Access denied: Only Super Admin or Main Incharge can configure rates."
        )

    if req.rate_per_km < 0:
        raise HTTPException(status_code=400, detail="Rate per KM cannot be negative.")

    actor = current_user.get("name") or current_user.get("username") or "Admin"
    rate_val = round(float(req.rate_per_km), 2)

    try:
        pg_execute_raw(
            """INSERT INTO travel_allowance_settings (id, rate_per_km, updated_at, updated_by)
               VALUES ('global', %s, NOW(), %s)
               ON CONFLICT (id) DO UPDATE SET rate_per_km = EXCLUDED.rate_per_km, updated_at = NOW(), updated_by = EXCLUDED.updated_by""",
            [rate_val, actor]
        )
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Postgres rate update error: {e}")
        raise HTTPException(status_code=500, detail="Failed to update TA rate, please retry.")

    # Mirror to active_db mock if present
    active_db = get_active_db()
    if active_db and hasattr(active_db, "collection"):
        try:
            active_db.collection("app_settings").document("travel_allowance").set(
                {"rate_per_km": rate_val, "updated_by": actor, "updated_at": datetime.utcnow().isoformat()},
                merge=True
            )
        except Exception as e:
            logger.debug(f"mock sync skipped: {e}")

    cache.set("ta_global_rate", {"rate_per_km": rate_val}, ttl=300)
    cache.delete_prefix("ta_roster_")
    cache.delete_prefix("ta_statewide_summary_")

    return {
        "status": "success",
        "message": f"Global TA rate updated to ₹{rate_val:.2f}/KM successfully.",
        "rate_per_km": rate_val
    }


# --- Super-Admin Prefill Permission Endpoints ---

@router.get("/admin/ta/prefill-access-list")
def get_prefill_access_list(current_user: dict = Depends(get_current_user)):
    """Returns list of admin users and their prefill access status. Super Admin only."""
    if current_user.get("role") != "SUPER_ADMIN":
        raise HTTPException(status_code=403, detail="Super Admin authorization required.")

    admin_list = []
    try:
        rows = pg_execute_raw(
            """SELECT u.id, u.user_id, u.username, u.name, u.role,
                      COALESCE(p.can_prefill, false) as can_prefill,
                      p.granted_by, p.updated_at
               FROM admin_users u
               LEFT JOIN travel_allowance_permissions p ON p.admin_id = u.id
               WHERE u.status = 'ACTIVE'
               ORDER BY u.role, u.name""",
            fetch=True
        )
        if rows:
            for r in rows:
                admin_list.append({
                    "id": r.get("id"),
                    "user_id": r.get("user_id"),
                    "username": r.get("username"),
                    "name": r.get("name"),
                    "role": r.get("role"),
                    "can_prefill": bool(r.get("can_prefill")),
                    "granted_by": r.get("granted_by"),
                    "updated_at": str(r.get("updated_at") or "")
                })
    except Exception as e:
        logger.warning(f"Error fetching prefill access list from postgres: {e}")

    # Fallback to mock store if postgres returned empty
    if not admin_list:
        active_db = get_active_db()
        if active_db and hasattr(active_db, "admin_users"):
            for u in active_db.admin_users:
                admin_list.append({
                    "id": u.get("id"),
                    "user_id": u.get("user_id"),
                    "username": u.get("username"),
                    "name": u.get("name"),
                    "role": u.get("role"),
                    "can_prefill": bool(u.get("can_prefill")),
                    "granted_by": u.get("granted_by"),
                    "updated_at": u.get("updated_at")
                })

    return {
        "status": "success",
        "data": admin_list
    }


@router.post("/admin/ta/grant-prefill-access")
def grant_prefill_access(
    req: TaPrefillAccessReq,
    current_user: dict = Depends(get_current_user)
):
    """Grants prefill permission to a specific admin user. Super Admin only."""
    if current_user.get("role") != "SUPER_ADMIN":
        raise HTTPException(status_code=403, detail="Super Admin authorization required.")

    actor = current_user.get("name") or current_user.get("username") or "Super Admin"
    target_id = req.admin_id

    # Resolve target ID from username if needed
    if not target_id and req.username:
        try:
            u_rows = pg_execute_raw(
                "SELECT id FROM admin_users WHERE LOWER(username) = LOWER(%s) LIMIT 1",
                [req.username.strip()],
                fetch=True
            )
            if u_rows:
                target_id = u_rows[0]["id"]
        except Exception as e:
            logger.debug(f"username resolution error: {e}")

    if not target_id:
        raise HTTPException(status_code=400, detail="Admin ID or username is required.")

    try:
        pg_execute_raw(
            """INSERT INTO travel_allowance_permissions (admin_id, can_prefill, granted_by, updated_at)
               VALUES (%s, true, %s, NOW())
               ON CONFLICT (admin_id) DO UPDATE SET can_prefill = true, granted_by = EXCLUDED.granted_by, updated_at = NOW()""",
            [target_id, actor]
        )
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Postgres grant prefill access error: {e}")
        raise HTTPException(status_code=500, detail="Failed to grant prefill permission, please retry.")

    # Sync in-memory store for fallback / unit tests
    _IN_MEMORY_PREFILL_PERMS[str(target_id)] = {"can_prefill": True, "granted_by": actor}
    if req.username:
        _IN_MEMORY_PREFILL_PERMS[str(req.username).lower()] = {"can_prefill": True, "granted_by": actor}

    return {
        "status": "success",
        "message": f"Prefill permission granted to admin {target_id}."
    }


@router.post("/admin/ta/revoke-prefill-access")
def revoke_prefill_access(
    req: TaPrefillAccessReq,
    current_user: dict = Depends(get_current_user)
):
    """Revokes prefill permission from a specific admin user. Super Admin only."""
    if current_user.get("role") != "SUPER_ADMIN":
        raise HTTPException(status_code=403, detail="Super Admin authorization required.")

    actor = current_user.get("name") or current_user.get("username") or "Super Admin"
    target_id = req.admin_id

    if not target_id and req.username:
        try:
            u_rows = pg_execute_raw(
                "SELECT id FROM admin_users WHERE LOWER(username) = LOWER(%s) LIMIT 1",
                [req.username.strip()],
                fetch=True
            )
            if u_rows:
                target_id = u_rows[0]["id"]
        except Exception as e:
            logger.debug(f"username resolution error: {e}")

    if not target_id:
        raise HTTPException(status_code=400, detail="Admin ID or username is required.")

    try:
        pg_execute_raw(
            """INSERT INTO travel_allowance_permissions (admin_id, can_prefill, granted_by, updated_at)
               VALUES (%s, false, %s, NOW())
               ON CONFLICT (admin_id) DO UPDATE SET can_prefill = false, granted_by = EXCLUDED.granted_by, updated_at = NOW()""",
            [target_id, actor]
        )
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Postgres revoke prefill access error: {e}")
        raise HTTPException(status_code=500, detail="Failed to revoke prefill permission, please retry.")

    # Sync in-memory store for fallback / unit tests
    _IN_MEMORY_PREFILL_PERMS[str(target_id)] = {"can_prefill": False, "granted_by": actor}
    if req.username:
        _IN_MEMORY_PREFILL_PERMS[str(req.username).lower()] = {"can_prefill": False, "granted_by": actor}

    return {
        "status": "success",
        "message": f"Prefill permission revoked from admin {target_id}."
    }


# --- Statewide Executive Summary ---

@router.get("/admin/ta/statewide-summary")
def get_statewide_ta_summary(
    month: str = Query(..., description="Target month in YYYY-MM format"),
    force_refresh: Optional[bool] = Query(False),
    current_user: dict = Depends(get_current_user)
):
    """
    Returns aggregated statewide TA statistics across all active districts for Super Admin and Main Incharge.
    Gated strictly to SUPER_ADMIN and MAIN_INCHARGE roles.
    Backed by 60s TTL cache with live invalidation on roster modifications.
    """
    role = current_user.get("role", "")
    if role not in ("SUPER_ADMIN", "MAIN_INCHARGE"):
        raise HTTPException(
            status_code=403,
            detail="Access denied. Only Super Admin and Main Incharge can view statewide TA."
        )

    month_clean = month.strip()[:7]
    cache_key = f"ta_statewide_summary_{month_clean}"

    if not force_refresh:
        cached = cache.get(cache_key)
        if cached is not None:
            return cached

    current_rate = get_current_ta_rate_value()

    # Query PostgreSQL for aggregated district figures across all active staff districts
    sql = """
        SELECT 
            d.district,
            COALESCE(s.staff_count, r.roster_count, 0) AS total_officers,
            COALESCE(r.total_km, 0) AS total_km,
            COALESCE(r.total_gross, 0) AS total_gross,
            COALESCE(r.total_deductions, 0) AS total_deductions,
            COALESCE(r.total_payable, 0) AS total_payable,
            COALESCE(r.approved_count, 0) AS approved_count,
            COALESCE(r.submitted_count, 0) AS submitted_count,
            COALESCE(r.reverted_count, 0) AS reverted_count,
            COALESCE(r.draft_count, 0) AS draft_count,
            COALESCE(r.dispute_count, 0) AS dispute_count,
            r.last_updated_at
        FROM (
            SELECT DISTINCT TRIM(district) as district 
            FROM staff_directory 
            WHERE is_active = true AND deleted_at IS NULL AND district IS NOT NULL AND TRIM(district) != ''
            UNION
            SELECT DISTINCT TRIM(district) as district 
            FROM travel_allowance_rosters 
            WHERE month = %s AND district IS NOT NULL AND TRIM(district) != ''
        ) d
        LEFT JOIN (
            SELECT TRIM(district) as district, COUNT(id) as staff_count
            FROM staff_directory
            WHERE is_active = true AND deleted_at IS NULL
            GROUP BY TRIM(district)
        ) s ON LOWER(s.district) = LOWER(d.district)
        LEFT JOIN (
            SELECT 
                TRIM(district) as district,
                COUNT(id) as roster_count,
                COALESCE(SUM(total_km), 0) AS total_km,
                COALESCE(SUM(gross_amount), 0) AS total_gross,
                COALESCE(SUM(deduction_amount), 0) AS total_deductions,
                COALESCE(SUM(final_payable_amount), 0) AS total_payable,
                COUNT(CASE WHEN status = 'APPROVED' THEN 1 END) AS approved_count,
                COUNT(CASE WHEN status = 'SUBMITTED' THEN 1 END) AS submitted_count,
                COUNT(CASE WHEN status = 'REVERTED' THEN 1 END) AS reverted_count,
                COUNT(CASE WHEN status = 'DRAFT' OR status IS NULL THEN 1 END) AS draft_count,
                COUNT(CASE WHEN dispute_status = 'PENDING' THEN 1 END) AS dispute_count,
                MAX(updated_at) AS last_updated_at
            FROM travel_allowance_rosters
            WHERE month = %s
            GROUP BY TRIM(district)
        ) r ON LOWER(r.district) = LOWER(d.district)
        ORDER BY d.district ASC
    """
    rows = []
    try:
        rows = pg_execute_raw(sql, [month_clean, month_clean], fetch=True) or []
    except Exception as e:
        logger.warning(f"Error querying statewide TA summary from postgres: {e}")

    # Build district summary objects
    districts = []
    tot_km = 0.0
    tot_gross = 0.0
    tot_ded = 0.0
    tot_payable = 0.0
    tot_officers = 0
    approved_districts = 0
    pending_districts = 0
    disputed_districts = 0

    for row in rows:
        d_name = row.get("district") or "Unknown"
        d_officers = int(row.get("total_officers") or 0)
        d_km = round(float(row.get("total_km") or 0.0), 2)
        d_gross = round(float(row.get("total_gross") or 0.0), 2)
        d_ded = round(float(row.get("total_deductions") or row.get("deduction_amount") or 0.0), 2)
        d_payable = round(float(row.get("total_payable") or row.get("final_payable_amount") or 0.0), 2)
        app_count = int(row.get("approved_count") or 0)
        sub_count = int(row.get("submitted_count") or 0)
        rev_count = int(row.get("reverted_count") or 0)
        dft_count = int(row.get("draft_count") or 0)
        disp_count = int(row.get("dispute_count") or 0)
        last_updated = row.get("last_updated_at")
        if hasattr(last_updated, "isoformat"):
            last_updated = last_updated.isoformat()
        elif last_updated is not None:
            last_updated = str(last_updated)

        completion_pct = round((app_count / d_officers * 100.0), 1) if d_officers > 0 else 0.0

        if app_count == d_officers and d_officers > 0:
            approved_districts += 1
        if sub_count > 0:
            pending_districts += 1
        if disp_count > 0:
            disputed_districts += 1

        if app_count == d_officers and d_officers > 0:
            d_status = "APPROVED"
        elif sub_count > 0:
            d_status = "SUBMITTED"
        elif rev_count > 0:
            d_status = "REVERTED"
        elif disp_count > 0:
            d_status = "DISPUTED"
        elif last_updated is None and dft_count == 0:
            d_status = "NOT_STARTED"
        else:
            d_status = "DRAFT"

        tot_officers += d_officers
        tot_km += d_km
        tot_gross += d_gross
        tot_ded += d_ded
        tot_payable += d_payable

        districts.append({
            "district": d_name,
            "total_officers": d_officers,
            "total_km": d_km,
            "gross_amount": d_gross,
            "deduction_amount": d_ded,
            "final_payable_amount": d_payable,
            "approved_count": app_count,
            "submitted_count": sub_count,
            "reverted_count": rev_count,
            "draft_count": dft_count,
            "dispute_count": disp_count,
            "completion_pct": completion_pct,
            "status": d_status,
            "last_updated_at": last_updated
        })

    total_districts = len(districts)
    overall_completion_pct = round((approved_districts / total_districts * 100.0), 1) if total_districts > 0 else 0.0

    result = {
        "status": "success",
        "month": month_clean,
        "rate_per_km": current_rate,
        "summary": {
            "total_districts": total_districts,
            "total_officers": tot_officers,
            "total_km": round(tot_km, 2),
            "total_gross": round(tot_gross, 2),
            "total_deductions": round(tot_ded, 2),
            "total_payable": round(tot_payable, 2),
            "approved_districts": approved_districts,
            "pending_districts": pending_districts,
            "disputed_districts": disputed_districts,
            "overall_completion_pct": overall_completion_pct
        },
        "districts": districts
    }

    cache.set(cache_key, result, ttl=60)
    return result


# --- Roster & Log Management ---

@router.get("/admin/ta/roster")
async def get_district_ta_roster(
    month: str = Query(..., description="Target month in YYYY-MM format"),
    district: str = Query(..., description="Target district name"),
    force_refresh: Optional[bool] = Query(False),
    current_user: dict = Depends(get_current_user)
):
    """
    Retrieves monthly TA roster for all active staff in a district.
    Validates Sub-Admin district scope and aggregates district-level KPIs.
    """
    canon_dist = check_district_access(current_user, district)
    cache_key = f"ta_roster_{month}_{canon_dist.lower()}"

    if not force_refresh:
        cached = cache.get(cache_key)
        if cached is not None:
            return cached

    current_rate = get_current_ta_rate_value()

    # 1. Fetch active staff directory
    district_staff = []
    try:
        staff_rows = pg_execute_raw(
            """SELECT id, name, district, designation 
               FROM staff_directory 
               WHERE (LOWER(TRIM(district)) = LOWER(TRIM(%s)) OR LOWER(TRIM(district)) = LOWER(TRIM(%s)))
                 AND is_active = true AND deleted_at IS NULL
               ORDER BY name ASC""",
            [canon_dist, district],
            fetch=True
        )
        for s in (staff_rows or []):
            name_clean = s.get("name", "").strip()
            if name_clean:
                district_staff.append({
                    "name": name_clean,
                    "designation": s.get("designation") or "Field Officer",
                    "id": str(s.get("id")),
                    "key": str(s.get("id"))
                })
    except Exception as e:
        logger.warning(f"Notice querying staff_directory from postgres: {e}")

    # Fallback to cached directory helper or mock
    if not district_staff:
        try:
            staff_records = await get_cached_staff_directory_raw()
            for s in (staff_records or []):
                if not isinstance(s, dict):
                    continue
                if s.get("is_active") is False or s.get("status") == "inactive":
                    continue
                s_dist = canonicalize_district(s.get("district") or "")
                if s_dist == canon_dist and s.get("name"):
                    name_clean = s.get("name").strip()
                    s_id = str(s.get("id") or f"{clean_alphanumeric(canon_dist)}_{clean_alphanumeric(name_clean)}")
                    district_staff.append({
                        "name": name_clean,
                        "designation": s.get("designation") or "Field Officer",
                        "id": s_id,
                        "key": s_id
                    })
        except Exception as e:
            logger.debug(f"cached directory lookup skipped: {e}")

    if not district_staff:
        try:
            dir_dict = await get_directory()
            names = dir_dict.get(canon_dist, [])
            for n in names:
                name_clean = str(n).strip()
                s_id = f"{clean_alphanumeric(canon_dist)}_{clean_alphanumeric(name_clean)}"
                district_staff.append({
                    "name": name_clean,
                    "designation": "Field Officer",
                    "id": s_id,
                    "key": s_id
                })
        except Exception as e:
            logger.debug(f"get_directory fallback skipped: {e}")

    # 2. Query saved travel allowance rosters and daily logs from PostgreSQL
    saved_list = []
    try:
        roster_rows = pg_execute_raw(
            """SELECT r.*,
                      COALESCE(
                          (SELECT json_agg(dl.* ORDER BY dl.day)
                           FROM travel_allowance_daily_logs dl
                           WHERE dl.roster_id = r.id),
                          '[]'::json
                      ) as daily_logs
               FROM travel_allowance_rosters r
               WHERE r.month = %s 
                 AND (LOWER(TRIM(r.district)) = LOWER(TRIM(%s)) 
                      OR r.district_id = (SELECT id FROM districts WHERE LOWER(TRIM(name)) = LOWER(TRIM(%s)) LIMIT 1))""",
            [month, canon_dist, canon_dist],
            fetch=True
        )
        for r in (roster_rows or []):
            d = dict(r)
            d["doc_id"] = str(d.get("id"))
            d["_doc_id"] = str(d.get("id"))
            d["days"] = d.get("daily_logs") or []
            saved_list.append(d)
    except Exception as e:
        logger.warning(f"Notice querying travel_allowance_rosters from postgres: {e}")

    # Fallback to mock store in pytest
    active_db = get_active_db()
    if not saved_list and active_db and hasattr(active_db, "store"):
        mock_logs = active_db.store.get("travel_allowance_logs", {})
        for mid, mval in mock_logs.items():
            md = dict(mval) if isinstance(mval, dict) else (mval.to_dict() if callable(mval.to_dict) else {})
            if md.get("month") == month and canonicalize_district(md.get("district") or "") == canon_dist:
                md["_doc_id"] = str(mid)
                md["doc_id"] = str(mid)
                saved_list.append(md)

    # 3. Match and Build deduplicated roster
    roster = []
    consumed_ids = set()
    seen_staff = set()

    for s in district_staff:
        s_id_clean = clean_alphanumeric(s["id"])
        s_name_clean = clean_alphanumeric(s["name"])
        if s_id_clean in seen_staff:
            continue
        seen_staff.add(s_id_clean)

        matched_log = None
        for sl in saved_list:
            if sl["_doc_id"] in consumed_ids:
                continue

            log_id_clean = clean_alphanumeric(str(sl.get("staff_id") or sl["_doc_id"]))
            log_key = clean_alphanumeric(sl.get("staff_key") or "")
            log_name = clean_alphanumeric(sl.get("staff_name") or "")

            is_match = (
                (s_id_clean and (s_id_clean == log_id_clean or s_id_clean == log_key or s_id_clean in log_id_clean)) or
                (s_name_clean and (s_name_clean == log_name or s_name_clean == log_key or s_name_clean in log_key)) or
                (sl.get("staff_name") and is_officer_name_match(s["name"], sl.get("staff_name"), canon_dist))
            )
            if is_match:
                matched_log = sl
                consumed_ids.add(sl["_doc_id"])
                break

        if matched_log:
            entry = dict(matched_log)
            if clean_alphanumeric(entry.get("staff_name")) in (s_id_clean, s_name_clean):
                entry["staff_name"] = s["name"]
            entry["doc_id"] = str(matched_log["_doc_id"])
            entry["staff_key"] = matched_log.get("staff_key") or s["id"]
            days_list = normalize_days_to_list(entry.get("days") or entry.get("daily_logs"), month=month)
            entry["days"] = days_list
            entry["active_days"] = sum(1 for d in days_list if float(d.get("total_km") or 0.0) > 0)

            recalc = calculate_log_totals(
                days_list,
                rate_per_km=entry.get("rate_per_km") or current_rate,
                deduction_amount=entry.get("deduction_amount") or 0.0
            )
            entry["total_km"] = recalc["total_km"]
            entry["gross_amount"] = recalc["gross_amount"]
            entry["deduction_amount"] = recalc["deduction_amount"]
            entry["final_payable_amount"] = recalc["final_payable_amount"]
            roster.append(entry)
        else:
            doc_id = f"{month}_{s['id']}"
            skeleton = {
                "doc_id": doc_id,
                "month": month,
                "district": canon_dist,
                "staff_name": s["name"],
                "staff_key": s["id"],
                "designation": s["designation"],
                "rate_per_km": current_rate,
                "total_km": 0.0,
                "gross_amount": 0.0,
                "deduction_amount": 0.0,
                "deduction_reason": "",
                "final_payable_amount": 0.0,
                "admin_remarks": "",
                "status": "DRAFT",
                "is_locked": False,
                "days": normalize_days_to_list([], month=month),
                "active_days": 0,
                "created_at": None,
                "updated_at": None
            }
            roster.append(skeleton)

    # 4. Include any unconsumed historical logs
    for sl in saved_list:
        if sl["_doc_id"] not in consumed_ids:
            entry = dict(sl)
            entry["doc_id"] = str(sl["_doc_id"])
            entry["staff_key"] = sl.get("staff_key") or normalize_staff_slug(sl.get("staff_name") or "")
            days_list = normalize_days_to_list(entry.get("days") or entry.get("daily_logs"), month=month)
            entry["days"] = days_list
            entry["active_days"] = sum(1 for d in days_list if float(d.get("total_km") or 0.0) > 0)
            recalc = calculate_log_totals(
                days_list,
                rate_per_km=entry.get("rate_per_km") or current_rate,
                deduction_amount=entry.get("deduction_amount") or 0.0
            )
            entry["total_km"] = recalc["total_km"]
            entry["gross_amount"] = recalc["gross_amount"]
            entry["deduction_amount"] = recalc["deduction_amount"]
            entry["final_payable_amount"] = recalc["final_payable_amount"]
            roster.append(entry)
            consumed_ids.add(sl["_doc_id"])

    # 5. District KPI Aggregations
    total_km = sum(float(r.get("total_km") or 0.0) for r in roster)
    gross_amount = sum(float(r.get("gross_amount") or 0.0) for r in roster)
    deduction_amount = sum(float(r.get("deduction_amount") or 0.0) for r in roster)
    final_payable_amount = sum(float(r.get("final_payable_amount") or 0.0) for r in roster)

    result = {
        "status": "success",
        "month": month,
        "district": canon_dist,
        "rate_per_km": current_rate,
        "kpis": {
            "total_km": round(total_km, 2),
            "gross_amount": round(gross_amount, 2),
            "deduction_amount": round(deduction_amount, 2),
            "final_payable_amount": round(final_payable_amount, 2),
            "total_staff": len(roster),
            "approved_count": sum(1 for r in roster if r.get("status") == "APPROVED"),
            "reverted_count": sum(1 for r in roster if r.get("status") == "REVERTED"),
            "submitted_count": sum(1 for r in roster if r.get("status") == "SUBMITTED"),
            "draft_count": sum(1 for r in roster if r.get("status") in ("DRAFT", None, ""))
        },
        "roster": roster
    }

    cache.set(cache_key, result, ttl=180)
    return result


@router.post("/admin/ta/prefill")
def prefill_district_ta_from_reports(
    req: TaPrefillReq,
    current_user: dict = Depends(get_current_user)
):
    """
    Pre-populates meter readings and visited locations from daily_field_reports.
    Protected by Super Admin permission gate check.
    Preserves manual overrides and deductions if an existing log exists.
    """
    check_prefill_permission(current_user)
    canon_dist = check_district_access(current_user, req.district)
    current_rate = get_current_ta_rate_value()

    staff_name = (req.staff_name or "").strip()
    staff_key = (req.staff_key or "").strip()
    if not staff_name and not staff_key:
        raise HTTPException(status_code=400, detail="staff_name or staff_key is required for prefill")

    try:
        year_str, month_str = req.month.split("-")
        year = int(year_str)
        month_num = int(month_str)
        _, num_days = calendar.monthrange(year, month_num)
    except Exception as e:
        logger.debug(f"invalid month format: {e}")
        raise HTTPException(status_code=400, detail="Invalid month format. Expected YYYY-MM.")

    # Read existing doc to preserve overrides
    existing_days_map = {}
    existing_deduction = 0.0
    existing_deduction_reason = ""
    existing_admin_remarks = ""

    try:
        ex_rows = pg_execute_raw(
            """SELECT r.*,
                      (SELECT json_agg(dl.* ORDER BY dl.day)
                       FROM travel_allowance_daily_logs dl
                       WHERE dl.roster_id = r.id) as daily_logs
               FROM travel_allowance_rosters r
               WHERE r.month = %s 
                 AND (LOWER(TRIM(r.district)) = LOWER(TRIM(%s)) OR LOWER(TRIM(r.district)) = LOWER(TRIM(%s)))
                 AND (LOWER(TRIM(r.staff_name)) = LOWER(TRIM(%s)) OR r.staff_key = %s OR r.staff_id::text = %s)
               LIMIT 1""",
            [req.month, canon_dist, req.district, staff_name, staff_key, staff_key],
            fetch=True
        )
        if ex_rows:
            ex_data = ex_rows[0]
            existing_deduction = float(ex_data.get("deduction_amount") or 0.0)
            existing_deduction_reason = ex_data.get("deduction_reason") or ""
            existing_admin_remarks = ex_data.get("admin_remarks") or ""
            ex_norm = normalize_days_to_list(ex_data.get("daily_logs"), req.month)
            for d in ex_norm:
                d_day = d.get("day")
                if d_day:
                    existing_days_map[int(d_day)] = d
    except Exception as e:
        logger.warning(f"Notice querying existing TA record for prefill: {e}")

    # Query daily_field_reports from PostgreSQL
    start_date = f"{req.month}-01"
    end_date = f"{req.month}-{num_days:02d}"

    reports_by_day = {}
    try:
        rep_rows = pg_execute_raw(
            """SELECT r.id, r.date_of_reporting, r.morning_km, r.evening_km, r.total_km, r.remark, r.fo_name,
                      COALESCE(
                          (SELECT string_agg(v.name, ', ' ORDER BY v.position)
                           FROM report_visited_names v
                           WHERE v.report_id = r.id),
                          ''
                      ) as visited_names
               FROM daily_field_reports r
               WHERE (LOWER(TRIM(r.working_place)) = LOWER(TRIM(%s)) OR LOWER(TRIM(r.working_place)) = LOWER(TRIM(%s)))
                 AND r.date_of_reporting >= %s::date
                 AND r.date_of_reporting <= %s::date
               ORDER BY r.date_of_reporting ASC""",
            [canon_dist, req.district, start_date, end_date],
            fetch=True
        )
        for r in (rep_rows or []):
            r_fo = str(r.get("fo_name") or "").strip()
            if staff_name and not is_officer_name_match(r_fo, staff_name, canon_dist):
                continue
            r_date = str(r.get("date_of_reporting") or "").strip()[:10]
            if r_date.startswith(req.month):
                try:
                    d_num = int(r_date.split("-")[2])
                    reports_by_day[d_num] = r
                except Exception as e:
                    logger.debug(f"prefill day num parse error: {e}")
    except Exception as e:
        logger.warning(f"Notice querying daily_field_reports during prefill: {e}")

    # Fallback to mock store if report_by_day is empty
    if not reports_by_day:
        active_db = get_active_db()
        if active_db and hasattr(active_db, "collection"):
            try:
                for doc in active_db.collection("daily_field_reports").stream():
                    r = doc.to_dict() if callable(doc.to_dict) else dict(doc)
                    r_wp = canonicalize_district(r.get("working_place") or "")
                    if r_wp != canon_dist:
                        continue
                    r_fo = str(r.get("fo_name") or "").strip()
                    if staff_name and not is_officer_name_match(r_fo, staff_name, canon_dist):
                        continue
                    r_date = str(r.get("date_of_reporting") or "").strip()[:10]
                    if r_date.startswith(req.month):
                        try:
                            d_num = int(r_date.split("-")[2])
                            reports_by_day[d_num] = r
                        except Exception as e:
                            logger.debug(f"mock date parse error: {e}")
            except Exception as e:
                logger.debug(f"mock daily reports stream skipped: {e}")

    # Build 31-day array
    days = []
    for day in range(1, num_days + 1):
        date_str = f"{year:04d}-{month_num:02d}-{day:02d}"
        ex_day = existing_days_map.get(day)

        if ex_day and ex_day.get("is_manual_override"):
            days.append(ex_day)
        else:
            rep = reports_by_day.get(day)
            if rep:
                m_km = float(rep.get("morning_km") or 0.0)
                e_km = float(rep.get("evening_km") or 0.0)
                t_km = max(0.0, e_km - m_km)
                visited = rep.get("visited_names", "")
                if isinstance(visited, list):
                    visited = ", ".join(str(v) for v in visited if v)
                purpose = str(rep.get("purpose") or rep.get("remark") or "")

                days.append({
                    "day": day,
                    "date": date_str,
                    "morning_km": m_km,
                    "evening_km": e_km,
                    "total_km": t_km,
                    "visited_names": visited,
                    "purpose": purpose,
                    "is_manual_override": False,
                    "admin_remarks": ex_day.get("admin_remarks", "") if ex_day else ""
                })
            else:
                days.append({
                    "day": day,
                    "date": date_str,
                    "morning_km": 0.0,
                    "evening_km": 0.0,
                    "total_km": 0.0,
                    "visited_names": "",
                    "purpose": "",
                    "is_manual_override": False,
                    "admin_remarks": ex_day.get("admin_remarks", "") if ex_day else ""
                })

    totals = calculate_log_totals(days, rate_per_km=current_rate, deduction_amount=existing_deduction)

    return {
        "status": "success",
        "message": f"Pre-filled readings for {req.month} from official daily reports.",
        "data": {
            "month": req.month,
            "district": canon_dist,
            "staff_name": staff_name,
            "staff_key": staff_key,
            "rate_per_km": current_rate,
            "total_km": totals["total_km"],
            "gross_amount": totals["gross_amount"],
            "deduction_amount": totals["deduction_amount"],
            "deduction_reason": existing_deduction_reason,
            "final_payable_amount": totals["final_payable_amount"],
            "admin_remarks": existing_admin_remarks,
            "days": days
        }
    }


@router.post("/admin/ta/save-log")
def save_travel_allowance_log(
    req: TaLogSaveReq,
    current_user: dict = Depends(get_current_user)
):
    """
    Saves meter readings, remarks, and deductions for an officer's monthly log.
    Re-calculates totals and enforces record locking guard.
    """
    canon_dist = check_district_access(current_user, req.district)
    staff_key = req.staff_key or normalize_staff_slug(req.staff_name)
    if not staff_key and not req.staff_name:
        raise HTTPException(status_code=400, detail="Staff name or key is required.")

    # 1. Resolve staff_id and district_id
    staff_id = None
    district_id = None
    try:
        s_rows = pg_execute_raw(
            """SELECT id, district_id FROM staff_directory 
               WHERE (LOWER(TRIM(name)) = LOWER(TRIM(%s)) OR id::text = %s)
                 AND deleted_at IS NULL LIMIT 1""",
            [req.staff_name, staff_key],
            fetch=True
        )
        if s_rows:
            staff_id = s_rows[0].get("id")
            district_id = s_rows[0].get("district_id")
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Postgres staff directory lookup error: {e}")
        raise HTTPException(status_code=500, detail="Failed to verify record lock status, please retry.")

    if staff_id is None:
        raise HTTPException(
            status_code=404,
            detail="Staff not found in staff_directory for the given name/key."
        )

    if not district_id:
        try:
            d_rows = pg_execute_raw(
                "SELECT id FROM districts WHERE LOWER(TRIM(name)) = LOWER(TRIM(%s)) LIMIT 1",
                [canon_dist],
                fetch=True
            )
            if d_rows:
                district_id = d_rows[0]["id"]
        except HTTPException:
            raise
        except Exception as e:
            logger.error(f"Postgres district lookup error: {e}")
            raise HTTPException(status_code=500, detail="Failed to verify record lock status, please retry.")

    if district_id is None:
        raise HTTPException(
            status_code=404,
            detail=f"District '{canon_dist}' not found in districts table."
        )

    # 2. Check existing record for lock/permission constraints
    existing_status = "DRAFT"
    existing_id = None
    try:
        ex_rows = pg_execute_raw(
            """SELECT id, status, is_locked FROM travel_allowance_rosters
               WHERE month = %s 
                 AND (staff_id = %s OR LOWER(TRIM(staff_name)) = LOWER(TRIM(%s)) OR staff_key = %s)
               LIMIT 1""",
            [req.month, staff_id, req.staff_name, staff_key],
            fetch=True
        )
        if ex_rows:
            ex = ex_rows[0]
            existing_id = ex.get("id")
            existing_status = ex.get("status", "DRAFT")
            allowed, err = validate_edit_permission(ex, current_user.get("role", ""))
            if not allowed:
                raise HTTPException(
                    status_code=423 if "locked" in err.lower() else 403,
                    detail=err
                )
        else:
            # Check mock store if active
            active_db = get_active_db()
            mock_locked = False
            if active_db and hasattr(active_db, "collection"):
                try:
                    m_id = f"{req.month}_{canon_dist.lower()}_{clean_alphanumeric(req.staff_key or req.staff_name)}"
                    doc = active_db.collection("travel_allowance_logs").document(m_id).get()
                    if hasattr(doc, "exists") and doc.exists:
                        d = doc.to_dict() if callable(doc.to_dict) else dict(doc)
                        existing_status = d.get("status", "DRAFT")
                        allowed, err = validate_edit_permission(d, current_user.get("role", ""))
                        if not allowed:
                            raise HTTPException(
                                status_code=423 if "locked" in err.lower() else 403,
                                detail=err
                            )
                        mock_locked = True
                except HTTPException:
                    raise
                except Exception as e:
                    logger.debug(f"mock lock check skipped: {e}")

            if not mock_locked:
                allowed, err = validate_edit_permission({"status": "DRAFT", "is_locked": False}, current_user.get("role", ""))
                if not allowed:
                    raise HTTPException(status_code=403, detail=err)
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Postgres lock check error: {e}")
        raise HTTPException(
            status_code=500,
            detail="Failed to verify record lock status, please retry."
        )

    current_rate = get_current_ta_rate_value()
    totals = calculate_log_totals(
        req.days,
        rate_per_km=current_rate,
        deduction_amount=req.deduction_amount or 0.0
    )

    actor_name = current_user.get("name") or current_user.get("username") or "Admin"
    final_staff_name = req.staff_name.strip()
    is_locked = True if existing_status == "APPROVED" else False

    # 3. Upsert into travel_allowance_rosters
    roster_pk = existing_id
    try:
        res = pg_execute_raw(
            """INSERT INTO travel_allowance_rosters (
                   month, district_id, district, staff_id, staff_name, staff_key, designation,
                   rate_per_km, total_km, gross_amount, deduction_amount, deduction_reason,
                   final_payable_amount, admin_remarks, status, is_locked, updated_at, updated_by
               ) VALUES (
                   %s, %s, %s, %s, %s, %s, %s,
                   %s, %s, %s, %s, %s,
                   %s, %s, %s, %s, NOW(), %s
               )
               ON CONFLICT (staff_id, month)
               DO UPDATE SET
                   district_id = EXCLUDED.district_id,
                   district = EXCLUDED.district,
                   staff_name = EXCLUDED.staff_name,
                   staff_key = EXCLUDED.staff_key,
                   designation = EXCLUDED.designation,
                   rate_per_km = EXCLUDED.rate_per_km,
                   total_km = EXCLUDED.total_km,
                   gross_amount = EXCLUDED.gross_amount,
                   deduction_amount = EXCLUDED.deduction_amount,
                   deduction_reason = EXCLUDED.deduction_reason,
                   final_payable_amount = EXCLUDED.final_payable_amount,
                   admin_remarks = EXCLUDED.admin_remarks,
                   updated_at = NOW(),
                   updated_by = EXCLUDED.updated_by
               RETURNING id""",
            [
                req.month, district_id, canon_dist, staff_id, final_staff_name, staff_key,
                req.designation or "Field Officer", current_rate, totals["total_km"],
                totals["gross_amount"], totals["deduction_amount"], req.deduction_reason or "",
                totals["final_payable_amount"], req.admin_remarks or "", existing_status,
                is_locked, actor_name
            ],
            fetch=True
        )
        if res:
            roster_pk = res[0].get("id")
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Postgres upsert roster error: {e}")
        raise HTTPException(status_code=500, detail="Failed to save travel allowance log, please retry.")

    # 4. Upsert child daily logs into travel_allowance_daily_logs
    if roster_pk:
        try:
            pg_execute_raw("DELETE FROM travel_allowance_daily_logs WHERE roster_id = %s", [roster_pk])
            daily_logs_batch = []
            for d in req.days:
                d_day = int(d.get("day", 1))
                d_date = str(d.get("date", f"{req.month}-{d_day:02d}"))
                m_km = float(d.get("morning_km") or 0.0)
                e_km = float(d.get("evening_km") or 0.0)
                calc_diff = max(0.0, e_km - m_km)
                is_over = bool(d.get("is_manual_override", False) or d.get("is_override", False))
                manual_t = d.get("manual_total_km")

                if is_over and manual_t is not None and str(manual_t).strip() != "":
                    try:
                        t_km = max(0.0, float(manual_t))
                    except (ValueError, TypeError):
                        t_km = calc_diff
                elif d.get("total_km") is not None and float(d.get("total_km") or 0.0) > 0:
                    try:
                        t_km = float(d.get("total_km"))
                    except (ValueError, TypeError):
                        t_km = calc_diff
                else:
                    t_km = calc_diff

                v_names = str(d.get("visited_names") or "")
                purp = str(d.get("purpose") or "")
                rem = str(d.get("admin_remarks") or "")

                daily_logs_batch.append((
                    roster_pk, d_day, d_date, m_km, e_km, t_km, v_names, purp, is_over, rem
                ))

            if daily_logs_batch:
                with get_db_connection() as conn:
                    if conn:
                        import psycopg2.extras
                        with conn.cursor() as cur:
                            psycopg2.extras.execute_values(
                                cur,
                                """INSERT INTO travel_allowance_daily_logs (
                                       roster_id, day, date, morning_km, evening_km, total_km,
                                       visited_names, purpose, is_manual_override, admin_remarks, updated_at
                                   ) VALUES %s""",
                                daily_logs_batch,
                                template="(%s, %s, %s::date, %s, %s, %s, %s, %s, %s, %s, NOW())"
                            )
                        conn.commit()
        except HTTPException:
            raise
        except Exception as e:
            logger.error(f"Postgres daily logs insert error: {e}", exc_info=True)
            raise HTTPException(status_code=500, detail="Failed to save daily log entries, please retry.")

    if not roster_pk:
        raise HTTPException(
            status_code=500,
            detail="Failed to save travel allowance log: unable to persist record."
        )

    # Mirror to active_db mock if running in test environment with active collection
    active_db = get_active_db()
    if active_db and hasattr(active_db, "collection"):
        try:
            m_id = f"{req.month}_{canon_dist.lower()}_{staff_key}"
            doc = active_db.collection("travel_allowance_logs").document(m_id).get()
            if hasattr(doc, "exists") and doc.exists:
                mock_doc = {
                    "doc_id": m_id,
                    "month": req.month,
                    "district": canon_dist,
                    "staff_name": final_staff_name,
                    "staff_key": staff_key,
                    "designation": req.designation or "Field Officer",
                    "rate_per_km": current_rate,
                    "total_km": totals["total_km"],
                    "gross_amount": totals["gross_amount"],
                    "deduction_amount": totals["deduction_amount"],
                    "deduction_reason": req.deduction_reason or "",
                    "final_payable_amount": totals["final_payable_amount"],
                    "admin_remarks": req.admin_remarks or "",
                    "status": existing_status,
                    "is_locked": is_locked,
                    "days": req.days,
                    "updated_at": datetime.utcnow().isoformat(),
                    "updated_by": actor_name
                }
                active_db.collection("travel_allowance_logs").document(m_id).set(mock_doc, merge=True)
        except Exception as e:
            logger.debug(f"mock sync skipped: {e}")

    cache.delete_prefix(f"ta_roster_{req.month}")
    cache.delete_prefix(f"ta_statewide_summary_{req.month[:7]}")

    return {
        "status": "success",
        "message": "Travel allowance log saved successfully.",
        "data": {
            "month": req.month,
            "district": canon_dist,
            "staff_name": final_staff_name,
            "staff_key": staff_key,
            "total_km": totals["total_km"],
            "gross_amount": totals["gross_amount"],
            "deduction_amount": totals["deduction_amount"],
            "final_payable_amount": totals["final_payable_amount"],
            "status": existing_status
        }
    }


@router.post("/admin/ta/submit-roster")
def submit_district_ta_roster(
    req: TaSubmitRosterReq,
    current_user: dict = Depends(get_current_user)
):
    """
    Submits draft or reverted staff records to Main Incharge for verification.
    Transitions records from DRAFT/REVERTED to SUBMITTED.
    """
    canon_dist = check_district_access(current_user, req.district)
    actor_name = current_user.get("name") or current_user.get("username") or "Sub-Admin"

    count = 0
    updated_rows = []
    try:
        sql = """UPDATE travel_allowance_rosters
                 SET status = 'SUBMITTED', is_locked = true, submitted_at = NOW(), submitted_by = %s, updated_at = NOW()
                 WHERE month = %s 
                   AND (LOWER(TRIM(district)) = LOWER(TRIM(%s)) OR LOWER(TRIM(district)) = LOWER(TRIM(%s)))
                   AND status IN ('DRAFT', 'REVERTED')"""
        params = [actor_name, req.month, canon_dist, req.district]
        if req.staff_keys:
            sql += " AND staff_key = ANY(%s)"
            params.append(req.staff_keys)
        sql += " RETURNING id"
        updated_rows = pg_execute_raw(sql, params, fetch=True)
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Postgres submit-roster error: {e}")
        raise HTTPException(status_code=500, detail="Failed to submit roster, please retry.")

    count = len(updated_rows) if updated_rows else 0

    # Mirror to active_db mock
    active_db = get_active_db()
    if active_db and hasattr(active_db, "collection"):
        try:
            coll = active_db.collection("travel_allowance_logs")
            docs = []
            try:
                docs = list(coll.where("month", "==", req.month).where("district", "==", canon_dist).stream())
            except Exception as e:
                logger.debug(f"where stream skipped: {e}")
                docs = []
            if not docs:
                try:
                    docs = list(coll.stream())
                except Exception as e:
                    logger.debug(f"stream skipped: {e}")
                    docs = []
            for doc in docs:
                d = doc.to_dict() if callable(doc.to_dict) else dict(doc)
                if (not d.get("month") or d.get("month") == req.month) and (not d.get("district") or canonicalize_district(d.get("district") or "") == canon_dist):
                    s_key = d.get("staff_key") or normalize_staff_slug(d.get("staff_name") or "")
                    if req.staff_keys and s_key not in req.staff_keys:
                        continue
                    if d.get("status") in ("DRAFT", "REVERTED", None, ""):
                        d["status"] = "SUBMITTED"
                        d["is_locked"] = True
                        d["submitted_at"] = datetime.utcnow().isoformat()
                        d["submitted_by"] = actor_name
                        getattr(doc, "reference", doc).set(d, merge=True)
                        if not updated_rows:
                            count += 1
        except Exception as e:
            logger.debug(f"mock sync skipped: {e}")

    cache.delete_prefix(f"ta_roster_{req.month}")
    cache.delete_prefix(f"ta_statewide_summary_{req.month[:7]}")
    return {
        "status": "success",
        "message": "Successfully submitted staff records for approval.",
        "count": count
    }


@router.post("/admin/ta/pass-staff")
def pass_staff_record(
    req: TaPassStaffReq,
    current_user: dict = Depends(get_current_user)
):
    """
    Granular per-staff pass/approval.
    Sets status = APPROVED, is_locked = True.
    """
    role = current_user.get("role", "")
    if role not in ("MAIN_INCHARGE", "SUPER_ADMIN", "ADMIN"):
        raise HTTPException(
            status_code=403,
            detail="Access denied. Only Incharge or Super Admin can pass staff records."
        )

    canon_dist = check_district_access(current_user, req.district)
    actor_name = current_user.get("name") or current_user.get("username") or "Incharge"

    # Batch passing
    if req.staff_key.upper() == "ALL" or req.staff_keys:
        target_keys = req.staff_keys or []
        updated_rows = []
        try:
            sql = """UPDATE travel_allowance_rosters
                     SET status = 'APPROVED', is_locked = true, approved_at = NOW(), approved_by = %s, updated_at = NOW()
                     WHERE month = %s 
                       AND (LOWER(TRIM(district)) = LOWER(TRIM(%s)) OR LOWER(TRIM(district)) = LOWER(TRIM(%s)))
                       AND status = 'SUBMITTED'"""
            params = [actor_name, req.month, canon_dist, req.district]
            if target_keys:
                sql += " AND staff_key = ANY(%s)"
                params.append(target_keys)
            sql += " RETURNING staff_key, id"
            updated_rows = pg_execute_raw(sql, params, fetch=True)
        except HTTPException:
            raise
        except Exception as e:
            logger.error(f"Postgres batch pass error: {e}")
            raise HTTPException(status_code=500, detail="Failed to update TA record, please retry.")

        # Sync mock store
        active_db = get_active_db()
        passed = [r["staff_key"] for r in updated_rows] if updated_rows else []
        mock_synced = False
        if active_db and hasattr(active_db, "collection"):
            try:
                for doc in active_db.collection("travel_allowance_logs").stream():
                    d = doc.to_dict() if callable(doc.to_dict) else dict(doc)
                    if d.get("month") == req.month and canonicalize_district(d.get("district") or "") == canon_dist:
                        s_key = d.get("staff_key") or normalize_staff_slug(d.get("staff_name") or "")
                        if target_keys and s_key not in target_keys:
                            continue
                        if d.get("status") == "SUBMITTED":
                            d["status"] = "APPROVED"
                            d["is_locked"] = True
                            d["approved_at"] = datetime.utcnow().isoformat()
                            d["approved_by"] = actor_name
                            getattr(doc, "reference", doc).set(d, merge=True)
                            if s_key not in passed:
                                passed.append(s_key)
                            mock_synced = True
            except Exception as e:
                logger.debug(f"mock sync skipped: {e}")

        if not updated_rows and not mock_synced:
            raise HTTPException(
                status_code=404,
                detail="Staff TA record not found for the given month/district/staff_key."
            )

        cache.delete_prefix(f"ta_roster_{req.month}")
        cache.delete_prefix(f"ta_statewide_summary_{req.month[:7]}")
        return {
            "status": "success",
            "message": "Staff records approved successfully.",
            "passed_keys": passed
        }

    # Single staff passing
    updated_rows = []
    try:
        updated_rows = pg_execute_raw(
            """UPDATE travel_allowance_rosters
               SET status = 'APPROVED', is_locked = true, approved_at = NOW(), approved_by = %s, updated_at = NOW()
               WHERE month = %s 
                 AND (LOWER(TRIM(district)) = LOWER(TRIM(%s)) OR LOWER(TRIM(district)) = LOWER(TRIM(%s)))
                 AND (staff_key = %s OR staff_id::text = %s OR LOWER(TRIM(staff_name)) = LOWER(TRIM(%s)))
               RETURNING id""",
            [actor_name, req.month, canon_dist, req.district, req.staff_key, req.staff_key, req.staff_key],
            fetch=True
        )
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Postgres single pass error: {e}")
        raise HTTPException(status_code=500, detail="Failed to update TA record, please retry.")

    # Sync mock store
    mock_synced = False
    active_db = get_active_db()
    if active_db and hasattr(active_db, "collection"):
        try:
            m_id = f"{req.month}_{canon_dist.lower()}_{clean_alphanumeric(req.staff_key)}"
            doc = active_db.collection("travel_allowance_logs").document(m_id).get()
            if hasattr(doc, "exists") and doc.exists:
                d = doc.to_dict() if callable(doc.to_dict) else dict(doc)
                d["status"] = "APPROVED"
                d["is_locked"] = True
                d["approved_at"] = datetime.utcnow().isoformat()
                d["approved_by"] = actor_name
                active_db.collection("travel_allowance_logs").document(m_id).set(d, merge=True)
                mock_synced = True
        except Exception as e:
            logger.debug(f"mock sync skipped: {e}")

    if not updated_rows and not mock_synced:
        raise HTTPException(
            status_code=404,
            detail="Staff TA record not found for the given month/district/staff_key."
        )

    cache.delete_prefix(f"ta_roster_{req.month}")
    cache.delete_prefix(f"ta_statewide_summary_{req.month[:7]}")
    return {
        "status": "success",
        "message": f"Staff record {req.staff_key} approved.",
        "data": {
            "staff_key": req.staff_key,
            "status": "APPROVED",
            "is_locked": True,
            "approved_by": actor_name
        }
    }


@router.post("/admin/ta/revert-staff")
def revert_staff_record(
    req: TaRevertStaffReq,
    current_user: dict = Depends(get_current_user)
):
    """
    Granular per-staff revert with mandatory reason.
    Sets status = REVERTED, is_locked = False.
    """
    role = current_user.get("role", "")
    if role not in ("MAIN_INCHARGE", "SUPER_ADMIN", "ADMIN"):
        raise HTTPException(
            status_code=403,
            detail="Access denied. Only Incharge or Super Admin can revert staff records."
        )

    if not req.revert_reason or not req.revert_reason.strip():
        raise HTTPException(
            status_code=400,
            detail="Revert reason is mandatory."
        )

    canon_dist = check_district_access(current_user, req.district)
    actor_name = current_user.get("name") or current_user.get("username") or "Incharge"

    updated_rows = []
    try:
        updated_rows = pg_execute_raw(
            """UPDATE travel_allowance_rosters
               SET status = 'REVERTED', is_locked = false, reverted_at = NOW(), reverted_by = %s,
                   revert_reason = %s, updated_at = NOW()
               WHERE month = %s 
                 AND (LOWER(TRIM(district)) = LOWER(TRIM(%s)) OR LOWER(TRIM(district)) = LOWER(TRIM(%s)))
                 AND (staff_key = %s OR staff_id::text = %s OR LOWER(TRIM(staff_name)) = LOWER(TRIM(%s)))
               RETURNING id""",
            [actor_name, req.revert_reason.strip(), req.month, canon_dist, req.district, req.staff_key, req.staff_key, req.staff_key],
            fetch=True
        )
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Postgres revert staff error: {e}")
        raise HTTPException(status_code=500, detail="Failed to update TA record, please retry.")

    mock_synced = False
    active_db = get_active_db()
    if active_db and hasattr(active_db, "collection"):
        try:
            m_id = f"{req.month}_{canon_dist.lower()}_{clean_alphanumeric(req.staff_key)}"
            doc = active_db.collection("travel_allowance_logs").document(m_id).get()
            if hasattr(doc, "exists") and doc.exists:
                d = doc.to_dict() if callable(doc.to_dict) else dict(doc)
                d["status"] = "REVERTED"
                d["is_locked"] = False
                d["reverted_at"] = datetime.utcnow().isoformat()
                d["reverted_by"] = actor_name
                d["revert_reason"] = req.revert_reason.strip()
                active_db.collection("travel_allowance_logs").document(m_id).set(d, merge=True)
                mock_synced = True
        except Exception as e:
            logger.debug(f"mock sync skipped: {e}")

    if not updated_rows and not mock_synced:
        raise HTTPException(
            status_code=404,
            detail="Staff TA record not found for the given month/district/staff_key."
        )

    cache.delete_prefix(f"ta_roster_{req.month}")
    cache.delete_prefix(f"ta_statewide_summary_{req.month[:7]}")
    return {
        "status": "success",
        "message": f"Staff record {req.staff_key} reverted.",
        "data": {
            "staff_key": req.staff_key,
            "status": "REVERTED",
            "is_locked": False,
            "revert_reason": req.revert_reason.strip(),
            "reverted_by": actor_name
        }
    }


@router.post("/admin/ta/unlock-staff")
def unlock_staff_record(
    req: TaUnlockStaffReq,
    current_user: dict = Depends(get_current_user)
):
    """
    Unlocks an approved/locked staff record so Sub-Admin can edit and resubmit.
    Sets is_locked = False, status = REVERTED.
    """
    role = current_user.get("role", "")
    if role not in ("MAIN_INCHARGE", "SUPER_ADMIN", "ADMIN"):
        raise HTTPException(
            status_code=403,
            detail="Access denied. Only Incharge or Super Admin can unlock staff records."
        )

    canon_dist = check_district_access(current_user, req.district)
    actor_name = current_user.get("name") or current_user.get("username") or "Incharge"

    updated_rows = []
    try:
        updated_rows = pg_execute_raw(
            """UPDATE travel_allowance_rosters
               SET status = 'REVERTED', is_locked = false, unlocked_at = NOW(), unlocked_by = %s, updated_at = NOW()
               WHERE month = %s 
                 AND (LOWER(TRIM(district)) = LOWER(TRIM(%s)) OR LOWER(TRIM(district)) = LOWER(TRIM(%s)))
                 AND (staff_key = %s OR staff_id::text = %s OR LOWER(TRIM(staff_name)) = LOWER(TRIM(%s)))
               RETURNING id""",
            [actor_name, req.month, canon_dist, req.district, req.staff_key, req.staff_key, req.staff_key],
            fetch=True
        )
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Postgres unlock staff error: {e}")
        raise HTTPException(status_code=500, detail="Failed to update TA record, please retry.")

    mock_synced = False
    active_db = get_active_db()
    if active_db and hasattr(active_db, "collection"):
        try:
            m_id = f"{req.month}_{canon_dist.lower()}_{clean_alphanumeric(req.staff_key)}"
            doc = active_db.collection("travel_allowance_logs").document(m_id).get()
            if hasattr(doc, "exists") and doc.exists:
                d = doc.to_dict() if callable(doc.to_dict) else dict(doc)
                d["status"] = "REVERTED"
                d["is_locked"] = False
                d["unlocked_at"] = datetime.utcnow().isoformat()
                d["unlocked_by"] = actor_name
                active_db.collection("travel_allowance_logs").document(m_id).set(d, merge=True)
                mock_synced = True
        except Exception as e:
            logger.debug(f"mock sync skipped: {e}")

    if not updated_rows and not mock_synced:
        raise HTTPException(
            status_code=404,
            detail="Staff TA record not found for the given month/district/staff_key."
        )

    cache.delete_prefix(f"ta_roster_{req.month}")
    cache.delete_prefix(f"ta_statewide_summary_{req.month[:7]}")
    return {
        "status": "success",
        "message": f"Staff record {req.staff_key} unlocked.",
        "data": {
            "staff_key": req.staff_key,
            "status": "REVERTED",
            "is_locked": False,
            "unlocked_by": actor_name
        }
    }


# --- Field Officer Mobile Endpoints ---

@router.get("/fo/ta/monthly-summary")
def get_fo_monthly_summary(
    month: str = Query(..., description="Target month YYYY-MM"),
    district: str = Query(..., description="Officer district"),
    fo_name: str = Query(..., description="Officer full name"),
    current_user: dict = Depends(get_current_user)
):
    """
    Field Officer mobile monthly TA summary.
    If record is not approved, returns UNDER_REVIEW with amounts masked.
    If approved, returns full financial breakdown and dispute window status.
    """
    canon_dist = canonicalize_district(district)
    record = None

    try:
        rows = pg_execute_raw(
            """SELECT r.*,
                      COALESCE(
                          (SELECT json_agg(dl.* ORDER BY dl.day)
                           FROM travel_allowance_daily_logs dl
                           WHERE dl.roster_id = r.id),
                          '[]'::json
                      ) as days
               FROM travel_allowance_rosters r
               WHERE r.month = %s 
                 AND (LOWER(TRIM(r.staff_name)) = LOWER(TRIM(%s)) OR LOWER(TRIM(r.staff_key)) = LOWER(TRIM(%s)))
                 AND (LOWER(TRIM(r.district)) = LOWER(TRIM(%s)) OR LOWER(TRIM(r.district)) = LOWER(TRIM(%s)))
               LIMIT 1""",
            [month, fo_name.strip(), normalize_staff_slug(fo_name), canon_dist, district],
            fetch=True
        )
        if rows:
            record = rows[0]
    except Exception as e:
        logger.warning(f"Notice querying FO monthly summary: {e}")

    if not record:
        active_db = get_active_db()
        if active_db and hasattr(active_db, "collection"):
            try:
                m_id = f"{month}_{canon_dist.lower()}_{clean_alphanumeric(fo_name)}"
                doc = active_db.collection("travel_allowance_logs").document(m_id).get()
                if hasattr(doc, "exists") and doc.exists:
                    record = doc.to_dict() if callable(doc.to_dict) else dict(doc)
            except Exception as e:
                logger.debug(f"mock sync skipped: {e}")

    if not record or record.get("status") != "APPROVED":
        return {
            "status": "UNDER_REVIEW",
            "message": "Monthly TA verification in progress",
            "data": None
        }

    dispute_active = is_dispute_window_open(record.get("approved_at"))
    return {
        "status": "APPROVED",
        "dispute_window_active": dispute_active,
        "data": {
            "doc_id": str(record.get("id") or f"{month}_{normalize_staff_slug(fo_name)}"),
            "month": record.get("month", month),
            "district": record.get("district", canon_dist),
            "staff_name": record.get("staff_name", fo_name),
            "staff_key": record.get("staff_key") or normalize_staff_slug(fo_name),
            "designation": record.get("designation", "Field Officer"),
            "rate_per_km": float(record.get("rate_per_km") or 4.0),
            "total_km": float(record.get("total_km") or 0.0),
            "gross_amount": float(record.get("gross_amount") or 0.0),
            "deduction_amount": float(record.get("deduction_amount") or 0.0),
            "deduction_reason": record.get("deduction_reason", ""),
            "final_payable_amount": float(record.get("final_payable_amount") or 0.0),
            "admin_remarks": record.get("admin_remarks", ""),
            "approved_at": str(record.get("approved_at") or ""),
            "approved_by": record.get("approved_by", ""),
            "dispute_status": record.get("dispute_status", "NONE"),
            "days": normalize_days_to_list(record.get("days") or record.get("daily_logs"), month=month)
        }
    }


@router.post("/fo/ta/dispute")
def raise_fo_ta_dispute(
    req: TaDisputeReq,
    current_user: dict = Depends(get_current_user)
):
    """
    Field Officer raises a dispute within the 24-hour approval window.
    Dispatches alert notification to Sub-Admin and Admin/Incharge.
    """
    effective_reason = (req.dispute_reason or req.reason or "").strip()
    if not effective_reason:
        raise HTTPException(status_code=400, detail="Dispute reason is mandatory.")

    canon_dist = canonicalize_district(req.district)
    record = None
    try:
        rows = pg_execute_raw(
            """SELECT * FROM travel_allowance_rosters
               WHERE month = %s 
                 AND (LOWER(TRIM(staff_name)) = LOWER(TRIM(%s)) OR staff_key = %s)
                 AND (LOWER(TRIM(district)) = LOWER(TRIM(%s)) OR LOWER(TRIM(district)) = LOWER(TRIM(%s)))
               LIMIT 1""",
            [req.month, req.fo_name.strip(), normalize_staff_slug(req.fo_name), canon_dist, req.district],
            fetch=True
        )
        if rows:
            record = rows[0]
    except HTTPException:
        raise
    except Exception as e:
        logger.debug(f"Postgres dispute lookup notice: {e}")

    if not record:
        active_db = get_active_db()
        if active_db and hasattr(active_db, "collection"):
            try:
                m_id = f"{req.month}_{canon_dist.lower()}_{clean_alphanumeric(req.fo_name)}"
                doc = active_db.collection("travel_allowance_logs").document(m_id).get()
                if hasattr(doc, "exists") and doc.exists:
                    record = doc.to_dict() if callable(doc.to_dict) else dict(doc)
            except Exception as e:
                logger.debug(f"mock lookup skipped: {e}")

    if not record:
        raise HTTPException(status_code=404, detail="TA record not found.")

    if record.get("status") != "APPROVED":
        raise HTTPException(status_code=400, detail="Cannot dispute a record that is not approved.")

    if not is_dispute_window_open(record.get("approved_at")):
        raise HTTPException(
            status_code=400,
            detail="Dispute window has closed. Disputes must be filed within 24 hours of approval."
        )

    now_str = datetime.utcnow().isoformat()
    updated_rows = []
    if record.get("id"):
        try:
            updated_rows = pg_execute_raw(
                """UPDATE travel_allowance_rosters
                   SET dispute_status = 'PENDING', dispute_reason = %s, disputed_at = NOW(), updated_at = NOW()
                   WHERE id = %s
                   RETURNING id""",
                [effective_reason, record.get("id")],
                fetch=True
            )
        except HTTPException:
            raise
        except Exception as e:
            logger.error(f"Postgres dispute update error: {e}")
            raise HTTPException(status_code=500, detail="Failed to update TA record, please retry.")

    # Dispatch notification to broadcast_alerts table
    try:
        notif_msg = f"{req.fo_name} raised a dispute for {req.month} TA: {effective_reason}"
        meta_json = json.dumps({
            "month": req.month,
            "district": canon_dist,
            "staff_name": req.fo_name,
            "dispute_reason": effective_reason,
            "type": "TA_DISPUTE"
        })
        pg_execute_raw(
            """INSERT INTO broadcast_alerts (title, message, alert_type, target_roles, target_district, metadata, is_active, created_at)
               VALUES (%s, %s, 'TA_DISPUTE', ARRAY['SUB_ADMIN', 'ADMIN', 'SUPER_ADMIN', 'MAIN_INCHARGE'], %s, %s::jsonb, true, NOW())""",
            [f"TA Dispute: {req.fo_name} ({canon_dist})", notif_msg, canon_dist, meta_json]
        )
    except Exception as e:
        logger.warning(f"Notice dispatching dispute alert: {e}")

    # Sync mock store
    mock_synced = False
    active_db = get_active_db()
    if active_db and hasattr(active_db, "collection"):
        try:
            m_id = f"{req.month}_{canon_dist.lower()}_{clean_alphanumeric(req.fo_name)}"
            doc = active_db.collection("travel_allowance_logs").document(m_id).get()
            if hasattr(doc, "exists") and doc.exists:
                record["dispute_status"] = "PENDING"
                record["dispute_reason"] = effective_reason
                record["disputed_at"] = now_str
                active_db.collection("travel_allowance_logs").document(m_id).set(record, merge=True)
                active_db.collection("broadcast_notifications").add({
                    "title": f"TA Dispute: {req.fo_name} ({canon_dist})",
                    "message": effective_reason,
                    "type": "TA_DISPUTE",
                    "target_roles": ["SUB_ADMIN", "ADMIN", "SUPER_ADMIN", "MAIN_INCHARGE"],
                    "target_district": canon_dist,
                    "metadata": {
                        "month": req.month,
                        "district": canon_dist,
                        "staff_name": req.fo_name,
                        "dispute_reason": effective_reason,
                        "type": "TA_DISPUTE"
                    }
                })
                mock_synced = True
        except Exception as e:
            logger.debug(f"mock sync skipped: {e}")

    if not updated_rows and not mock_synced:
        raise HTTPException(
            status_code=500,
            detail="Failed to update TA record, please retry."
        )

    cache.delete_prefix(f"ta_roster_{req.month}")
    cache.delete_prefix(f"ta_statewide_summary_{req.month[:7]}")

    return {
        "status": "success",
        "message": "Dispute submitted successfully to Incharge and Sub-Admin."
    }


@router.post("/admin/ta/resolve-dispute")
def resolve_ta_dispute(
    req: TaResolveDisputeReq,
    current_user: dict = Depends(get_current_user)
):
    """
    Main Incharge resolves a pending FO dispute:
    - ACCEPT: Unlocks record (is_locked = False, status = REVERTED) for Sub-Admin correction.
    - REJECT: Closes dispute (dispute_status = REJECTED, status = APPROVED, is_locked = True).
    """
    role = current_user.get("role", "")
    if role not in ("MAIN_INCHARGE", "SUPER_ADMIN", "ADMIN"):
        raise HTTPException(status_code=403, detail="Access denied. Only Incharge or Super Admin can resolve disputes.")

    action = req.action.upper().strip()
    if action not in ("ACCEPT", "REJECT"):
        raise HTTPException(status_code=400, detail="Action must be either ACCEPT or REJECT.")

    canon_dist = check_district_access(current_user, req.district)
    actor_name = current_user.get("name") or current_user.get("username") or "Incharge"

    new_status = "REVERTED" if action == "ACCEPT" else "APPROVED"
    new_disp_status = "RESOLVED" if action == "ACCEPT" else "REJECTED"
    new_locked = False if action == "ACCEPT" else True

    updated_rows = []
    try:
        updated_rows = pg_execute_raw(
            """UPDATE travel_allowance_rosters
               SET status = %s, is_locked = %s, dispute_status = %s,
                   dispute_resolved_at = NOW(), dispute_resolved_by = %s,
                   dispute_resolution_remarks = %s, updated_at = NOW()
               WHERE month = %s 
                 AND (LOWER(TRIM(district)) = LOWER(TRIM(%s)) OR LOWER(TRIM(district)) = LOWER(TRIM(%s)))
                 AND (staff_key = %s OR staff_id::text = %s OR LOWER(TRIM(staff_name)) = LOWER(TRIM(%s)))
               RETURNING id""",
            [
                new_status, new_locked, new_disp_status, actor_name,
                req.resolution_remarks or "", req.month, canon_dist, req.district,
                req.staff_key, req.staff_key, req.staff_key
            ],
            fetch=True
        )
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Postgres resolve dispute error: {e}")
        raise HTTPException(status_code=500, detail="Failed to update TA record, please retry.")

    mock_synced = False
    active_db = get_active_db()
    if active_db and hasattr(active_db, "collection"):
        try:
            m_id = f"{req.month}_{canon_dist.lower()}_{clean_alphanumeric(req.staff_key)}"
            doc = active_db.collection("travel_allowance_logs").document(m_id).get()
            if hasattr(doc, "exists") and doc.exists:
                d = doc.to_dict() if callable(doc.to_dict) else dict(doc)
                d["status"] = new_status
                d["is_locked"] = new_locked
                d["dispute_status"] = new_disp_status
                d["dispute_resolved_at"] = datetime.utcnow().isoformat()
                d["dispute_resolved_by"] = actor_name
                d["dispute_resolution_remarks"] = req.resolution_remarks or ""
                active_db.collection("travel_allowance_logs").document(m_id).set(d, merge=True)
                mock_synced = True
        except Exception as e:
            logger.debug(f"mock sync skipped: {e}")

    if not updated_rows and not mock_synced:
        raise HTTPException(
            status_code=404,
            detail="Staff TA record not found for the given month/district/staff_key."
        )

    cache.delete_prefix(f"ta_roster_{req.month}")
    cache.delete_prefix(f"ta_statewide_summary_{req.month[:7]}")
    return {
        "status": "success",
        "message": f"Dispute {action.lower()}ed successfully.",
        "data": {
            "staff_key": req.staff_key,
            "status": new_status,
            "is_locked": new_locked,
            "dispute_status": new_disp_status,
            "dispute_resolved_by": actor_name,
            "resolution_remarks": req.resolution_remarks or ""
        }
    }


# --- Excel Payroll Export ---

def build_ta_excel_workbook(district_data: dict) -> bytes:
    """
    Generates multi-sheet openpyxl workbook:
    - Sheet 1: DASHBOARD with =SUM(...) formulas.
    - Sheets 2..N: Individual 31-day bike log registers.
    """
    wb = openpyxl.Workbook()
    district = district_data.get("district", "Unknown")
    month = district_data.get("month", "")
    rate_per_km = float(district_data.get("rate_per_km", 4.0))
    roster = district_data.get("roster", [])

    # 1. Master DASHBOARD Sheet
    ws_dash = wb.active
    ws_dash.title = "DASHBOARD"

    # Title Banner
    ws_dash.merge_cells("A1:I1")
    title_cell = ws_dash.cell(row=1, column=1, value="BIHAR TB ELIMINATION MISSION - TRAVEL ALLOWANCE & BIKE LOG SUMMARY")
    title_cell.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws_dash.row_dimensions[1].height = 30

    # Subtitle Info
    ws_dash.merge_cells("A2:I2")
    sub_cell = ws_dash.cell(
        row=2, column=1,
        value=f"District: {district}  |  Month: {month}  |  Rate: ₹{rate_per_km:.2f}/KM  |  Generated on: {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}"
    )
    sub_cell.font = Font(name="Calibri", size=10, italic=True, color="475569")
    sub_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws_dash.row_dimensions[2].height = 20

    # Table Headers
    headers = ["SL", "Field Officer Name", "Designation", "Total KM", "Rate (₹/KM)", "Gross Amount (₹)", "Deductions (₹)", "Net Payable (₹)", "Status"]
    header_row_idx = 4
    ws_dash.row_dimensions[header_row_idx].height = 25
    for col_idx, h_text in enumerate(headers, start=1):
        c = ws_dash.cell(row=header_row_idx, column=col_idx, value=h_text)
        c.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = EXCEL_HEADER_BORDER

    # Data Rows
    start_row = 5
    for idx, staff in enumerate(roster, start=1):
        r = start_row + idx - 1
        ws_dash.row_dimensions[r].height = 20
        total_km = float(staff.get("total_km") or 0.0)
        ded = float(staff.get("deduction_amount") or 0.0)
        status = staff.get("status") or "DRAFT"

        ws_dash.cell(row=r, column=1, value=idx).alignment = Alignment(horizontal="center", vertical="center")
        ws_dash.cell(row=r, column=2, value=staff.get("staff_name", "")).alignment = Alignment(horizontal="left", vertical="center")
        ws_dash.cell(row=r, column=3, value=staff.get("designation", "Field Officer")).alignment = Alignment(horizontal="center", vertical="center")

        c_km = ws_dash.cell(row=r, column=4, value=total_km)
        c_km.number_format = "#,##0.00"
        c_km.alignment = Alignment(horizontal="right", vertical="center")

        c_rate = ws_dash.cell(row=r, column=5, value=rate_per_km)
        c_rate.number_format = "₹#,##0.00"
        c_rate.alignment = Alignment(horizontal="right", vertical="center")

        c_gross = ws_dash.cell(row=r, column=6, value=f"=D{r}*E{r}")
        c_gross.number_format = "₹#,##0.00"
        c_gross.alignment = Alignment(horizontal="right", vertical="center")

        c_ded = ws_dash.cell(row=r, column=7, value=ded)
        c_ded.number_format = "₹#,##0.00"
        c_ded.alignment = Alignment(horizontal="right", vertical="center")

        c_net = ws_dash.cell(row=r, column=8, value=f"=F{r}-G{r}")
        c_net.number_format = "₹#,##0.00"
        c_net.alignment = Alignment(horizontal="right", vertical="center")
        c_net.font = Font(name="Calibri", size=10, bold=True)

        c_status = ws_dash.cell(row=r, column=9, value=status)
        c_status.alignment = Alignment(horizontal="center", vertical="center")

        for col_i in range(1, 10):
            ws_dash.cell(row=r, column=col_i).border = EXCEL_THIN_BORDER

    # Total Row
    tot_r = start_row + len(roster)
    ws_dash.row_dimensions[tot_r].height = 24
    ws_dash.cell(row=tot_r, column=1, value="TOTAL").alignment = Alignment(horizontal="center", vertical="center")
    ws_dash.cell(row=tot_r, column=2, value=f"{len(roster)} Officers").alignment = Alignment(horizontal="left", vertical="center")
    ws_dash.cell(row=tot_r, column=3, value="")

    if len(roster) > 0:
        c_tot_km = ws_dash.cell(row=tot_r, column=4, value=f"=SUM(D{start_row}:D{tot_r - 1})")
        c_tot_gross = ws_dash.cell(row=tot_r, column=6, value=f"=SUM(F{start_row}:F{tot_r - 1})")
        c_tot_ded = ws_dash.cell(row=tot_r, column=7, value=f"=SUM(G{start_row}:G{tot_r - 1})")
        c_tot_net = ws_dash.cell(row=tot_r, column=8, value=f"=SUM(H{start_row}:H{tot_r - 1})")
    else:
        c_tot_km = ws_dash.cell(row=tot_r, column=4, value=0.0)
        c_tot_gross = ws_dash.cell(row=tot_r, column=6, value=0.0)
        c_tot_ded = ws_dash.cell(row=tot_r, column=7, value=0.0)
        c_tot_net = ws_dash.cell(row=tot_r, column=8, value=0.0)

    ws_dash.cell(row=tot_r, column=5, value="")
    ws_dash.cell(row=tot_r, column=9, value="")

    c_tot_km.number_format = "#,##0.00"
    c_tot_km.alignment = Alignment(horizontal="right", vertical="center")
    c_tot_gross.number_format = "₹#,##0.00"
    c_tot_gross.alignment = Alignment(horizontal="right", vertical="center")
    c_tot_ded.number_format = "₹#,##0.00"
    c_tot_ded.alignment = Alignment(horizontal="right", vertical="center")
    c_tot_net.number_format = "₹#,##0.00"
    c_tot_net.alignment = Alignment(horizontal="right", vertical="center")

    for col_i in range(1, 10):
        cell_t = ws_dash.cell(row=tot_r, column=col_i)
        cell_t.font = Font(name="Calibri", size=10, bold=True)
        cell_t.border = EXCEL_TOTAL_ROW_BORDER

    # Auto column widths
    for col in ws_dash.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or '')) for cell in col if not str(cell.value or '').startswith("BIHAR"))
        ws_dash.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 35)

    # 2. Individual Staff Sheets
    used_sheet_names = {"DASHBOARD"}
    for staff in roster:
        raw_name = (staff.get("staff_name") or "Field Officer").strip()
        clean_name = re.sub(r'[\[\]:*?/\\]', '_', raw_name).strip()[:24] or "Officer"
        sheet_title = clean_name
        counter = 1
        while sheet_title in used_sheet_names:
            sheet_title = f"{clean_name[:20]}_{counter}"
            counter += 1
        used_sheet_names.add(sheet_title)

        ws_staff = wb.create_sheet(title=sheet_title)

        ws_staff.merge_cells("A1:H1")
        st_title = ws_staff.cell(row=1, column=1, value=f"BIHAR TB ELIMINATION MISSION - MONTHLY BIKE LOG REGISTER ({month})")
        st_title.font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
        st_title.fill = PatternFill(start_color="065F46", end_color="065F46", fill_type="solid")
        st_title.alignment = Alignment(horizontal="center", vertical="center")
        ws_staff.row_dimensions[1].height = 28

        ws_staff.merge_cells("A2:H2")
        st_sub = ws_staff.cell(
            row=2, column=1,
            value=f"Officer: {raw_name}  |  Designation: {staff.get('designation', 'Field Officer')}  |  District: {district}  |  Status: {staff.get('status', 'DRAFT')}"
        )
        st_sub.font = Font(name="Calibri", size=10, italic=True, color="334155")
        st_sub.alignment = Alignment(horizontal="center", vertical="center")
        ws_staff.row_dimensions[2].height = 20

        day_headers = ["Day", "Date", "Morning KM", "Evening KM", "Daily KM", "Places Visited (PHC / Clinic / Field)", "Purpose / Remarks", "Manual Override"]
        ws_staff.row_dimensions[4].height = 24
        for col_idx, h_text in enumerate(day_headers, start=1):
            c = ws_staff.cell(row=4, column=col_idx, value=h_text)
            c.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
            c.fill = PatternFill(start_color="047857", end_color="047857", fill_type="solid")
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = EXCEL_HEADER_BORDER

        days = staff.get("days") or []
        st_start_row = 5
        for d_idx, d in enumerate(days, start=1):
            dr = st_start_row + d_idx - 1
            ws_staff.row_dimensions[dr].height = 18
            d_num = d.get("day", d_idx)
            d_date = d.get("date", "")
            m_km = float(d.get("morning_km") or 0.0)
            e_km = float(d.get("evening_km") or 0.0)
            t_km = float(d.get("total_km") or 0.0)
            visited = str(d.get("visited_names") or "")
            purpose = str(d.get("purpose") or d.get("admin_remarks") or "")
            override = "YES" if d.get("is_manual_override") else "NO"

            ws_staff.cell(row=dr, column=1, value=d_num).alignment = Alignment(horizontal="center", vertical="center")
            ws_staff.cell(row=dr, column=2, value=d_date).alignment = Alignment(horizontal="center", vertical="center")
            ws_staff.cell(row=dr, column=3, value=m_km).alignment = Alignment(horizontal="right", vertical="center")
            ws_staff.cell(row=dr, column=4, value=e_km).alignment = Alignment(horizontal="right", vertical="center")

            c_dkm = ws_staff.cell(row=dr, column=5, value=t_km)
            c_dkm.number_format = "#,##0.00"
            c_dkm.alignment = Alignment(horizontal="right", vertical="center")

            ws_staff.cell(row=dr, column=6, value=visited).alignment = Alignment(horizontal="left", vertical="center")
            ws_staff.cell(row=dr, column=7, value=purpose).alignment = Alignment(horizontal="left", vertical="center")
            ws_staff.cell(row=dr, column=8, value=override).alignment = Alignment(horizontal="center", vertical="center")

            for col_i in range(1, 9):
                ws_staff.cell(row=dr, column=col_i).border = EXCEL_THIN_BORDER

        st_tot_r = st_start_row + len(days)
        ws_staff.row_dimensions[st_tot_r].height = 22
        ws_staff.cell(row=st_tot_r, column=1, value="TOTAL").alignment = Alignment(horizontal="center", vertical="center")
        ws_staff.cell(row=st_tot_r, column=2, value="")
        ws_staff.cell(row=st_tot_r, column=3, value="")
        ws_staff.cell(row=st_tot_r, column=4, value="")

        if len(days) > 0:
            c_st_tot_km = ws_staff.cell(row=st_tot_r, column=5, value=f"=SUM(E{st_start_row}:E{st_tot_r - 1})")
        else:
            c_st_tot_km = ws_staff.cell(row=st_tot_r, column=5, value=float(staff.get("total_km") or 0.0))

        c_st_tot_km.number_format = "#,##0.00"
        c_st_tot_km.alignment = Alignment(horizontal="right", vertical="center")
        ws_staff.cell(row=st_tot_r, column=6, value="")
        ws_staff.cell(row=st_tot_r, column=7, value="")
        ws_staff.cell(row=st_tot_r, column=8, value="")

        for col_i in range(1, 9):
            cell_st = ws_staff.cell(row=st_tot_r, column=col_i)
            cell_st.font = Font(name="Calibri", size=10, bold=True)
            cell_st.border = EXCEL_TOTAL_ROW_BORDER

        # Summary Block Below Table
        sb_row = st_tot_r + 2
        ws_staff.cell(row=sb_row, column=1, value="ACCOUNTING SUMMARY").font = Font(name="Calibri", size=11, bold=True, color="047857")

        summary_rows = [
            ("Total Distance (KM)", f"{float(staff.get('total_km') or 0.0):.2f} KM"),
            ("Approved Rate", f"₹{rate_per_km:.2f} / KM"),
            ("Gross Payable Amount", f"₹{float(staff.get('gross_amount') or 0.0):.2f}"),
            ("Admin Deduction", f"₹{float(staff.get('deduction_amount') or 0.0):.2f}"),
            ("Deduction Reason", str(staff.get('deduction_reason') or "None")),
            ("Net Payable Amount", f"₹{float(staff.get('final_payable_amount') or 0.0):.2f}"),
            ("Verification Status", str(staff.get('status') or "DRAFT")),
        ]

        for s_idx, (s_label, s_val) in enumerate(summary_rows, start=1):
            curr_sb = sb_row + s_idx
            ws_staff.cell(row=curr_sb, column=1, value=s_label).font = Font(name="Calibri", size=10, bold=(s_label == "Net Payable Amount"))
            ws_staff.cell(row=curr_sb, column=2, value=s_val).font = Font(name="Calibri", size=10, bold=(s_label == "Net Payable Amount"))

        for col in ws_staff.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = max(len(str(cell.value or '')) for cell in col if not str(cell.value or '').startswith("BIHAR"))
            ws_staff.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)

    output = io.BytesIO()
    wb.save(output)
    content_bytes = output.getvalue()
    output.close()
    return content_bytes


@router.get("/admin/ta/export-excel")
async def export_travel_allowance_excel(
    month: str = Query(..., description="Target month in YYYY-MM format"),
    district: str = Query(..., description="Target district name"),
    current_user: dict = Depends(get_current_user)
):
    """
    Generates and streams multi-sheet openpyxl Excel workbook for district TA payroll.
    Protected by Sub-Admin district scope check.
    """
    canon_dist = check_district_access(current_user, district)

    roster_res = await get_district_ta_roster(
        month=month,
        district=canon_dist,
        force_refresh=True,
        current_user=current_user
    )

    wb_bytes = build_ta_excel_workbook(roster_res)
    gc.collect()

    filename = f"TA_{safe_filename(canon_dist)}_{month}.xlsx"
    return ExcelStreamingResponse(
        wb_bytes,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )
