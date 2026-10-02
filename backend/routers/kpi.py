import os
import io
import re
import gc
import sys
import zipfile
import tempfile
import logging
import asyncio
import calendar
from datetime import datetime, timedelta, date as dt_date
from typing import Optional, List, Dict, Any
from fastapi import APIRouter, HTTPException, Depends, BackgroundTasks
from fastapi.responses import StreamingResponse, FileResponse

from backend.core.database import db
from backend.core.cache import cache
from backend.core.security import get_current_admin
from backend.core.helpers import (
    get_ist_now,
    canonicalize_district,
    is_officer_name_match,
    get_previous_month,
    log_admin_activity,
    DEFAULT_BIHAR_DISTRICTS
)
from backend.core.master_ledger import (
    get_cached_staff_targets_for_month,
    get_raw_monthly_reports
)
from backend.core.styles import (
    safe_filename,
    ExcelStreamingResponse,
    style_excel_worksheet,
    EXCEL_HEADER_BORDER,
    EXCEL_THIN_BORDER,
    EXCEL_TOTAL_ROW_BORDER,
    EXCEL_CLUSTER_DIVIDER
)

router = APIRouter(tags=["kpi"])
logger = logging.getLogger("kpi")
KPI_EXCEL_SEMAPHORE = asyncio.Semaphore(1)


@router.get("/download-excel")
async def download_excel(month: Optional[str] = None, admin: dict = Depends(get_current_admin)):
    try:
        target_month = month.strip() if month and month.strip() else get_ist_now().strftime("%Y-%m")
        main_mod = sys.modules.get("main")
        get_raw = getattr(main_mod, "get_raw_monthly_reports", get_raw_monthly_reports) if main_mod else get_raw_monthly_reports
        docs = await get_raw(target_month)
        consolidated_data = []
        
        list_fields_mapping = {
            "notification_ids": "Notification",
            "hiv_dm_ids": "HIV & DM",
            "dbt_ids": "DBT",
            "sample_collection_ids": "Sample Collection",
            "sample_tested_ids": "Sample Tested",
            "outcome_assigned_ids": "Outcome Assigned",
            "home_visit_ids": "Home Visit",
            "contact_tracing_ids": "Contact Tracing",
            "follow_up_ids": "Follow Up",
            "face_to_face_ids": "Face to Face",
            "presumptive_ids": "Presumptive",
            "documents_ids": "Documents",
            "fdc_provided_ids": "FDC Provided",
            "kit_consumption_ids": "Kit Consumption",
            "differentiated_tb_ids": "Differentiated TB",
            "tpt_treatment_start_ids": "TPT Treatment Start",
            "tpt_presumptive_ids": "TPT Presumptive",
            "adhar_face_authentication_ids": "Adhar Face Authentication",
            "consent_with_id_ids": "Consent with ID"
        }
        
        for doc in (docs or []):
            data = doc if isinstance(doc, dict) else (doc.to_dict() if hasattr(doc, "to_dict") else dict(doc))
            
            # Find the maximum length among all ID arrays
            max_len = 1  # At least 1 row per report
            for key in list_fields_mapping.keys():
                ids = data.get(key) or []
                if len(ids) > max_len:
                    max_len = len(ids)
                    
            for i in range(max_len):
                row = {
                    "Date": data.get("date_of_reporting", ""),
                    "Name": data.get("fo_name", ""),
                    "Designation": data.get("designation", ""),
                    "Block": data.get("working_place", ""),
                }
                
                # Fill array IDs
                for db_key, excel_col in list_fields_mapping.items():
                    ids = data.get(db_key) or []
                    row[excel_col] = ids[i] if i < len(ids) else ""
                    
                # Static data only on first row
                if i == 0:
                    row["Morning KM"] = data.get("morning_km", 0)
                    row["Evening KM"] = data.get("evening_km", 0)
                    row["Total KM"] = data.get("total_km", 0)
                    row["Doctors Visited"] = ", ".join(data.get("visited_names", []))
                    row["Morning KM Photo"] = data.get("morning_km_photo_url", "")
                    row["Evening KM Photo"] = data.get("evening_km_photo_url", "")
                    user_remark = data.get("remark", "")
                    override_str = "[Adjusted]" if data.get("is_override_used") else ""
                    row["Remarks"] = f"{override_str} {user_remark}".strip()
                else:
                    row["Morning KM"] = ""
                    row["Evening KM"] = ""
                    row["Total KM"] = ""
                    row["Doctors Visited"] = ""
                    row["Morning KM Photo"] = ""
                    row["Evening KM Photo"] = ""
                    row["Remarks"] = ""
                    
                consolidated_data.append(row)

        import pandas as pd  # lazy import
        import openpyxl
        df = pd.DataFrame(consolidated_data)
        df.loc[len(df)] = pd.Series({'Date': 'Designed by Insomniac'})
        output = io.BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            df.to_excel(writer, index=False, sheet_name='Consolidated Report')
        output.seek(0)
        
        actor_name = admin.get("name") or admin.get("username", "Admin")
        actor_id = admin.get("user_id") or admin.get("username", "admin")
        actor_role = admin.get("role", "SUB_ADMIN")
        await log_admin_activity(
            action_type="REPORT_DOWNLOADED",
            details=f"Admin {actor_name} downloaded Consolidated Excel for {target_month}",
            user_name=actor_name,
            user_id=actor_id,
            role=actor_role,
            diff={"report_type": "Consolidated Excel", "month": target_month}
        )
        
        return StreamingResponse(
            output, 
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", 
            headers={"Content-Disposition": "attachment; filename=DFY_Consolidated_Report.xlsx"}
        )
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))





# 17 Standard KPI Categories Definition (Exact Master Blueprint)
EXCEL_KPI_CATEGORIES = [
    ("NOTIFICATION", "notification_ids", 3),
    ("HIV & DM", "hiv_dm_ids", 4),
    ("DBT", "dbt_ids", 5),
    ("SAMPLE COLLECTION", "sample_collection_ids", 6),
    ("SAMPLE TESTED", "sample_tested_ids", 7),
    ("Outcome Assigned", "outcome_assigned_ids", 8),
    ("Home Visit", "home_visit_ids", 9),
    ("Contact Tracing", "contact_tracing_ids", 10),
    ("Follow Up", "follow_up_ids", 11),
    ("Face to Face", "face_to_face_ids", 12),
    ("Presumptive", "presumptive_ids", 13),
    ("Documents", "documents_ids", 14),
    ("FDC Provided", "fdc_provided_ids", 15),
    ("Kit Consumption", "kit_consumption_ids", 16),
    ("DIFF TB", "differentiated_tb_ids", 17),
    ("TPT START", "tpt_treatment_start_ids", 18),
    ("TPT PRESUMTIVE", "tpt_presumptive_ids", 19)
]

def get_kpi_tab_name(day: int) -> str:
    if day == 1:
        return "1ST"
    elif day == 2:
        return "2nd"
    elif day == 3:
        return "3rd"
    elif day in [21, 31]:
        return f"{day}st"
    elif day == 22:
        return "22nd"
    elif day == 23:
        return "23rd"
    else:
        return f"{day}th"

def generate_district_kpi_bytes(
    district: str, 
    month_prefix: Optional[str] = None, 
    raw_reports: Optional[list] = None, 
    target_records: Optional[list] = None,
    district_targets_map: Optional[dict] = None
) -> Optional[bytes]:
    import pandas as pd  # lazy — cached in sys.modules after first call
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    if not month_prefix:
        month_prefix = datetime.now().strftime("%Y-%m")
        
    safe_dist = safe_filename(district)
    template_path = f"templates/template_{safe_dist}.xlsx"
    if not os.path.exists(template_path):
        alt_dist = None
        if district in ["AURANGABAD-BI", "Aurangabad"]:
            alt_dist = "Aurangabad"
        elif district in ["Purba Champaran", "East Champaran"]:
            alt_dist = "East Champaran"
        elif district in ["BHOJPUR", "Bhojpur"]:
            alt_dist = "Bhojpur"
        if alt_dist:
            alt_path = f"templates/template_{safe_filename(alt_dist)}.xlsx"
            if os.path.exists(alt_path):
                template_path = alt_path

    if not os.path.exists(template_path):
        return None
        
    wb = None
    try:
        # Load workbook preserving all formulas
        wb = openpyxl.load_workbook(template_path, data_only=False)
        sheet_map = {name.strip().lower(): name for name in wb.sheetnames}

        # Configure auto-recalculation
        try:
            wb.calculation.calcMode = 'auto'
            wb.calculation.fullCalcOnLoad = True
        except Exception:
            pass

        # 1. Fetch Targets (Prioritizing Month-Scoped Target)
        target_map = {}

        c_dist = canonicalize_district(district)
        # Fetch Official District Target — use pre-fetched map (ZIP path) or 3-fallback Firestore reads
        dt_doc_id = f"{month_prefix}_{c_dist}".replace(" ", "").lower()
        official_target = None
        try:
            if district_targets_map is not None:
                # O(1) dict lookups — no Firestore reads (ZIP export path)
                prev_m = get_previous_month(month_prefix)
                fb_id = c_dist.replace(" ", "").lower()
                prev_id = f"{prev_m}_{fb_id}" if prev_m else None
                raw_val = (
                    district_targets_map.get(dt_doc_id) or
                    (district_targets_map.get(prev_id) if prev_id else None) or
                    district_targets_map.get(fb_id)
                )
                if isinstance(raw_val, (int, float)) and raw_val > 0:
                    official_target = int(raw_val)
            else:
                # Single-district download — 3-fallback Firestore reads (original behavior)
                dt_doc = db.collection("district_targets").document(dt_doc_id).get()
                if dt_doc and getattr(dt_doc, "exists", False) is True:
                    raw_val = dt_doc.to_dict().get("official_target") if hasattr(dt_doc, "to_dict") and callable(dt_doc.to_dict) and dt_doc.to_dict() else None
                    if isinstance(raw_val, (int, float)) and raw_val > 0:
                        official_target = int(raw_val)
                if official_target is None or official_target <= 0:
                    prev_m = get_previous_month(month_prefix)
                    if prev_m:
                        prev_id = f"{prev_m}_{c_dist}".replace(" ", "").lower()
                        prev_doc = db.collection("district_targets").document(prev_id).get()
                        if prev_doc and getattr(prev_doc, "exists", False) is True:
                            raw_val = prev_doc.to_dict().get("official_target") if hasattr(prev_doc, "to_dict") and callable(prev_doc.to_dict) and prev_doc.to_dict() else None
                            if isinstance(raw_val, (int, float)) and raw_val > 0:
                                official_target = int(raw_val)
                if official_target is None or official_target <= 0:
                    fb_id = c_dist.replace(" ", "").lower()
                    fb_doc = db.collection("district_targets").document(fb_id).get()
                    if fb_doc and getattr(fb_doc, "exists", False) is True:
                        raw_val = fb_doc.to_dict().get("official_target") if hasattr(fb_doc, "to_dict") and callable(fb_doc.to_dict) and fb_doc.to_dict() else None
                        if isinstance(raw_val, (int, float)) and raw_val > 0:
                            official_target = int(raw_val)
        except Exception as e:
            print(f"Notice: Target fetch exception for {district}: {e}")
            official_target = None

        if not isinstance(official_target, (int, float)) or official_target <= 0:
            official_target = None

        alias_dists = {c_dist.lower(), district.lower()}
        if "aurangabad" in c_dist.lower():
            alias_dists.update(["aurangabad-bi", "aurangabad"])
        elif "champaran" in c_dist.lower():
            alias_dists.update(["purba champaran", "east champaran"])
        elif "bhojpur" in c_dist.lower():
            alias_dists.update(["bhojpur"])

        if target_records is not None:
            prev_m = get_previous_month(month_prefix)
            for td in target_records:
                t_data = td if isinstance(td, dict) else (td.to_dict() if hasattr(td, "to_dict") else dict(td))
                td_dist = canonicalize_district(t_data.get("district", "")).lower()
                raw_td_dist = str(t_data.get("district", "")).strip().lower()
                if td_dist in alias_dists or raw_td_dist in alias_dists:
                    f_name = re.sub(r'\s+', ' ', str(t_data.get("fo_name") or t_data.get("name") or "")).strip().lower()
                    if f_name:
                        t_val = int(t_data.get("target", 50)) if str(t_data.get("target", "")).isdigit() else 50
                        td_month = t_data.get("month")
                        if td_month == month_prefix:
                            target_map[f_name] = t_val
                        elif td_month == prev_m and f_name not in target_map:
                            target_map[f_name] = t_val
                        elif not td_month and f_name not in target_map:
                            target_map[f_name] = t_val
        else:
            try:
                t_docs = db.collection("staff_targets").where("district", "==", district).stream()
                for td in t_docs:
                    t_data = td.to_dict() if hasattr(td, "to_dict") else dict(td)
                    f_name = re.sub(r'\s+', ' ', str(t_data.get("fo_name", ""))).strip().lower()
                    if f_name:
                        t_val = int(t_data.get("target", 50)) if str(t_data.get("target", "")).isdigit() else 50
                        if t_data.get("month") == month_prefix:
                            target_map[f_name] = t_val
                        elif f_name not in target_map:
                            target_map[f_name] = t_val
            except Exception as e:
                print(f"Target fetch notice for {district}: {e}")

        # 2. Fetch and Sort Daily Field Reports for this District and Month
        reports = []
        if raw_reports is not None:
            seen_report_ids = set()
            for r in raw_reports:
                r_data = r if isinstance(r, dict) else (r.to_dict() if hasattr(r, "to_dict") else dict(r))
                r_wp = canonicalize_district(r_data.get("working_place", "") or r_data.get("district", "")).lower()
                r_raw_wp = str(r_data.get("working_place", "") or r_data.get("district", "")).strip().lower()
                if r_wp in alias_dists or r_raw_wp in alias_dists:
                    r_date = str(r_data.get("date_of_reporting") or r_data.get("date") or "").strip()
                    if not month_prefix or r_date.startswith(month_prefix):
                        r_id = r_data.get("id") or r_data.get("doc_id") or f"{r_wp}_{r_data.get('fo_name')}_{r_date}"
                        if r_id not in seen_report_ids:
                            seen_report_ids.add(r_id)
                            reports.append(r_data)
        else:
            alias_queries = [c_dist, district]
            if "aurangabad" in c_dist.lower():
                alias_queries.extend(["AURANGABAD-BI", "Aurangabad"])
            elif "champaran" in c_dist.lower():
                alias_queries.extend(["Purba Champaran", "East Champaran"])
            elif "bhojpur" in c_dist.lower():
                alias_queries.extend(["BHOJPUR", "Bhojpur"])
            alias_queries = list(dict.fromkeys(alias_queries))

            seen_report_ids = set()
            start_date = f"{month_prefix}-01"
            end_date = f"{month_prefix}-31"
            for aq in alias_queries:
                docs = db.collection("daily_field_reports")\
                    .where("working_place", "==", aq)\
                    .where("date_of_reporting", ">=", start_date)\
                    .where("date_of_reporting", "<=", end_date)\
                    .stream()
                for doc in docs:
                    doc_id = getattr(doc, "id", None)
                    if not doc_id:
                        d_dict = doc.to_dict() if hasattr(doc, "to_dict") else dict(doc)
                        doc_id = d_dict.get("id") or d_dict.get("doc_id") or str(d_dict)
                    if doc_id not in seen_report_ids:
                        seen_report_ids.add(doc_id)
                        reports.append(doc.to_dict() if hasattr(doc, "to_dict") else dict(doc))

        reports.sort(key=lambda x: str(x.get("date_of_reporting", "")))

        # Map staff names to their index in this template
        staff_name_to_idx = {}
        ws_day1 = wb["1ST"] if "1ST" in wb.sheetnames else wb[sheet_map.get("1st")] if "1st" in sheet_map else None
        if ws_day1:
            s_idx = 0
            for r in range(2, ws_day1.max_row + 1, 40):
                val = ws_day1.cell(row=r, column=1).value
                if val and str(val).strip():
                    staff_name_to_idx[re.sub(r'\s+', ' ', str(val)).strip().lower()] = s_idx
                    s_idx += 1

        # Pre-calculate counts for each FO and each KPI
        num_staff = len(staff_name_to_idx)
        num_kpis = len(EXCEL_KPI_CATEGORIES)
        staff_counts = { s_idx: { k_idx: 0 for k_idx in range(num_kpis) } for s_idx in range(num_staff) }
        left_clusters = [kpi for kpi in EXCEL_KPI_CATEGORIES if kpi[0] != "Kit Consumption"]
        district_cluster_counts = { c_idx: 0 for c_idx in range(len(left_clusters)) }
        category_to_cluster_idx = {}
        c_counter = 0
        for k_idx, (cat_name, _, _) in enumerate(EXCEL_KPI_CATEGORIES):
            if cat_name != "Kit Consumption":
                category_to_cluster_idx[k_idx] = c_counter
                c_counter += 1

        # 3. Populate Tabs 3 to 33 ('1ST' to '31st')
        for rep in reports:
            date_str = str(rep.get("date_of_reporting", "")).strip()
            try:
                day_int = int(date_str.split('-')[2])
            except Exception:
                continue

            raw_tab_key = get_kpi_tab_name(day_int).lower()
            actual_tab_name = sheet_map.get(raw_tab_key)
            fo_norm = re.sub(r'\s+', ' ', str(rep.get("fo_name", ""))).strip().lower()

            if fo_norm in staff_name_to_idx:
                s_idx = staff_name_to_idx[fo_norm]

                # Accumulate staff and district counts
                for k_idx, (_, cat_key, _) in enumerate(EXCEL_KPI_CATEGORIES):
                    ids = rep.get(cat_key) or []
                    if isinstance(ids, list):
                        valid_ids = [str(pid).strip() for pid in ids if str(pid).strip()]
                        staff_counts[s_idx][k_idx] += len(valid_ids)
                        c_idx = category_to_cluster_idx.get(k_idx)
                        if c_idx is not None:
                            district_cluster_counts[c_idx] += len(valid_ids)

                if actual_tab_name and actual_tab_name in wb.sheetnames:
                    ws_day = wb[actual_tab_name]
                    start_row = 2 + (s_idx * 40)

                    # Populate IDs vertically within 40-row bounds
                    for kpi_name, cat_key, col_idx in EXCEL_KPI_CATEGORIES:
                        ids = rep.get(cat_key) or []
                        if isinstance(ids, list):
                            valid_ids = [str(pid).strip() for pid in ids if str(pid).strip()]
                            for i in range(min(40, len(valid_ids))):
                                ws_day.cell(row=start_row + i, column=col_idx).value = valid_ids[i]

        # 4. Populate Tab 2: 'CONSOLIDATED SHEET'
        if "CONSOLIDATED SHEET" in wb.sheetnames:
            ws_cons = wb["CONSOLIDATED SHEET"]

            # Wing 1: Left Side (District Master Rollup & Master Log) -- Columns A to AV (Cols 1 to 48)
            cluster_row_ptrs = { c_idx: 4 for c_idx in range(len(left_clusters)) }

            # Write Left Wing Row 2 Grand Totals
            for c_idx in range(len(left_clusters)):
                start_c = 1 + (c_idx * 3)
                # Pre-compute exact total count for immediate display across all viewers
                ws_cons.cell(row=2, column=start_c).value = district_cluster_counts[c_idx]

            for rep in reports:
                rep_date = str(rep.get("date_of_reporting", "")).strip()
                rep_fo = str(rep.get("fo_name", "")).strip()

                for c_idx, (_, cat_key, _) in enumerate(left_clusters):
                    ids = rep.get(cat_key) or []
                    if isinstance(ids, list):
                        start_c = 1 + (c_idx * 3)
                        for patient_id in ids:
                            pid_str = str(patient_id).strip()
                            if pid_str:
                                r = cluster_row_ptrs[c_idx]
                                c1 = ws_cons.cell(row=r, column=start_c, value=pid_str)
                                c2 = ws_cons.cell(row=r, column=start_c + 1, value=rep_date)
                                c3 = ws_cons.cell(row=r, column=start_c + 2, value=rep_fo)
                                c1.border = EXCEL_THIN_BORDER
                                c2.border = EXCEL_THIN_BORDER
                                c3.border = EXCEL_CLUSTER_DIVIDER
                                c1.alignment = Alignment(horizontal="center", vertical="center")
                                c2.alignment = Alignment(horizontal="center", vertical="center")
                                c3.alignment = Alignment(horizontal="left", vertical="center")
                                cluster_row_ptrs[c_idx] += 1

            # Wing 2: Right Side (Staff-Wise Performance & Indicator Wing) -- Column AW (Col 49) onwards
            staff_kpi_row_ptrs = {}
            for s_idx in range(num_staff):
                staff_base_col = 49 + (s_idx * 17)
                # Write Row 3 Staff Totals
                for k_idx in range(num_kpis):
                    staff_kpi_row_ptrs[(s_idx, k_idx)] = 4
                    ws_cons.cell(row=3, column=staff_base_col + k_idx).value = staff_counts[s_idx][k_idx]

            for rep in reports:
                fo_norm = re.sub(r'\s+', ' ', str(rep.get("fo_name", ""))).strip().lower()
                if fo_norm in staff_name_to_idx:
                    s_idx = staff_name_to_idx[fo_norm]
                    staff_base_col = 49 + (s_idx * 17)

                    for k_idx, (_, cat_key, _) in enumerate(EXCEL_KPI_CATEGORIES):
                        ids = rep.get(cat_key) or []
                        if isinstance(ids, list):
                            col = staff_base_col + k_idx
                            for patient_id in ids:
                                pid_str = str(patient_id).strip()
                                if pid_str:
                                    r = staff_kpi_row_ptrs[(s_idx, k_idx)]
                                    c_cell = ws_cons.cell(row=r, column=col, value=pid_str)
                                    if col == staff_base_col + 16:
                                        c_cell.border = EXCEL_CLUSTER_DIVIDER
                                    else:
                                        c_cell.border = EXCEL_THIN_BORDER
                                    c_cell.alignment = Alignment(horizontal="center", vertical="center")
                                    staff_kpi_row_ptrs[(s_idx, k_idx)] += 1

        # 5. Populate Tab 1: 'Performance sheet' with Cohort Breakdown
        if "Performance sheet" in wb.sheetnames:
            ws_perf = wb["Performance sheet"]
            ref_h_cell = ws_perf.cell(row=4, column=6)
            ref_data_cell = ws_perf.cell(row=5, column=6)

            # Pre-calculate staff cohort breakdown for key cascade indicators
            current_month_notif_ids = set()
            for rep in reports:
                for nid in (rep.get("notification_ids") or []):
                    nid_clean = str(nid).strip()
                    if nid_clean:
                        current_month_notif_ids.add(nid_clean)

            staff_cohort_counts = {
                s_idx: { 'hiv_cur': 0, 'hiv_prev': 0, 'udst_cur': 0, 'udst_prev': 0, 'con_cur': 0, 'con_prev': 0 }
                for s_idx in range(num_staff)
            }

            for rep in reports:
                fo_norm = re.sub(r'\s+', ' ', str(rep.get("fo_name", ""))).strip().lower()
                if fo_norm in staff_name_to_idx:
                    s_idx = staff_name_to_idx[fo_norm]
                    for pid in (rep.get("hiv_dm_ids") or []):
                        pid_clean = str(pid).strip()
                        if pid_clean:
                            if pid_clean in current_month_notif_ids:
                                staff_cohort_counts[s_idx]['hiv_cur'] += 1
                            else:
                                staff_cohort_counts[s_idx]['hiv_prev'] += 1
                    for pid in (rep.get("sample_tested_ids") or []):
                        pid_clean = str(pid).strip()
                        if pid_clean:
                            if pid_clean in current_month_notif_ids:
                                staff_cohort_counts[s_idx]['udst_cur'] += 1
                            else:
                                staff_cohort_counts[s_idx]['udst_prev'] += 1
                    for pid in (rep.get("contact_tracing_ids") or []):
                        pid_clean = str(pid).strip()
                        if pid_clean:
                            if pid_clean in current_month_notif_ids:
                                staff_cohort_counts[s_idx]['con_cur'] += 1
                            else:
                                staff_cohort_counts[s_idx]['con_prev'] += 1

            cohort_col_defs = [
                (22, 'HIV\n(Cur Month)', 'hiv_cur'),
                (23, 'HIV\n(Prev Backlog)', 'hiv_prev'),
                (24, 'UDST\n(Cur Month)', 'udst_cur'),
                (25, 'UDST\n(Prev Backlog)', 'udst_prev'),
                (26, 'Contact Tr\n(Cur Month)', 'con_cur'),
                (27, 'Contact Tr\n(Prev Backlog)', 'con_prev')
            ]

            # Populate headers in Row 4
            for col_idx, header_title, _ in cohort_col_defs:
                c = ws_perf.cell(row=4, column=col_idx, value=header_title)
                if ref_h_cell.font:
                    c.font = Font(name=ref_h_cell.font.name or "Calibri", size=ref_h_cell.font.size or 10, bold=True, color=getattr(ref_h_cell.font.color, 'rgb', '00FFFFFF') or '00FFFFFF')
                if ref_h_cell.fill:
                    fill_color = getattr(ref_h_cell.fill.start_color, 'rgb', '00374151') or '00374151'
                    c.fill = PatternFill(fill_type="solid", start_color=fill_color, end_color=fill_color)
                c.border = EXCEL_THIN_BORDER
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                col_letter = get_column_letter(col_idx)
                ws_perf.column_dimensions[col_letter].width = 14

            for r_idx in range(5, 5 + num_staff):
                cell_name = ws_perf.cell(row=r_idx, column=1).value
                if cell_name and str(cell_name).strip() not in ["GRAND TOTAL", ""]:
                    norm_name = re.sub(r'\s+', ' ', str(cell_name)).strip().lower()
                    if norm_name in staff_name_to_idx:
                        s_idx = staff_name_to_idx[norm_name]

                        # Col 3: Target
                        target_val = target_map.get(norm_name, 50)
                        ws_perf.cell(row=r_idx, column=3).value = target_val

                        # Col 4: NOTIFICATION
                        notif_count = staff_counts[s_idx][0]
                        ws_perf.cell(row=r_idx, column=4).value = notif_count

                        # Col 5: % Achieved (Formula)
                        cell_pct = ws_perf.cell(row=r_idx, column=5)
                        cell_pct.value = f"=IF(C{r_idx}>0, D{r_idx}/C{r_idx}, 0)"
                        cell_pct.number_format = "0.0%"

                        # Cols 6 to 21: Remaining 16 KPIs (including DIFF TB at 19, TPT START at 20, TPT PRESUMTIVE at 21)
                        for k_idx in range(1, num_kpis):
                            kpi_val = staff_counts[s_idx][k_idx]
                            ws_perf.cell(row=r_idx, column=5 + k_idx).value = kpi_val

                        # Cols 22 to 27: Cohort breakdown
                        for col_idx, _, field_key in cohort_col_defs:
                            c_val = staff_cohort_counts[s_idx].get(field_key, 0)
                            cc = ws_perf.cell(row=r_idx, column=col_idx, value=c_val)
                            if ref_data_cell.font:
                                cc.font = Font(name=ref_data_cell.font.name or "Calibri", size=ref_data_cell.font.size or 10, bold=False)
                            cc.alignment = Alignment(horizontal="center", vertical="center")
                            cc.border = EXCEL_THIN_BORDER

            # Locate GRAND TOTAL row and set formulas
            gt_row = None
            for r in range(5, ws_perf.max_row + 1):
                val = ws_perf.cell(row=r, column=1).value
                if val and "GRAND TOTAL" in str(val).upper():
                    gt_row = r
                    break

            if gt_row:
                if isinstance(official_target, (int, float)) and official_target > 0:
                    ws_perf.cell(row=gt_row, column=3).value = official_target
                    cell_gt_pct = ws_perf.cell(row=gt_row, column=5)
                    cell_gt_pct.value = f"=IF(C{gt_row}>0, D{gt_row}/C{gt_row}, 0)"
                    cell_gt_pct.number_format = "0.0%"
                ref_gt_cell = ws_perf.cell(row=gt_row, column=6)
                gt_fill_color = "001E3A8A"
                if ref_gt_cell and ref_gt_cell.fill and getattr(ref_gt_cell.fill.start_color, 'rgb', None):
                    gt_fill_color = getattr(ref_gt_cell.fill.start_color, 'rgb', '001E3A8A') or '001E3A8A'

                for col_idx in range(19, 28):
                    col_ltr = get_column_letter(col_idx)
                    gt_c = ws_perf.cell(row=gt_row, column=col_idx, value=f"=SUM({col_ltr}5:{col_ltr}{gt_row - 1})")
                    if ref_gt_cell and ref_gt_cell.font:
                        gt_c.font = Font(name=ref_gt_cell.font.name or "Calibri", size=ref_gt_cell.font.size or 10, bold=True, color=getattr(ref_gt_cell.font.color, 'rgb', '00FFFFFF') or '00FFFFFF')
                    else:
                        gt_c.font = Font(name="Calibri", size=10, bold=True, color="00FFFFFF")
                    gt_c.fill = PatternFill(fill_type="solid", start_color=gt_fill_color, end_color=gt_fill_color)
                    gt_c.border = EXCEL_THIN_BORDER
                    gt_c.alignment = Alignment(horizontal="center", vertical="center")

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        res_bytes = output.getvalue()
        return res_bytes
    finally:
        if wb is not None:
            try:
                wb.close()
            except Exception:
                pass
            del wb
        gc.collect()

async def generate_district_kpi_bytes_async(district: str, month_prefix: Optional[str] = None) -> Optional[bytes]:
    main_mod = sys.modules.get("main")
    get_raw = getattr(main_mod, "get_raw_monthly_reports", get_raw_monthly_reports) if main_mod else get_raw_monthly_reports
    get_targets = getattr(main_mod, "get_cached_staff_targets_for_month", get_cached_staff_targets_for_month) if main_mod else get_cached_staff_targets_for_month
    gen_bytes = getattr(main_mod, "generate_district_kpi_bytes", generate_district_kpi_bytes) if main_mod else generate_district_kpi_bytes

    clean_month = month_prefix or get_ist_now().strftime("%Y-%m")
    raw_monthly = await get_raw(clean_month)
    cached_targets = await get_targets(clean_month)
    return await asyncio.to_thread(
        gen_bytes,
        district,
        month_prefix=clean_month,
        raw_reports=raw_monthly,
        target_records=cached_targets
    )

# Concurrency Semaphore to protect Render memory/CPU from multi-tap or parallel heavy Excel exports
KPI_EXCEL_SEMAPHORE = asyncio.Semaphore(1)
attendance_excel_semaphore = asyncio.Semaphore(1)
ATTENDANCE_EXCEL_SEMAPHORE = attendance_excel_semaphore


@router.get("/download-kpi-workbook")
async def download_kpi_workbook(district: str, month: Optional[str] = None, admin: dict = Depends(get_current_admin)):
    try:
        # Sub-Admin RBAC check
        admin_role = admin.get("role", "SUB_ADMIN")
        admin_allowed = admin.get("allowed_districts", [])
        allowed_c = [canonicalize_district(a).lower() for a in admin_allowed]
        c_dist = canonicalize_district(district).lower()

        if admin_role == "SUB_ADMIN" and "all" not in allowed_c:
            if c_dist not in allowed_c and district.lower() not in allowed_c:
                raise HTTPException(
                    status_code=403, 
                    detail="Access denied: You do not have permission to download KPI reports for this district."
                )

        async with KPI_EXCEL_SEMAPHORE:
            main_mod = sys.modules.get("main")
            gen_async = getattr(main_mod, "generate_district_kpi_bytes_async", generate_district_kpi_bytes_async) if main_mod else generate_district_kpi_bytes_async
            excel_bytes = await gen_async(district, month)
            if not excel_bytes:
                raise HTTPException(status_code=404, detail=f"Template for {district} not found on server.")
                
            safe_dist = safe_filename(district)
            month_tag = month or datetime.now().strftime("%Y-%m")
            headers = {
                'Content-Disposition': f'attachment; filename="KPI_Report_{safe_dist}_{month_tag}.xlsx"'
            }
            gc.collect()

            actor_name = admin.get("name") or admin.get("username", "Admin")
            actor_id = admin.get("user_id") or admin.get("username", "admin")
            actor_role = admin.get("role", "SUB_ADMIN")
            await log_admin_activity(
                action_type="REPORT_DOWNLOADED",
                details=f"Admin {actor_name} downloaded KPI Workbook for {district} ({month_tag})",
                district=district if district and district != "All" else "",
                user_name=actor_name,
                user_id=actor_id,
                role=actor_role,
                diff={"report_type": "KPI Workbook", "district": district, "month": month_tag}
            )

            return StreamingResponse(
                io.BytesIO(excel_bytes), 
                headers=headers,
                media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@router.get("/admin/reports/medicine-consumption")
async def download_medicine_consumption(
    month: Optional[str] = None, 
    district: Optional[str] = None, 
    districts: Optional[str] = None,
    admin: dict = Depends(get_current_admin)
):
    try:
        admin_role = admin.get("role", "SUB_ADMIN")
        allowed = admin.get("allowed_districts", [])
        
        target_month = month or datetime.now().strftime("%Y-%m")
        
        # Sub-Admin RBAC check
        allowed_c = [canonicalize_district(a).lower() for a in allowed]
        
        # Determine requested district set
        req_dist_set = None
        if districts and districts.strip() and districts.strip() != "All":
            req_dist_set = set([canonicalize_district(d.strip()).lower() for d in districts.split(",") if d.strip()])
        elif district and district.strip() and district.strip() != "All":
            req_dist_set = {canonicalize_district(district.strip()).lower()}

        if admin_role == "SUB_ADMIN" and allowed and "All" not in allowed:
            if req_dist_set:
                unauthorized = req_dist_set - set(allowed_c)
                if unauthorized:
                    raise HTTPException(status_code=403, detail=f"Permission denied for district(s): {', '.join(unauthorized)}.")
            else:
                req_dist_set = set(allowed_c)
        
        async with KPI_EXCEL_SEMAPHORE:
            raw_reports = await get_raw_monthly_reports(target_month)
            
            # Filter reports by district
            filtered_reports = []
            for r in raw_reports:
                wp = canonicalize_district(r.get("working_place", "") or r.get("district", ""))
                wp_lower = wp.lower()
                
                # Check Sub-admin restriction
                if admin_role == "SUB_ADMIN" and allowed and "All" not in allowed:
                    if wp_lower not in allowed_c:
                        continue
                        
                # Check requested district filter
                if req_dist_set is not None:
                    if wp_lower not in req_dist_set:
                        continue
                        
                filtered_reports.append(r)
                
            # Build openpyxl workbook
            import openpyxl  # lazy import
            wb = openpyxl.Workbook()
            
            # Sheet 1: Detailed Patient Consumption
            ws1 = wb.active
            ws1.title = "Detailed Patient Consumption"
            ws1.append([
                "Date", "District", "FO Name", "Nikshay ID", "Patient Name",
                "Category", "Weight (kg)", "Phase", "Regimen", "Daily Dose", "Strips Issued"
            ])
            
            # Sheet 2: FO & District Summary
            ws2 = wb.create_sheet(title="FO & District Summary")
            ws2.append([
                "District", "FO Name", "Total Patients",
                "Adult IP", "Adult CP", "Pediatric IP", "Pediatric CP", "Total Strips Issued"
            ])
            
            summary_map = {}
            
            for r in filtered_reports:
                wp = canonicalize_district(r.get("working_place", "") or r.get("district", ""))
                fo = r.get("fo_name", "")
                date = r.get("date_of_reporting", "") or r.get("date", "")
                fdc_details = r.get("fdc_details", [])
                fdc_ids = r.get("fdc_provided_ids", [])
                
                # Fallback if fdc_details was not saved but fdc_provided_ids exist
                if not fdc_details and fdc_ids:
                    fdc_details = [{
                        "id": str(pid),
                        "patient_name": f"Patient #{pid}",
                        "patient_type": "adult",
                        "weight_kg": None,
                        "phase": "IP",
                        "regimen_name": "4 FDC (HRZE)",
                        "daily_dose_text": "4 FDC (Standard)",
                        "strips": 1
                    } for pid in fdc_ids]
                    
                for det in fdc_details:
                    pid = str(det.get("id", ""))
                    pname = det.get("patient_name", "")
                    ptype = str(det.get("patient_type", "adult")).lower()
                    weight = det.get("weight_kg") or ""
                    phase = str(det.get("phase", "IP")).upper()
                    regimen = det.get("regimen_name") or det.get("fdc_type") or "FDC 4"
                    daily_dose = det.get("daily_dose_text", "")
                    strips = int(det.get("strips") or 1)
                    
                    ws1.append([
                        date, wp, fo, pid, pname,
                        ptype.capitalize(), weight, phase, regimen, daily_dose, strips
                    ])
                    
                    sum_key = (wp, fo)
                    if sum_key not in summary_map:
                        summary_map[sum_key] = {
                            "total_patients": 0,
                            "adult_ip": 0,
                            "adult_cp": 0,
                            "pedia_ip": 0,
                            "pedia_cp": 0,
                            "total_strips": 0
                        }
                    s = summary_map[sum_key]
                    s["total_patients"] += 1
                    s["total_strips"] += strips
                    if ptype == "adult":
                        if phase == "IP":
                            s["adult_ip"] += 1
                        else:
                            s["adult_cp"] += 1
                    else:
                        if phase == "IP":
                            s["pedia_ip"] += 1
                        else:
                            s["pedia_cp"] += 1
                            
            for (wp, fo), s in sorted(summary_map.items()):
                ws2.append([
                    wp, fo, s["total_patients"],
                    s["adult_ip"], s["adult_cp"], s["pedia_ip"], s["pedia_cp"], s["total_strips"]
                ])
                
            style_excel_worksheet(ws1, header_fill_color="0D9488")
            style_excel_worksheet(ws2, header_fill_color="4F46E5")
            
            buf = io.BytesIO()
            wb.save(buf)
            wb.close()
            buf.seek(0)
            excel_bytes = buf.getvalue()
            
            if req_dist_set and len(req_dist_set) == 1:
                single_dist = list(req_dist_set)[0].title()
                safe_dist = safe_filename(single_dist)
            elif req_dist_set and len(req_dist_set) > 1:
                safe_dist = f"Selected_{len(req_dist_set)}_Districts"
            else:
                safe_dist = safe_filename(district or "All")

            headers = {
                'Content-Disposition': f'attachment; filename="Medicine_Consumption_{safe_dist}_{target_month}.xlsx"'
            }
            import gc
            gc.collect()
            
            actor_name = admin.get("name") or admin.get("username", "Admin")
            actor_id = admin.get("user_id") or admin.get("username", "admin")
            actor_role = admin.get("role", "SUB_ADMIN")
            await log_admin_activity(
                action_type="REPORT_DOWNLOADED",
                details=f"Admin {actor_name} downloaded KPI Workbook for {district or 'All'} ({target_month})",
                district=district if district and district != "All" else "",
                user_name=actor_name,
                user_id=actor_id,
                role=actor_role,
                diff={"report_type": "KPI Workbook", "district": district or "All", "month": target_month}
            )

            return StreamingResponse(
                io.BytesIO(excel_bytes),
                headers=headers,
                media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@router.get("/download-all-kpi-workbooks")
async def download_all_kpi_workbooks(background_tasks: BackgroundTasks, month: Optional[str] = None, districts: Optional[str] = None, admin: dict = Depends(get_current_admin)):
    try:
        admin_role = admin.get("role", "SUB_ADMIN")
        admin_allowed = admin.get("allowed_districts", [])
        allowed_c = [canonicalize_district(a).lower() for a in admin_allowed]

        if admin_role == "SUB_ADMIN" and "all" not in allowed_c and len(admin_allowed) <= 1:
            raise HTTPException(
                status_code=403, 
                detail="Bulk ZIP download is restricted to multi-district administrators. Please download your individual district KPI workbook."
            )

        all_bihar = DEFAULT_BIHAR_DISTRICTS
        if districts and districts.strip() and districts.strip() != "All":
            requested_set = set([canonicalize_district(d.strip()) for d in districts.split(",") if d.strip()])
            bihar_districts = [d for d in all_bihar if d in requested_set or canonicalize_district(d) in requested_set]
        else:
            bihar_districts = all_bihar

        # Enforce Sub-Admin permission filter
        if admin_role == "SUB_ADMIN" and admin_allowed and "all" not in allowed_c:
            bihar_districts = [d for d in bihar_districts if canonicalize_district(d).lower() in allowed_c or d.lower() in allowed_c]

        if not bihar_districts:
            raise HTTPException(status_code=400, detail="No valid districts selected or permitted.")

        month_tag = month or datetime.now().strftime("%Y-%m")
        archive_name = "DFY_KPI_Selected_Districts" if (districts and districts != "All") else "DFY_Master_KPI_All_Districts"

        with tempfile.NamedTemporaryFile(suffix=".zip", delete=False) as tmp_file:
            tmp_zip_path = tmp_file.name

        try:
            async with KPI_EXCEL_SEMAPHORE:
                main_mod = sys.modules.get("main")
                gen_bytes = getattr(main_mod, "generate_district_kpi_bytes", generate_district_kpi_bytes) if main_mod else generate_district_kpi_bytes
                get_raw = getattr(main_mod, "get_raw_monthly_reports", get_raw_monthly_reports) if main_mod else get_raw_monthly_reports
                get_targets = getattr(main_mod, "get_cached_staff_targets_for_month", get_cached_staff_targets_for_month) if main_mod else get_cached_staff_targets_for_month

                raw_monthly = await get_raw(month_tag)
                cached_targets = await get_targets(month_tag)

                # B1 fix: pre-fetch all district_targets in ONE stream before the loop
                # eliminates up to 99 blocking Firestore reads (3 per district × 33 districts)
                dt_cache_key = f"district_targets_map_{month_tag}"
                district_targets_map = cache.get(dt_cache_key)
                if district_targets_map is None:
                    dt_docs = await asyncio.to_thread(
                        lambda: list(db.collection("district_targets").stream())
                    )
                    district_targets_map = {}
                    for d in dt_docs:
                        dd = d.to_dict() if hasattr(d, "to_dict") else {}
                        district_targets_map[d.id] = dd.get("official_target", 0)
                    cache.set(dt_cache_key, district_targets_map, ttl=600)

                with zipfile.ZipFile(tmp_zip_path, "w", zipfile.ZIP_DEFLATED) as zip_file:
                    for dist in bihar_districts:
                        excel_bytes = await asyncio.to_thread(
                            gen_bytes,
                            dist,
                            month_prefix=month_tag,
                            raw_reports=raw_monthly,
                            target_records=cached_targets,
                            district_targets_map=district_targets_map
                        )
                        if excel_bytes:
                            zip_file.writestr(f"KPI_Report_{safe_filename(dist)}_{month_tag}.xlsx", excel_bytes)
                            del excel_bytes
                            gc.collect()
                        # 750ms queue relaxation to let Render CPU and event loop breathe
                        await asyncio.sleep(0.75)
        except Exception as e:
            if os.path.exists(tmp_zip_path):
                try:
                    os.remove(tmp_zip_path)
                except Exception:
                    pass
            raise HTTPException(status_code=500, detail=str(e))

        def cleanup_tmp():
            try:
                if os.path.exists(tmp_zip_path):
                    os.remove(tmp_zip_path)
            except Exception:
                pass

        background_tasks.add_task(cleanup_tmp)

        actor_name = admin.get("name") or admin.get("username", "Admin")
        actor_id = admin.get("user_id") or admin.get("username", "admin")
        actor_role = admin.get("role", "SUB_ADMIN")
        await log_admin_activity(
            action_type="REPORT_DOWNLOADED",
            details=f"Admin {actor_name} downloaded Bulk 33-District KPI ZIP for {month_tag}",
            user_name=actor_name,
            user_id=actor_id,
            role=actor_role,
            diff={"report_type": "Bulk KPI ZIP", "month": month_tag}
        )

        return FileResponse(
            tmp_zip_path,
            media_type="application/zip",
            filename=f"{archive_name}_{month_tag}.zip"
        )
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


