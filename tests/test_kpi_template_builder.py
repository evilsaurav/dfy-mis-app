# -*- coding: utf-8 -*-
"""
tests/test_kpi_template_builder.py
Unit tests for the dynamic KPI template skeleton builder (Backlog #88 Option B Phase A).
"""
import io
import sys
from pathlib import Path
import pytest
import openpyxl
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parent.parent
if str(BASE_DIR) not in sys.path:
    sys.path.insert(0, str(BASE_DIR))

from backend.core.kpi_template_builder import (
    build_district_skeleton_bytes,
    compute_roster_fingerprint,
)
from backend.routers.kpi import EXCEL_KPI_CATEGORIES, get_kpi_tab_name


def test_compute_roster_fingerprint_deterministic_and_order_independent():
    """Verify that fingerprint is deterministic, order-independent, and returns a SHA-256 hex string."""
    ids_asc = [12, 45, 99, 102]
    ids_desc = [102, 99, 45, 12]
    ids_scrambled = [45, 12, 102, 99]

    fp_asc = compute_roster_fingerprint(ids_asc)
    fp_desc = compute_roster_fingerprint(ids_desc)
    fp_scrambled = compute_roster_fingerprint(ids_scrambled)

    assert fp_asc == fp_desc == fp_scrambled
    assert isinstance(fp_asc, str)
    assert len(fp_asc) == 64  # SHA-256 hex digest length


def test_compute_roster_fingerprint_changes_on_roster_diff():
    """Verify fingerprint detects roster mutations (adding, removing, or changing staff IDs)."""
    base = [10, 20, 30]
    fp_base = compute_roster_fingerprint(base)

    # Added staff
    fp_added = compute_roster_fingerprint([10, 20, 30, 40])
    assert fp_base != fp_added

    # Removed staff
    fp_removed = compute_roster_fingerprint([10, 20])
    assert fp_base != fp_removed

    # Replaced staff
    fp_replaced = compute_roster_fingerprint([10, 20, 31])
    assert fp_base != fp_replaced


def test_build_district_skeleton_bytes_structure_and_sheet_names():
    """Verify generated skeleton contains valid bytes, exactly 33 sheets, and all expected sheet titles."""
    fake_staff = ["Officer Alpha", "Officer Beta", "Officer Gamma"]
    skeleton_bytes = build_district_skeleton_bytes("Vaishali", fake_staff)

    assert isinstance(skeleton_bytes, bytes)
    assert len(skeleton_bytes) > 0

    wb = openpyxl.load_workbook(io.BytesIO(skeleton_bytes), data_only=False)
    sheet_names = wb.sheetnames

    # Must contain Performance sheet, CONSOLIDATED SHEET, and 31 daily tabs (total 33)
    assert len(sheet_names) == 33
    assert sheet_names[0] == "Performance sheet"
    assert sheet_names[1] == "CONSOLIDATED SHEET"

    # Verify all 31 daily tabs exist with canonical ordinal names (1ST, 2nd..31st)
    for day in range(1, 32):
        expected_tab = get_kpi_tab_name(day)
        assert expected_tab in sheet_names, f"Missing daily tab {expected_tab}"


def test_performance_sheet_grand_total_row_math():
    """
    Verify Performance sheet rows for 3 staff:
    - Staff rows are rows 5, 6, 7
    - GRAND TOTAL lands exactly on row 8 (start_r 5 + len(staff) 3)
    - SUM formulas reference range C5:C7, D5:D7, etc.
    """
    fake_staff = ["Amit Kumar", "Priya Singh", "Ramesh Verma"]
    skeleton_bytes = build_district_skeleton_bytes("Muzaffarpur", fake_staff)
    wb = openpyxl.load_workbook(io.BytesIO(skeleton_bytes), data_only=False)
    ws = wb["Performance sheet"]

    # Title Banner
    assert "MUZAFFARPUR" in str(ws["A1"].value)

    # Staff rows (5, 6, 7)
    assert ws["A5"].value == "Amit Kumar"
    assert ws["B5"].value == "Field Officer"
    assert ws["C5"].value == 50
    assert ws["D5"].value == 0
    assert ws["E5"].value == "=IF(C5>0, D5/C5, 0)"

    assert ws["A6"].value == "Priya Singh"
    assert ws["E6"].value == "=IF(C6>0, D6/C6, 0)"

    assert ws["A7"].value == "Ramesh Verma"
    assert ws["E7"].value == "=IF(C7>0, D7/C7, 0)"

    # Grand Total row (start_r=5 + 3 staff = row 8)
    total_r = 8
    assert ws[f"A{total_r}"].value == "GRAND TOTAL"
    assert ws[f"C{total_r}"].value == "=SUM(C5:C7)"
    assert ws[f"D{total_r}"].value == "=SUM(D5:D7)"
    assert ws[f"E{total_r}"].value == "=IF(C8>0, D8/C8, 0)"

    # Last KPI column (Col U, 21st column)
    assert ws[f"U{total_r}"].value == "=SUM(U5:U7)"


def test_consolidated_sheet_wing2_column_math():
    """
    Verify CONSOLIDATED SHEET column math:
    - Wing 1: Left 48 columns (16 clusters x 3 cols)
    - Wing 2: 17 columns per staff starting at Col 49 (AW):
      Staff 1 (index 0): col 49 (AW) to 65 (BM)
      Staff 2 (index 1): col 66 (BN) to 82 (CD)
      Staff 3 (index 2): col 83 (CE) to 99 (CU)
    """
    fake_staff = ["Officer One", "Officer Two", "Officer Three"]
    skeleton_bytes = build_district_skeleton_bytes("Gaya", fake_staff)
    wb = openpyxl.load_workbook(io.BytesIO(skeleton_bytes), data_only=False)
    ws = wb["CONSOLIDATED SHEET"]

    # Wing 1 check: Col 1 is NOTIFICATION, Col 4 is HIV & DM
    assert ws.cell(row=1, column=1).value == "NOTIFICATION"
    assert ws.cell(row=1, column=4).value == "HIV & DM"

    # Wing 2 check:
    # Staff 1 (index 0): start col 49
    c_s1 = 49
    assert ws.cell(row=1, column=c_s1).value == "Officer One"
    assert ws.cell(row=2, column=c_s1).value == EXCEL_KPI_CATEGORIES[0][0]  # "NOTIFICATION"
    assert ws.cell(row=3, column=c_s1).value == 0

    # Staff 2 (index 1): start col 49 + 17 = 66
    c_s2 = 49 + 17
    assert ws.cell(row=1, column=c_s2).value == "Officer Two"
    assert ws.cell(row=2, column=c_s2).value == EXCEL_KPI_CATEGORIES[0][0]

    # Staff 3 (index 2): start col 49 + 2*17 = 83
    c_s3 = 49 + 34
    assert ws.cell(row=1, column=c_s3).value == "Officer Three"
    assert ws.cell(row=2, column=c_s3).value == EXCEL_KPI_CATEGORIES[0][0]


def test_daily_tabs_staff_blocks_row_math():
    """
    Verify daily tab row math:
    fo_start_r = 2 + (s_idx * 40)
    Staff 1 (idx 0): row 2..41
    Staff 2 (idx 1): row 42..81
    Staff 3 (idx 2): row 82..121
    """
    fake_staff = ["FO Alpha", "FO Beta", "FO Gamma"]
    skeleton_bytes = build_district_skeleton_bytes("Buxar", fake_staff)
    wb = openpyxl.load_workbook(io.BytesIO(skeleton_bytes), data_only=False)
    ws_1st = wb["1ST"]

    # Headers at row 1
    assert ws_1st.cell(row=1, column=1).value == "NAME"
    assert ws_1st.cell(row=1, column=2).value == "DESIGNATION"
    assert ws_1st.cell(row=1, column=3).value == EXCEL_KPI_CATEGORIES[0][0]

    # Staff 1: row 2
    assert ws_1st.cell(row=2, column=1).value == "FO Alpha"
    assert ws_1st.cell(row=2, column=2).value == "Field Officer"

    # Staff 2: row 42
    assert ws_1st.cell(row=42, column=1).value == "FO Beta"
    assert ws_1st.cell(row=42, column=2).value == "Field Officer"

    # Staff 3: row 82
    assert ws_1st.cell(row=82, column=1).value == "FO Gamma"
    assert ws_1st.cell(row=82, column=2).value == "Field Officer"
