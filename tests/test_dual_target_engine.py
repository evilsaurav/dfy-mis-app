import pytest
from unittest.mock import MagicMock, patch, AsyncMock
from fastapi.testclient import TestClient
from main import app, cache, create_access_token

client = TestClient(app)

def test_subadmin_district_target_isolation():
    # Sub-admin permitted for Jamui only
    token_jamui = create_access_token({
        "sub": "subadmin_jamui",
        "username": "subadmin_jamui",
        "role": "SUB_ADMIN",
        "name": "SubAdmin Jamui",
        "allowed_districts": ["Jamui"]
    })
    headers_jamui = {"Authorization": f"Bearer {token_jamui}"}

    # 1. Updating Jamui should succeed (200)
    with patch("main.db.collection") as mock_coll:
        mock_doc = MagicMock()
        mock_coll.return_value.document.return_value = mock_doc
        
        res = client.post("/update-district-target", json={
            "month": "2026-10",
            "district": "Jamui",
            "official_target": 100
        }, headers=headers_jamui)
        assert res.status_code == 200, res.text
        assert res.json()["success"] is True

    # 2. Updating Gaya (unauthorized) should return 403 Forbidden
    res_forbidden = client.post("/update-district-target", json={
        "month": "2026-10",
        "district": "Gaya",
        "official_target": 150
    }, headers=headers_jamui)
    assert res_forbidden.status_code == 403
    assert "cross-district target modification forbidden" in res_forbidden.text.lower()

def test_targets_endpoint_returns_dual_targets_and_buffer():
    token = create_access_token({"sub": "admin", "role": "SUPER_ADMIN", "name": "Super Admin"})
    headers = {"Authorization": f"Bearer {token}"}

    with patch("main.get_directory") as mock_dir, \
         patch("main.db.collection") as mock_coll:
        mock_dir.return_value = {"Jamui": ["Rajiv Kumar", "Bablu Kumar"]}
        
        # Mock staff targets
        doc1 = MagicMock()
        doc1.to_dict.return_value = {"district": "Jamui", "fo_name": "Rajiv Kumar", "target": 35, "month": "2026-10"}
        doc2 = MagicMock()
        doc2.to_dict.return_value = {"district": "Jamui", "fo_name": "Bablu Kumar", "target": 35, "month": "2026-10"}
        
        # Mock district target doc
        dt_doc = MagicMock()
        dt_doc.exists = True
        dt_doc.to_dict.return_value = {"district": "Jamui", "month": "2026-10", "official_target": 50}

        def mock_coll_side_effect(name):
            m = MagicMock()
            if name == "district_targets":
                m.document.return_value.get.return_value = dt_doc
            elif name == "staff_targets":
                m.stream.return_value = [doc1, doc2]
            return m

        mock_coll.side_effect = mock_coll_side_effect
        cache.clear()

        res = client.get("/targets?month=2026-10&district=Jamui", headers=headers)
        assert res.status_code == 200
        data = res.json()
        assert data["official_district_target"] == 50
        assert data["staff_targets_sum"] == 70
        assert data["buffer_percent"] == 40.0
        assert data["buffer_count"] == 20

def test_targets_endpoint_fallback_when_district_target_unset():
    token = create_access_token({"sub": "admin", "role": "SUPER_ADMIN", "name": "Super Admin"})
    headers = {"Authorization": f"Bearer {token}"}

    with patch("main.get_directory") as mock_dir, \
         patch("main.db.collection") as mock_coll:
        mock_dir.return_value = {"Jamui": ["Rajiv Kumar", "Bablu Kumar"]}
        
        doc1 = MagicMock()
        doc1.to_dict.return_value = {"district": "Jamui", "fo_name": "Rajiv Kumar", "target": 35, "month": "2026-10"}
        doc2 = MagicMock()
        doc2.to_dict.return_value = {"district": "Jamui", "fo_name": "Bablu Kumar", "target": 35, "month": "2026-10"}
        
        dt_doc = MagicMock()
        dt_doc.exists = False
        dt_doc.to_dict.return_value = None

        def mock_coll_side_effect(name):
            m = MagicMock()
            if name == "district_targets":
                m.document.return_value.get.return_value = dt_doc
            elif name == "staff_targets":
                m.stream.return_value = [doc1, doc2]
            return m

        mock_coll.side_effect = mock_coll_side_effect
        cache.clear()

        res = client.get("/targets?month=2026-10&district=Jamui", headers=headers)
        assert res.status_code == 200
        data = res.json()
        # Fallback to staff targets sum (70)
        assert data["official_district_target"] == 70
        assert data["staff_targets_sum"] == 70
        assert data["buffer_percent"] == 0.0
        assert data["buffer_count"] == 0

def test_subadmin_empty_allowed_districts_forbidden():
    token_empty = create_access_token({
        "sub": "subadmin_empty",
        "username": "subadmin_empty",
        "role": "SUB_ADMIN",
        "name": "SubAdmin Empty",
        "allowed_districts": []
    })
    headers = {"Authorization": f"Bearer {token_empty}"}

    res = client.post("/update-district-target", json={
        "month": "2026-10",
        "district": "Jamui",
        "official_target": 100
    }, headers=headers)
    assert res.status_code == 403

def test_negative_target_validation():
    token = create_access_token({"sub": "admin", "role": "SUPER_ADMIN", "name": "Super Admin"})
    headers = {"Authorization": f"Bearer {token}"}

    res = client.post("/update-district-target", json={
        "month": "2026-10",
        "district": "Jamui",
        "official_target": -10
    }, headers=headers)
    assert res.status_code == 400
    assert "cannot be negative" in res.text.lower()

def test_kpi_excel_engine_uses_official_district_target():
    import io
    import openpyxl
    from main import generate_district_kpi_bytes
    
    raw_reports = [{"working_place": "Jamui", "date_of_reporting": "2026-10-05", "fo_name": "Rajiv Kumar", "total_notifications": 10}]
    target_records = [{"district": "Jamui", "fo_name": "Rajiv Kumar", "target": 35}]
    
    with patch("main.db.collection") as mock_coll:
        dt_doc = MagicMock()
        dt_doc.exists = True
        dt_doc.to_dict.return_value = {"district": "Jamui", "month": "2026-10", "official_target": 100}
        mock_coll.return_value.document.return_value.get.return_value = dt_doc
        
        excel_bytes = generate_district_kpi_bytes("Jamui", "2026-10", raw_reports=raw_reports, target_records=target_records)
        assert excel_bytes is not None
        assert len(excel_bytes) > 0
        
        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes), data_only=False)
        assert "Performance sheet" in wb.sheetnames
        ws_perf = wb["Performance sheet"]
        
        # Verify Grand Total target cell has the official district target (100)
        gt_found = False
        for r in range(5, ws_perf.max_row + 1):
            val = ws_perf.cell(row=r, column=1).value
            if val and "GRAND TOTAL" in str(val).upper():
                assert ws_perf.cell(row=r, column=3).value == 100
                gt_found = True
                break
        assert gt_found, "GRAND TOTAL row should be present in Performance sheet"

def test_bulk_district_target_update():
    token = create_access_token({"sub": "admin", "role": "SUPER_ADMIN", "name": "Super Admin"})
    headers = {"Authorization": f"Bearer {token}"}

    with patch("main.db.collection") as mock_coll:
        mock_doc = MagicMock()
        mock_coll.return_value.document.return_value = mock_doc

        res = client.post("/update-district-targets-bulk", json={
            "month": "2026-10",
            "targets": [
                {"district": "Jehanabad", "official_target": 53},
                {"district": "Jamui", "official_target": 70}
            ]
        }, headers=headers)
        assert res.status_code == 200, res.text
        data = res.json()
        assert data["success"] is True
        assert data["count"] == 2
        assert "Jehanabad" in data["updated_districts"]
        assert "Jamui" in data["updated_districts"]

def test_get_previous_month():
    from main import get_previous_month
    assert get_previous_month("2026-10") == "2026-09"
    assert get_previous_month("2026-01") == "2025-12"
    assert get_previous_month("2026-03") == "2026-02"
    assert get_previous_month("") == ""
    assert get_previous_month(None) == ""

def test_targets_endpoint_inherits_from_previous_month():
    token = create_access_token({"sub": "admin", "role": "SUPER_ADMIN", "name": "Super Admin"})
    headers = {"Authorization": f"Bearer {token}"}

    with patch("main.get_directory") as mock_dir, \
         patch("main.db.collection") as mock_coll:
        mock_dir.return_value = {"Jamui": ["Rajiv Kumar", "Bablu Kumar"]}
        
        # Staff targets exist only for 2026-09 (previous month)
        doc1 = MagicMock()
        doc1.to_dict.return_value = {"district": "Jamui", "fo_name": "Rajiv Kumar", "target": 45, "month": "2026-09"}
        doc2 = MagicMock()
        doc2.to_dict.return_value = {"district": "Jamui", "fo_name": "Bablu Kumar", "target": 45, "month": "2026-09"}
        
        # Official district target exists only for 2026-09
        dt_doc_sep = MagicMock()
        dt_doc_sep.exists = True
        dt_doc_sep.to_dict.return_value = {"district": "Jamui", "month": "2026-09", "official_target": 80}

        dt_doc_oct = MagicMock()
        dt_doc_oct.exists = False
        dt_doc_oct.to_dict.return_value = None

        def mock_coll_side_effect(name):
            m = MagicMock()
            if name == "district_targets":
                def doc_get(doc_id):
                    dm = MagicMock()
                    if "2026-09" in doc_id:
                        dm.get.return_value = dt_doc_sep
                    else:
                        dm.get.return_value = dt_doc_oct
                    return dm
                m.document.side_effect = doc_get
            elif name == "staff_targets":
                m.stream.return_value = [doc1, doc2]
            return m

        mock_coll.side_effect = mock_coll_side_effect
        cache.clear()

        res = client.get("/targets?month=2026-10&district=Jamui", headers=headers)
        assert res.status_code == 200
        data = res.json()
        assert data["official_district_target"] == 80  # Inherited from 2026-09!
        assert data["staff_targets_sum"] == 90        # 45 + 45 inherited from 2026-09!
        assert data["buffer_count"] == 10
        assert data["targets"][0]["target"] == 45
        assert data["targets"][0].get("inherited_from") == "2026-09"

def test_bulk_staff_target_update():
    token = create_access_token({"sub": "admin", "role": "SUPER_ADMIN", "name": "Super Admin"})
    headers = {"Authorization": f"Bearer {token}"}

    with patch("main.db.collection") as mock_coll:
        mock_doc = MagicMock()
        mock_coll.return_value.document.return_value = mock_doc

        res = client.post("/update-targets-bulk", json={
            "month": "2026-10",
            "targets": [
                {"district": "Jehanabad", "fo_name": "Rajiv Kumar", "target": 20},
                {"district": "Jehanabad", "fo_name": "Bablu Kumar", "target": 25}
            ]
        }, headers=headers)
        assert res.status_code == 200, res.text
        data = res.json()
        assert data["success"] is True
        assert data["count"] == 2

def test_statewide_top_performers_dual_target_mode():
    token = create_access_token({"sub": "admin", "role": "SUPER_ADMIN", "name": "Super Admin"})
    headers = {"Authorization": f"Bearer {token}"}

    mock_reports = [
        {
            "doc_id": "rep_1",
            "working_place": "Jamui",
            "fo_name": "Rajiv Kumar",
            "date_of_reporting": "2026-10-15",
            "notification_ids": [f"N_{i}" for i in range(25)],
            "sample_tested_ids": [],
            "sample_collection_ids": []
        }
    ]

    mock_targets_payload = {
        "success": True,
        "official_targets_by_district": {
            "Jamui": 50
        },
        "targets": [
            {"district": "Jamui", "fo_name": "Rajiv Kumar", "target": 35},
            {"district": "Jamui", "fo_name": "Bablu Kumar", "target": 35}
        ]
    }

    cache.clear()

    with patch("main.get_raw_monthly_reports", new=AsyncMock(return_value=mock_reports)), \
         patch("main.get_targets", new=AsyncMock(return_value=mock_targets_payload)):

        # 1. Official mode: Jamui target should be 50, percentage = 25/50 * 100 = 50%
        res_off = client.get("/api/statewide-top-performers?month=2026-10&period=monthly&target_mode=official", headers=headers)
        assert res_off.status_code == 200
        d_off = res_off.json()
        assert d_off["success"] is True
        assert d_off["target_mode"] == "official"
        jamui_off = [d for d in d_off["top_districts"] if d["district"] == "Jamui"][0]
        assert jamui_off["target"] == 50
        assert jamui_off["notifications"] == 25
        assert jamui_off["percentage"] == 50.0

        # 2. Frontline mode: Jamui target should be 70 (35 + 35), percentage = 25/70 * 100 ~ 35.7%
        res_front = client.get("/api/statewide-top-performers?month=2026-10&period=monthly&target_mode=frontline", headers=headers)
        assert res_front.status_code == 200
        d_front = res_front.json()
        assert d_front["success"] is True
        assert d_front["target_mode"] == "frontline"
        jamui_front = [d for d in d_front["top_districts"] if d["district"] == "Jamui"][0]
        assert jamui_front["target"] == 70
        assert jamui_front["notifications"] == 25
        assert round(jamui_front["percentage"], 1) == 35.7

        # 3. Verify distinct cache keys exist
        assert cache.get("statewide_top_2026-10_monthly_official") is not None
        assert cache.get("statewide_top_2026-10_monthly_frontline") is not None

    # 4. Updating district target evicts statewide_top_ caches
    with patch("main.db.collection") as mock_coll:
        mock_coll.return_value.document.return_value = MagicMock()
        update_res = client.post("/update-district-target", json={
            "month": "2026-10",
            "district": "Jamui",
            "official_target": 60
        }, headers=headers)
        assert update_res.status_code == 200

        assert cache.get("statewide_top_2026-10_monthly_official") is None
        assert cache.get("statewide_top_2026-10_monthly_frontline") is None




