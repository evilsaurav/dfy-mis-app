# -*- coding: utf-8 -*-
"""
tests/test_kpi_template_generator.py
Unit and integration tests for the Master KPI Template Generator (17 Clinical Indicators).
"""
import os
import sys
import openpyxl
import pytest

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))
import generate_templates
from openpyxl.utils import get_column_letter

EXPECTED_17_CATEGORIES = [
    ("NOTIFICATION", "notification_ids"),
    ("HIV & DM", "hiv_dm_ids"),
    ("DBT", "dbt_ids"),
    ("SAMPLE COLLECTION", "sample_collection_ids"),
    ("SAMPLE TESTED", "sample_tested_ids"),
    ("Outcome Assigned", "outcome_assigned_ids"),
    ("Home Visit", "home_visit_ids"),
    ("Contact Tracing", "contact_tracing_ids"),
    ("Follow Up", "follow_up_ids"),
    ("Face to Face", "face_to_face_ids"),
    ("Presumptive", "presumptive_ids"),
    ("Documents", "documents_ids"),
    ("FDC Provided", "fdc_provided_ids"),
    ("Kit Consumption", "kit_consumption_ids"),
    ("DIFF TB", "differentiated_tb_ids"),
    ("TPT START", "tpt_treatment_start_ids"),
    ("TPT PRESUMTIVE", "tpt_presumptive_ids"),
]


def test_kpi_categories_definition():
    """Verify that KPI_CATEGORIES contains all 17 clinical indicators in order."""
    assert len(generate_templates.KPI_CATEGORIES) == 17
    assert generate_templates.KPI_CATEGORIES == EXPECTED_17_CATEGORIES


def test_generate_single_district_template_structure(tmp_path):
    """Test generating a template in-memory/temp directory and verify structure across all 33 sheets."""
    district = "TestDistrict"
    staff = ["Staff Alpha", "Staff Beta"]
    out_file = str(tmp_path / "template_TestDistrict.xlsx")

    generate_templates.generate_district_template(district, staff, out_file)
    assert os.path.exists(out_file)

    wb = openpyxl.load_workbook(out_file)

    # 1. Total Sheet Count must be 33
    assert len(wb.sheetnames) == 33
    assert wb.sheetnames[0] == "Performance sheet"
    assert wb.sheetnames[1] == "CONSOLIDATED SHEET"
    assert wb.sheetnames[2] == "1ST"
    assert wb.sheetnames[-1] == "31st"

    # 2. Performance Sheet Verification
    ws_perf = wb["Performance sheet"]
    perf_headers = [ws_perf.cell(row=4, column=c).value for c in range(1, 22)]
    expected_perf_headers = [
        "Employee Name",
        "DESIG.",
        "Target",
        "NOTIFICATION",
        "% Achieved",
    ] + [k[0] for k in EXPECTED_17_CATEGORIES[1:]]
    assert len(perf_headers) == 21
    assert perf_headers == expected_perf_headers

    # Check staff rows in Performance sheet
    for idx, s_name in enumerate(staff):
        r = 5 + idx
        assert ws_perf.cell(row=r, column=1).value == s_name
        assert ws_perf.cell(row=r, column=2).value == "Field Officer"
        assert ws_perf.cell(row=r, column=3).value == 50
        assert ws_perf.cell(row=r, column=4).value == 0
        assert ws_perf.cell(row=r, column=5).value == f"=IF(C{r}>0, D{r}/C{r}, 0)"
        for c in range(6, 22):
            assert ws_perf.cell(row=r, column=c).value == 0

    # Grand total row
    tot_r = 5 + len(staff)
    assert ws_perf.cell(row=tot_r, column=1).value == "GRAND TOTAL"
    assert ws_perf.cell(row=tot_r, column=3).value == f"=SUM(C5:C{tot_r-1})"
    assert ws_perf.cell(row=tot_r, column=4).value == f"=SUM(D5:D{tot_r-1})"
    assert ws_perf.cell(row=tot_r, column=5).value == f"=IF(C{tot_r}>0, D{tot_r}/C{tot_r}, 0)"
    for c in range(6, 22):
        col_letter = get_column_letter(c)
        assert ws_perf.cell(row=tot_r, column=c).value == f"=SUM({col_letter}5:{col_letter}{tot_r-1})"

    # Column widths in Performance sheet
    for c in range(1, 22):
        col_letter = get_column_letter(c)
        assert ws_perf.column_dimensions[col_letter].width is not None

    # 3. CONSOLIDATED SHEET Verification
    ws_cons = wb["CONSOLIDATED SHEET"]
    # Left Wing: 16 clusters (Kit Consumption excluded), 3 cols each -> Cols 1 to 48
    expected_left_clusters = [k[0] for k in EXPECTED_17_CATEGORIES if k[0] != "Kit Consumption"]
    assert len(expected_left_clusters) == 16

    for c_idx, cluster_name in enumerate(expected_left_clusters):
        start_c = 1 + (c_idx * 3)
        assert ws_cons.cell(row=1, column=start_c).value == cluster_name
        assert ws_cons.cell(row=2, column=start_c).value == 0
        assert ws_cons.cell(row=3, column=start_c).value == "Patient ID"
        assert ws_cons.cell(row=3, column=start_c + 1).value == "Date"
        assert ws_cons.cell(row=3, column=start_c + 2).value == "Reported by"

    # Specific check for new clusters 14, 15, 16
    # Cluster 14 (Cols 40-42): DIFF TB
    assert ws_cons.cell(row=1, column=40).value == "DIFF TB"
    assert ws_cons.cell(row=3, column=40).value == "Patient ID"
    assert ws_cons.cell(row=3, column=41).value == "Date"
    assert ws_cons.cell(row=3, column=42).value == "Reported by"

    # Cluster 15 (Cols 43-45): TPT START
    assert ws_cons.cell(row=1, column=43).value == "TPT START"
    assert ws_cons.cell(row=3, column=43).value == "Patient ID"
    assert ws_cons.cell(row=3, column=44).value == "Date"
    assert ws_cons.cell(row=3, column=45).value == "Reported by"

    # Cluster 16 (Cols 46-48): TPT PRESUMTIVE
    assert ws_cons.cell(row=1, column=46).value == "TPT PRESUMTIVE"
    assert ws_cons.cell(row=3, column=46).value == "Patient ID"
    assert ws_cons.cell(row=3, column=47).value == "Date"
    assert ws_cons.cell(row=3, column=48).value == "Reported by"

    # Right Wing: Staff breakdown starting at Column 49
    for s_idx, s_name in enumerate(staff):
        staff_start_c = 49 + (s_idx * 17)
        staff_end_c = staff_start_c + 16
        assert ws_cons.cell(row=1, column=staff_start_c).value == s_name
        for k_idx, (kpi_name, _) in enumerate(EXPECTED_17_CATEGORIES):
            c_num = staff_start_c + k_idx
            assert ws_cons.cell(row=2, column=c_num).value == kpi_name
            assert ws_cons.cell(row=3, column=c_num).value == 0

    # 4. Daily Sheets Verification (e.g. '1ST' and '31st')
    for test_day in ["1ST", "15th", "31st"]:
        ws_day = wb[test_day]
        headers = [ws_day.cell(row=1, column=c).value for c in range(1, 20)]
        expected_daily_headers = ["NAME", "DESIGNATION"] + [k[0] for k in EXPECTED_17_CATEGORIES]
        assert len(headers) == 19
        assert headers == expected_daily_headers
        assert headers[16] == "DIFF TB"
        assert headers[17] == "TPT START"
        assert headers[18] == "TPT PRESUMTIVE"

        # Check staff rows for Daily sheet
        for s_idx, s_name in enumerate(staff):
            fo_start_r = 2 + (s_idx * 40)
            assert ws_day.cell(row=fo_start_r, column=1).value == s_name
            assert ws_day.cell(row=fo_start_r, column=2).value == "Field Officer"

    wb.close()


def test_existing_templates_conform_to_17_kpis():
    """Verify that templates in the templates/ folder conform to the 17 KPI structure."""
    templates_dir = os.path.join(os.path.dirname(__file__), "..", "templates")
    template_files = [f for f in os.listdir(templates_dir) if f.startswith("template_") and f.endswith(".xlsx")]
    assert len(template_files) == 22, f"Expected 22 templates, found {len(template_files)}"

    # Sample test across 3 districts: small, medium, large
    sample_templates = ["template_Khagaria.xlsx", "template_Begusarai.xlsx", "template_Muzaffarpur.xlsx"]
    for t_file in sample_templates:
        fpath = os.path.join(templates_dir, t_file)
        if not os.path.exists(fpath):
            continue
        wb = openpyxl.load_workbook(fpath, data_only=False)
        assert len(wb.sheetnames) == 33

        # Check Performance sheet has 21 columns
        ws_perf = wb["Performance sheet"]
        assert ws_perf.cell(row=4, column=19).value == "DIFF TB"
        assert ws_perf.cell(row=4, column=20).value == "TPT START"
        assert ws_perf.cell(row=4, column=21).value == "TPT PRESUMTIVE"

        # Check Consolidated sheet has Wing 1 ending at col 48 and Wing 2 starting at 49
        ws_cons = wb["CONSOLIDATED SHEET"]
        assert ws_cons.cell(row=1, column=40).value == "DIFF TB"
        assert ws_cons.cell(row=1, column=43).value == "TPT START"
        assert ws_cons.cell(row=1, column=46).value == "TPT PRESUMTIVE"
        assert ws_cons.cell(row=2, column=49).value == "NOTIFICATION"

        # Check Daily sheet '1ST' has 19 columns
        ws_day = wb["1ST"]
        assert ws_day.cell(row=1, column=17).value == "DIFF TB"
        assert ws_day.cell(row=1, column=18).value == "TPT START"
        assert ws_day.cell(row=1, column=19).value == "TPT PRESUMTIVE"
        wb.close()
