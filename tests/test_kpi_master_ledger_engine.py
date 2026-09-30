# -*- coding: utf-8 -*-
"""
tests/test_kpi_master_ledger_engine.py
Unit and integration tests for Zero-Read Bulk KPI Excel Engine & LRU 2-Month RAM Watchdog.
"""
import io
import os
import sys
from pathlib import Path
from unittest.mock import patch, AsyncMock, MagicMock
import openpyxl
import pytest
from fastapi.testclient import TestClient

BASE_DIR = Path(__file__).resolve().parent.parent
if str(BASE_DIR) not in sys.path:
    sys.path.insert(0, str(BASE_DIR))

import main
from main import app, cache, create_access_token

client = TestClient(app)

@pytest.fixture(autouse=True)
def clean_cache_and_watchdog():
    cache.clear()
    if hasattr(main, "ACTIVE_MONTHLY_CACHE_KEYS"):
        main.ACTIVE_MONTHLY_CACHE_KEYS.clear()
    yield
    cache.clear()
    if hasattr(main, "ACTIVE_MONTHLY_CACHE_KEYS"):
        main.ACTIVE_MONTHLY_CACHE_KEYS.clear()

def make_admin_token(role: str = "SUPER_ADMIN", allowed_districts=None, username="super_admin", user_id="admin_123"):
    return create_access_token({
        "user_id": user_id,
        "username": username,
        "name": "Super Admin",
        "role": role,
        "allowed_districts": allowed_districts if allowed_districts is not None else ["All"]
    })


def test_generate_district_kpi_bytes_accepts_raw_reports_and_targets_zero_firestore_reads():
    """
    Verify generate_district_kpi_bytes accepts raw_reports and target_records and computes
    workbook bytes entirely in-memory with ZERO Firestore queries (db.collection calls = 0).
    """
    district = "Khagaria"
    month_prefix = "2026-09"

    mock_raw_reports = [
        {
            "id": "rep1",
            "working_place": "Khagaria",
            "fo_name": "Rambilash Paswan",
            "date_of_reporting": "2026-09-01",
            "notification_ids": ["NOTIF001", "NOTIF002"],
            "hiv_dm_ids": ["NOTIF001", "PREV001"],
            "dbt_ids": ["NOTIF001"],
            "sample_collection_ids": ["NOTIF001", "NOTIF002", "NOTIF003"],
            "sample_tested_ids": ["NOTIF001", "PREV002"],
            "outcome_assigned_ids": ["NOTIF001"],
            "home_visit_ids": ["NOTIF001"],
            "contact_tracing_ids": ["NOTIF001"],
            "follow_up_ids": ["NOTIF001"],
            "face_to_face_ids": ["NOTIF001"],
            "presumptive_ids": ["PRESUMP001"],
            "documents_ids": ["NOTIF001"],
            "fdc_provided_ids": ["NOTIF001"],
            "kit_consumption_ids": ["KIT001", "KIT002"],
            "differentiated_tb_ids": ["NOTIF001"],
            "tpt_treatment_start_ids": ["TPT_ST_001"],
            "tpt_presumptive_ids": ["TPT_PRE_001"],
        }
    ]

    mock_target_records = [
        {
            "district": "Khagaria",
            "fo_name": "Rambilash Paswan",
            "month": "2026-09",
            "target": 65
        }
    ]

    with patch.object(main, "db") as mock_db:
        mock_db.collection.side_effect = AssertionError("db.collection must NOT be called when raw_reports and target_records are passed!")

        excel_bytes = main.generate_district_kpi_bytes(
            district=district,
            month_prefix=month_prefix,
            raw_reports=mock_raw_reports,
            target_records=mock_target_records
        )

        assert excel_bytes is not None
        assert isinstance(excel_bytes, bytes)
        assert len(excel_bytes) > 0

        # Load generated workbook and verify target and notification values
        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes), data_only=False)
        assert "Performance sheet" in wb.sheetnames
        ws_perf = wb["Performance sheet"]

        # Check target row for Rambilash Paswan (Col 3 Target = 65, Col 4 Notification = 2)
        found_fo = False
        for row_idx in range(5, ws_perf.max_row + 1):
            val = ws_perf.cell(row=row_idx, column=1).value
            if val and "rambilash" in str(val).lower():
                found_fo = True
                assert ws_perf.cell(row=row_idx, column=3).value == 65
                assert ws_perf.cell(row=row_idx, column=4).value == 2
                break
        assert found_fo, "Rambilash Paswan not found in Performance sheet!"


@pytest.mark.asyncio
async def test_generate_district_kpi_bytes_async():
    """Verify generate_district_kpi_bytes_async fetches raw monthly and cached targets and delegates to thread."""
    mock_raw_reports = [
        {
            "id": "rep1",
            "working_place": "Khagaria",
            "fo_name": "Rambilash Paswan",
            "date_of_reporting": "2026-09-02",
            "notification_ids": ["NOTIF100"],
        }
    ]
    mock_targets = [
        {"district": "Khagaria", "fo_name": "Rambilash Paswan", "month": "2026-09", "target": 75}
    ]

    assert hasattr(main, "generate_district_kpi_bytes_async"), "main must implement generate_district_kpi_bytes_async"

    with patch("main.get_raw_monthly_reports", AsyncMock(return_value=mock_raw_reports)) as mock_raw, \
         patch("main.get_cached_staff_targets_for_month", AsyncMock(return_value=mock_targets)) as mock_targ, \
         patch.object(main, "db") as mock_db:

        mock_db.collection.side_effect = AssertionError("db.collection should not be called!")

        excel_bytes = await main.generate_district_kpi_bytes_async("Khagaria", "2026-09")
        assert excel_bytes is not None
        assert isinstance(excel_bytes, bytes)
        mock_raw.assert_awaited_once_with("2026-09")
        mock_targ.assert_awaited_once_with("2026-09")


def test_download_all_kpi_workbooks_prefetches_once_and_zero_per_district_reads():
    """Verify /download-all-kpi-workbooks pre-fetches monthly reports and targets ONCE before district loop."""
    mock_raw_monthly = [
        {
            "id": "rep_khagaria",
            "working_place": "Khagaria",
            "fo_name": "Rambilash Paswan",
            "date_of_reporting": "2026-09-05",
            "notification_ids": ["NOTIF_K1"],
        },
        {
            "id": "rep_begusarai",
            "working_place": "Begusarai",
            "fo_name": "Sample Officer",
            "date_of_reporting": "2026-09-05",
            "notification_ids": ["NOTIF_B1"],
        }
    ]
    mock_cached_targets = [
        {"district": "Khagaria", "fo_name": "Rambilash Paswan", "month": "2026-09", "target": 50},
        {"district": "Begusarai", "fo_name": "Sample Officer", "month": "2026-09", "target": 50}
    ]

    with patch("main.get_raw_monthly_reports", AsyncMock(return_value=mock_raw_monthly)) as mock_raw, \
         patch("main.get_cached_staff_targets_for_month", AsyncMock(return_value=mock_cached_targets)) as mock_targets, \
         patch.object(main, "db") as mock_db:

        def forbidden_collection(name):
            if name in ("daily_field_reports", "staff_targets"):
                raise AssertionError(f"Collection '{name}' must NOT be queried during bulk KPI generation!")
            col = MagicMock()
            col.where.return_value = col
            col.stream.return_value = []
            return col

        mock_db.collection.side_effect = forbidden_collection

        token = make_admin_token()
        # Request 2 districts
        resp = client.get(
            "/download-all-kpi-workbooks?month=2026-09&districts=Khagaria,Begusarai",
            headers={"Authorization": f"Bearer {token}"}
        )

        assert resp.status_code == 200, resp.text
        assert resp.headers["content-type"] == "application/zip"
        assert len(resp.content) > 0

        # Assert get_raw_monthly_reports and get_cached_staff_targets_for_month were called exactly ONCE outside the loop
        assert mock_raw.call_count == 1
        assert mock_targets.call_count == 1


@pytest.mark.asyncio
async def test_ram_watchdog_lru_two_month_cache_bound():
    """
    Verify LRU 2-Month Memory Bound & RAM Watchdog:
    ACTIVE_MONTHLY_CACHE_KEYS tracks loaded months, enforces upper bound of at most 2 active months,
    evicting the oldest key and its sub-keys with gc.collect().
    """
    assert hasattr(main, "ACTIVE_MONTHLY_CACHE_KEYS"), "main must define ACTIVE_MONTHLY_CACHE_KEYS"
    main.ACTIVE_MONTHLY_CACHE_KEYS.clear()
    cache.clear()

    # Create mock firestore return for 3 different months
    def make_doc_snap(m_str):
        doc = MagicMock()
        doc.id = f"doc_{m_str}"
        doc.to_dict.return_value = {
            "id": f"doc_{m_str}",
            "working_place": "Patna",
            "fo_name": "Test FO",
            "date_of_reporting": f"{m_str}-15"
        }
        return doc

    with patch.object(main, "db") as mock_db, \
         patch("gc.collect") as mock_gc:

        mock_coll = MagicMock()
        mock_coll.where.return_value = mock_coll
        mock_coll.stream.side_effect = lambda: [make_doc_snap("curr")]
        mock_db.collection.return_value = mock_coll

        # Month 1: 2026-07
        await main.get_raw_monthly_reports("2026-07")
        assert "shared_raw_month_2026-07" in main.ACTIVE_MONTHLY_CACHE_KEYS
        assert len(main.ACTIVE_MONTHLY_CACHE_KEYS) == 1
        assert cache.get("shared_raw_month_2026-07") is not None

        # Month 2: 2026-08
        await main.get_raw_monthly_reports("2026-08")
        assert "shared_raw_month_2026-08" in main.ACTIVE_MONTHLY_CACHE_KEYS
        assert len(main.ACTIVE_MONTHLY_CACHE_KEYS) == 2
        assert cache.get("shared_raw_month_2026-07") is not None
        assert cache.get("shared_raw_month_2026-08") is not None

        # Month 3: 2026-09 -> Should trigger eviction of Month 1 (2026-07)
        await main.get_raw_monthly_reports("2026-09")
        assert len(main.ACTIVE_MONTHLY_CACHE_KEYS) == 2
        assert "shared_raw_month_2026-07" not in main.ACTIVE_MONTHLY_CACHE_KEYS
        assert "shared_raw_month_2026-08" in main.ACTIVE_MONTHLY_CACHE_KEYS
        assert "shared_raw_month_2026-09" in main.ACTIVE_MONTHLY_CACHE_KEYS

        # Verify 2026-07 cache was evicted from cache
        assert cache.get("shared_raw_month_2026-07") is None
        assert cache.get("shared_raw_month_2026-08") is not None
        assert cache.get("shared_raw_month_2026-09") is not None

        # Verify gc.collect() was called during eviction
        assert mock_gc.called

        # Touch 2026-08 (MRU update)
        await main.get_raw_monthly_reports("2026-08")
        # Order should now be ["shared_raw_month_2026-09", "shared_raw_month_2026-08"]
        assert main.ACTIVE_MONTHLY_CACHE_KEYS == ["shared_raw_month_2026-09", "shared_raw_month_2026-08"]

        # Month 4: 2026-10 -> Should evict 2026-09, preserving 2026-08
        await main.get_raw_monthly_reports("2026-10")
        assert len(main.ACTIVE_MONTHLY_CACHE_KEYS) == 2
        assert "shared_raw_month_2026-09" not in main.ACTIVE_MONTHLY_CACHE_KEYS
        assert "shared_raw_month_2026-08" in main.ACTIVE_MONTHLY_CACHE_KEYS
        assert "shared_raw_month_2026-10" in main.ACTIVE_MONTHLY_CACHE_KEYS
        assert cache.get("shared_raw_month_2026-09") is None


def test_download_excel_derives_from_get_raw_monthly_reports():
    """Verify /download-excel derives data from get_raw_monthly_reports(target_month) instead of raw stream."""
    mock_reports = [
        {
            "id": "doc1",
            "working_place": "Khagaria",
            "fo_name": "Rambilash Paswan",
            "designation": "Field Officer",
            "date_of_reporting": "2026-09-10",
            "notification_ids": ["N1"],
            "morning_km": 10,
            "evening_km": 15,
            "total_km": 25,
            "visited_names": ["Dr. Smith"]
        }
    ]

    with patch("main.get_raw_monthly_reports", AsyncMock(return_value=mock_reports)) as mock_raw, \
         patch.object(main, "db") as mock_db:

        mock_db.collection.side_effect = AssertionError("db.collection('daily_field_reports') must NOT be called in /download-excel!")

        token = make_admin_token()
        resp = client.get("/download-excel?month=2026-09", headers={"Authorization": f"Bearer {token}"})

        assert resp.status_code == 200, resp.text
        assert resp.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        assert len(resp.content) > 0
        mock_raw.assert_awaited_once_with("2026-09")


def test_download_kpi_workbook_endpoint_uses_async_master_ledger():
    """Verify calling /download-kpi-workbook routes through generate_district_kpi_bytes_async with zero raw report firestore streams."""
    dummy_bytes = b"PK\x03\x04mockkpibytes"
    token = make_admin_token()

    with patch("main.generate_district_kpi_bytes_async", AsyncMock(return_value=dummy_bytes)) as mock_async_gen, \
         patch.object(main, "db") as mock_db:

        def forbidden_coll(name):
            if name in ("daily_field_reports", "staff_targets"):
                raise AssertionError(f"Collection '{name}' must NOT be streamed directly during /download-kpi-workbook!")
            col = MagicMock()
            col.where.return_value = col
            col.stream.return_value = []
            return col

        mock_db.collection.side_effect = forbidden_coll

        resp = client.get(
            "/download-kpi-workbook?district=Patna&month=2026-09",
            headers={"Authorization": f"Bearer {token}"}
        )

        assert resp.status_code == 200, resp.text
        assert resp.content == dummy_bytes
        assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in resp.headers["content-type"]
        assert "KPI_Report_Patna_2026-09.xlsx" in resp.headers["content-disposition"]
        mock_async_gen.assert_awaited_once_with("Patna", "2026-09")

