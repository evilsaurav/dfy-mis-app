# -*- coding: utf-8 -*-
"""
tests/test_kpi_excel_engine_17_indicators.py
Unit and integration tests for Backend KPI Excel Data Engine & Dynamic Population (17 Clinical Indicators).
"""
import io
import os
import sys
from pathlib import Path
from unittest.mock import patch, MagicMock
import openpyxl
import pytest

BASE_DIR = Path(__file__).resolve().parent.parent
if str(BASE_DIR) not in sys.path:
    sys.path.insert(0, str(BASE_DIR))

import main


EXPECTED_17_CATEGORIES = [
    ("NOTIFICATION", "notification_ids", 3),
    ("HIV & DM", "hiv_dm_ids", 4),
    ("DBT", "dbt_ids", 5),
    ("SAMPLE COLLECTION", "sample_collection_ids", 6),
    ("SAMPLE TESTED", "sample_tested_ids", 7),
    ("Outcome Assigned", "outcome_assigned_ids", 8),
    ("Home Visit", "home_visit_ids", 9),
    ("Contact Tracing", "contact_tracing_ids", 10),
    ("Follow Up", "follow_up_ids", 11),
    ("Face to Face", "face_to_face_ids", 12),
    ("Presumptive", "presumptive_ids", 13),
    ("Documents", "documents_ids", 14),
    ("FDC Provided", "fdc_provided_ids", 15),
    ("Kit Consumption", "kit_consumption_ids", 16),
    ("DIFF TB", "differentiated_tb_ids", 17),
    ("TPT START", "tpt_treatment_start_ids", 18),
    ("TPT PRESUMTIVE", "tpt_presumptive_ids", 19),
]


def test_excel_kpi_categories_definition():
    """Verify that EXCEL_KPI_CATEGORIES contains all 17 clinical indicators with columns 3 to 19."""
    assert len(main.EXCEL_KPI_CATEGORIES) == 17
    assert main.EXCEL_KPI_CATEGORIES == EXPECTED_17_CATEGORIES
    for idx, (_, _, col_idx) in enumerate(main.EXCEL_KPI_CATEGORIES):
        assert col_idx == 3 + idx


class MockDocSnap:
    def __init__(self, doc_id: str, data: dict):
        self.id = doc_id
        self._data = data

    def to_dict(self):
        return self._data


class MockQuery:
    def __init__(self, coll_name: str, docs: list, filters=None):
        self.coll_name = coll_name
        self.docs = docs
        self.filters = filters or []

    def where(self, field, op, val):
        new_filters = list(self.filters) + [(field, op, val)]
        return MockQuery(self.coll_name, self.docs, new_filters)

    def stream(self):
        matched = []
        for doc in self.docs:
            data = doc.to_dict()
            match = True
            for field, op, val in self.filters:
                if op == "==" and data.get(field) != val:
                    match = False
                    break
                elif op == ">=" and str(data.get(field, "")) < str(val):
                    match = False
                    break
                elif op == "<=" and str(data.get(field, "")) > str(val):
                    match = False
                    break
            if match:
                matched.append(doc)
        return matched


def test_generate_district_kpi_bytes_population():
    """Verify that generate_district_kpi_bytes accurately populates all 17 indicators across sheets."""
    district = "Khagaria"
    month_prefix = "2026-09"

    mock_reports = [
        MockDocSnap("rep1", {
            "working_place": "Khagaria",
            "fo_name": "Rambilash Paswan",
            "date_of_reporting": "2026-09-01",
            "notification_ids": ["NOTIF001", "NOTIF002"],
            "hiv_dm_ids": ["NOTIF001", "PREV001"],
            "dbt_ids": ["NOTIF001"],
            "sample_collection_ids": ["SAMP001"],
            "sample_tested_ids": ["NOTIF002", "PREV002"],
            "outcome_assigned_ids": ["OUT001"],
            "home_visit_ids": ["HV001"],
            "contact_tracing_ids": ["PREV003"],
            "follow_up_ids": ["FU001"],
            "face_to_face_ids": ["F2F001"],
            "presumptive_ids": ["PRES001"],
            "documents_ids": ["DOC001"],
            "fdc_provided_ids": ["FDC001"],
            "kit_consumption_ids": ["KIT1", "KIT2", "KIT3", "KIT4"],
            "differentiated_tb_ids": ["DIFF001", "DIFF002"],
            "tpt_treatment_start_ids": ["TPTS001"],
            "tpt_presumptive_ids": ["TPTP001", "TPTP002", "TPTP003"],
        }),
        MockDocSnap("rep2", {
            "working_place": "Khagaria",
            "fo_name": "Ghanshyam kumar",
            "date_of_reporting": "2026-09-02",
            "notification_ids": ["NOTIF003"],
            "differentiated_tb_ids": ["DIFF003"],
            "tpt_treatment_start_ids": ["TPTS002"],
            "tpt_presumptive_ids": ["TPTP004"],
        })
    ]

    mock_targets = [
        MockDocSnap("tgt1", {"district": "Khagaria", "fo_name": "Rambilash Paswan", "month": "2026-09", "target": 60}),
        MockDocSnap("tgt2", {"district": "Khagaria", "fo_name": "Ghanshyam kumar", "month": "2026-09", "target": 45}),
        MockDocSnap("tgt3", {"district": "Khagaria", "fo_name": "Rajnish Kumar", "month": "2026-09", "target": 50}),
    ]

    def mock_collection_impl(name):
        if name == "daily_field_reports":
            return MockQuery("daily_field_reports", mock_reports)
        elif name == "staff_targets":
            return MockQuery("staff_targets", mock_targets)
        return MockQuery(name, [])

    with patch.object(main.db, "collection", side_effect=mock_collection_impl):
        kpi_bytes = main.generate_district_kpi_bytes(district, month_prefix)

    assert kpi_bytes is not None
    assert len(kpi_bytes) > 0

    wb = openpyxl.load_workbook(io.BytesIO(kpi_bytes), data_only=False)

    # 1. Verify Daily Tab '1ST' (Day 1 for Rambilash Paswan)
    assert "1ST" in wb.sheetnames
    ws_day1 = wb["1ST"]

    # Staff 0 starts at row 2
    assert ws_day1.cell(row=2, column=1).value == "Rambilash Paswan"
    # Col 17: DIFF TB (DIFF001, DIFF002)
    assert ws_day1.cell(row=2, column=17).value == "DIFF001"
    assert ws_day1.cell(row=3, column=17).value == "DIFF002"
    # Col 18: TPT START (TPTS001)
    assert ws_day1.cell(row=2, column=18).value == "TPTS001"
    # Col 19: TPT PRESUMTIVE (TPTP001, TPTP002, TPTP003)
    assert ws_day1.cell(row=2, column=19).value == "TPTP001"
    assert ws_day1.cell(row=3, column=19).value == "TPTP002"
    assert ws_day1.cell(row=4, column=19).value == "TPTP003"

    # 2. Verify 'CONSOLIDATED SHEET'
    assert "CONSOLIDATED SHEET" in wb.sheetnames
    ws_cons = wb["CONSOLIDATED SHEET"]

    # Wing 1 (District Master Rollup - 16 clusters in Cols 1-48)
    # Row 2 Grand Totals:
    # Cluster 13 (DIFF TB, Cols 40-42): Rambilash (2) + Ghanshyam (1) = 3
    assert ws_cons.cell(row=2, column=40).value == 3
    # Cluster 14 (TPT START, Cols 43-45): Rambilash (1) + Ghanshyam (1) = 2
    assert ws_cons.cell(row=2, column=43).value == 2
    # Cluster 15 (TPT PRESUMTIVE, Cols 46-48): Rambilash (3) + Ghanshyam (1) = 4
    assert ws_cons.cell(row=2, column=46).value == 4

    # Patient rows in Wing 1:
    # DIFF TB rows starting at row 4
    assert ws_cons.cell(row=4, column=40).value == "DIFF001"
    assert ws_cons.cell(row=4, column=41).value == "2026-09-01"
    assert ws_cons.cell(row=4, column=42).value == "Rambilash Paswan"
    assert ws_cons.cell(row=5, column=40).value == "DIFF002"
    assert ws_cons.cell(row=6, column=40).value == "DIFF003"

    # TPT START rows starting at row 4
    assert ws_cons.cell(row=4, column=43).value == "TPTS001"
    assert ws_cons.cell(row=5, column=43).value == "TPTS002"

    # TPT PRESUMTIVE rows starting at row 4
    assert ws_cons.cell(row=4, column=46).value == "TPTP001"
    assert ws_cons.cell(row=5, column=46).value == "TPTP002"
    assert ws_cons.cell(row=6, column=46).value == "TPTP003"
    assert ws_cons.cell(row=7, column=46).value == "TPTP004"

    # Wing 2 (Staff Breakdown - starting at Col 49, 17 cols per staff)
    # Staff 0 (Rambilash Paswan): base col = 49
    # Row 3 Staff KPI Totals:
    # Col 49+14 = 63 (DIFF TB): 2
    assert ws_cons.cell(row=3, column=63).value == 2
    # Col 49+15 = 64 (TPT START): 1
    assert ws_cons.cell(row=3, column=64).value == 1
    # Col 49+16 = 65 (TPT PRESUMTIVE): 3
    assert ws_cons.cell(row=3, column=65).value == 3

    # Patient IDs logged under Staff 0
    assert ws_cons.cell(row=4, column=63).value == "DIFF001"
    assert ws_cons.cell(row=5, column=63).value == "DIFF002"
    assert ws_cons.cell(row=4, column=64).value == "TPTS001"
    assert ws_cons.cell(row=4, column=65).value == "TPTP001"
    assert ws_cons.cell(row=5, column=65).value == "TPTP002"
    assert ws_cons.cell(row=6, column=65).value == "TPTP003"

    # Boundary border on 17th column (Col 65)
    c65 = ws_cons.cell(row=4, column=65)
    assert c65.border.right.style == "medium"

    # 3. Verify 'Performance sheet'
    assert "Performance sheet" in wb.sheetnames
    ws_perf = wb["Performance sheet"]

    # Shifted cohort headers in Row 4
    assert ws_perf.cell(row=4, column=19).value == "DIFF TB"
    assert ws_perf.cell(row=4, column=20).value == "TPT START"
    assert ws_perf.cell(row=4, column=21).value == "TPT PRESUMTIVE"
    assert ws_perf.cell(row=4, column=22).value == "HIV\n(Cur Month)"
    assert ws_perf.cell(row=4, column=23).value == "HIV\n(Prev Backlog)"
    assert ws_perf.cell(row=4, column=24).value == "UDST\n(Cur Month)"
    assert ws_perf.cell(row=4, column=25).value == "UDST\n(Prev Backlog)"
    assert ws_perf.cell(row=4, column=26).value == "Contact Tr\n(Cur Month)"
    assert ws_perf.cell(row=4, column=27).value == "Contact Tr\n(Prev Backlog)"

    # Staff 0 (Row 5 - Rambilash Paswan)
    assert ws_perf.cell(row=5, column=1).value == "Rambilash Paswan"
    assert ws_perf.cell(row=5, column=3).value == 60  # Target from target_map
    assert ws_perf.cell(row=5, column=4).value == 2   # Notification count
    assert ws_perf.cell(row=5, column=19).value == 2  # DIFF TB count
    assert ws_perf.cell(row=5, column=20).value == 1  # TPT START count
    assert ws_perf.cell(row=5, column=21).value == 3  # TPT PRESUMTIVE count

    # Staff 0 Cohort counts (Cols 22 to 27)
    # HIV: NOTIF001 is current month, PREV001 is backlog
    assert ws_perf.cell(row=5, column=22).value == 1  # hiv_cur
    assert ws_perf.cell(row=5, column=23).value == 1  # hiv_prev
    # UDST: NOTIF002 is current month, PREV002 is backlog
    assert ws_perf.cell(row=5, column=24).value == 1  # udst_cur
    assert ws_perf.cell(row=5, column=25).value == 1  # udst_prev
    # Contact Tracing: PREV003 is backlog
    assert ws_perf.cell(row=5, column=26).value == 0  # con_cur
    assert ws_perf.cell(row=5, column=27).value == 1  # con_prev

    # Grand Total Row (Row 8 in Khagaria template)
    gt_row = 8
    assert ws_perf.cell(row=gt_row, column=1).value == "GRAND TOTAL"
    assert ws_perf.cell(row=gt_row, column=19).value == "=SUM(S5:S7)"
    assert ws_perf.cell(row=gt_row, column=20).value == "=SUM(T5:T7)"
    assert ws_perf.cell(row=gt_row, column=21).value == "=SUM(U5:U7)"
    assert ws_perf.cell(row=gt_row, column=22).value == "=SUM(V5:V7)"
    assert ws_perf.cell(row=gt_row, column=23).value == "=SUM(W5:W7)"
    assert ws_perf.cell(row=gt_row, column=24).value == "=SUM(X5:X7)"
    assert ws_perf.cell(row=gt_row, column=25).value == "=SUM(Y5:Y7)"
    assert ws_perf.cell(row=gt_row, column=26).value == "=SUM(Z5:Z7)"
    assert ws_perf.cell(row=gt_row, column=27).value == "=SUM(AA5:AA7)"

    # Verify styling on Grand Total cohort cells
    for c in range(19, 28):
        cell = ws_perf.cell(row=gt_row, column=c)
        assert cell.font.bold is True
        fill_rgb = getattr(cell.fill.start_color, "rgb", "")
        assert fill_rgb in ["001E3A8A", "1E3A8A", "ff1e3a8a", "FF1E3A8A"]

    wb.close()


def test_in_loop_memory_release():
    """Verify that gc.collect is invoked to release openpyxl workbook memory."""
    district = "Khagaria"
    with patch.object(main.db, "collection", return_value=MockQuery("dummy", [])), \
         patch("gc.collect") as mock_gc_collect:
        main.generate_district_kpi_bytes(district, "2026-09")
        assert mock_gc_collect.called
