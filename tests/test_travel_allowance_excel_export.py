import io
import openpyxl
import pytest
from unittest.mock import patch, MagicMock
from fastapi.testclient import TestClient
from main import app, cache, create_access_token

client = TestClient(app)

class MockDocRef:
    def __init__(self, coll_name: str, doc_id: str, store: dict):
        self.coll_name = coll_name
        self.doc_id = doc_id
        self.store = store

    def get(self):
        snap = MagicMock()
        exists = self.coll_name in self.store and self.doc_id in self.store[self.coll_name]
        snap.exists = exists
        snap.id = self.doc_id
        snap.to_dict.return_value = dict(self.store.get(self.coll_name, {}).get(self.doc_id, {}))
        snap.reference = self
        return snap

class MockCollection:
    def __init__(self, coll_name: str, store: dict, filters=None):
        self.coll_name = coll_name
        self.store = store
        self.filters = filters or []

    def document(self, doc_id: str):
        return MockDocRef(self.coll_name, doc_id, self.store)

    def where(self, field, op, val):
        new_filters = list(self.filters) + [(field, op, val)]
        return MockCollection(self.coll_name, self.store, new_filters)

    def stream(self):
        docs = self.store.get(self.coll_name, {})
        for doc_id, data in docs.items():
            matches = True
            for field, op, val in self.filters:
                doc_val = data.get(field)
                if op == "==" and doc_val != val:
                    matches = False
                    break
            if matches:
                snap = MagicMock()
                snap.exists = True
                snap.id = doc_id
                snap.to_dict.return_value = dict(data)
                yield snap

@pytest.fixture(autouse=True)
def clear_cache():
    cache.clear()
    yield
    cache.clear()

def test_export_ta_excel_multi_sheet():
    store = {
        "admin_users": {
            "dc_gaya": {
                "name": "DC Gaya",
                "role": "SUB_ADMIN",
                "allowed_districts": ["Gaya"]
            }
        },
        "staff_directory": {
            "gaya_rameshkumar": {
                "district": "Gaya",
                "name": "Ramesh Kumar",
                "designation": "Field Officer",
                "status": "Active"
            },
            "gaya_anilsharma": {
                "district": "Gaya",
                "name": "Anil Sharma",
                "designation": "Lab Technician",
                "status": "Active"
            }
        },
        "travel_allowance_logs": {
            "2026-09_gaya_rameshkumar": {
                "month": "2026-09",
                "district": "Gaya",
                "staff_name": "Ramesh Kumar",
                "staff_key": "gaya_rameshkumar",
                "designation": "Field Officer",
                "total_km": 60,
                "gross_amount": 240.0,
                "deduction_amount": 40.0,
                "deduction_reason": "Excess route",
                "final_payable_amount": 200.0,
                "admin_final_remarks": "Verified and passed",
                "daily_logs": {
                    "2026-09-01": {
                        "initial_reading": 1000,
                        "final_reading": 1030,
                        "total_km": 30,
                        "is_override": False,
                        "rate": 4.0,
                        "amount": 120.0,
                        "from_location": "Gaya Sadar",
                        "to_location": "Bodhgaya",
                        "purpose": "Sample Collection",
                        "remarks": ""
                    },
                    "2026-09-02": {
                        "initial_reading": 1030,
                        "final_reading": 1060,
                        "total_km": 30,
                        "is_override": False,
                        "rate": 4.0,
                        "amount": 120.0,
                        "from_location": "Tekari",
                        "to_location": "Konch",
                        "purpose": "Doctor Visit",
                        "remarks": ""
                    }
                }
            }
        }
    }

    mock_db = MagicMock()
    mock_db.collection.side_effect = lambda c: MockCollection(c, store)

    token = create_access_token({
        "sub": "dc_gaya",
        "role": "SUB_ADMIN",
        "name": "DC Gaya",
        "allowed_districts": ["Gaya"]
    })
    headers = {"Authorization": f"Bearer {token}"}

    with patch("main.db", mock_db):
        res = client.get("/api/ta-logs/export-excel?month=2026-09&district=Gaya", headers=headers)
        assert res.status_code == 200, res.text
        assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in res.headers["content-type"]
        assert "attachment; filename=" in res.headers["content-disposition"]
        assert "DFY_TA_Bike_Log_Gaya_2026-09.xlsx" in res.headers["content-disposition"]

        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        assert "DASHBOARD" in wb.sheetnames
        assert wb.sheetnames[0] == "DASHBOARD"

        # Check DASHBOARD sheet structure
        ws_dash = wb["DASHBOARD"]
        assert "DOCTORS FOR YOU" in str(ws_dash["A1"].value)
        # Headers should be at row 4
        headers_row = [cell.value for cell in ws_dash[4]]
        assert "Sl. No" in headers_row
        assert "Employee Name" in headers_row
        assert "Designation" in headers_row
        assert "Total KM" in headers_row
        assert "Gross Amount (₹)" in headers_row
        assert "Deductions (₹)" in headers_row
        assert "Final Payable Amount (₹)" in headers_row

        # Check Staff Sheet exists
        assert "Ramesh Kumar" in wb.sheetnames
        ws_staff = wb["Ramesh Kumar"]
        assert "DOCTORS FOR YOU" in str(ws_staff["A1"].value)
        assert "Ramesh Kumar" in str(ws_staff["B3"].value)
        assert "4.00" in str(ws_staff["H3"].value) or "4" in str(ws_staff["H3"].value)

        # Verify daily table has entries
        # Date column should contain 2026-09-01
        found_date = False
        for row in ws_staff.iter_rows(min_row=7, max_col=1, values_only=True):
            if row[0] == "2026-09-01":
                found_date = True
                break
        assert found_date, "Expected 2026-09-01 in staff daily log sheet"

def test_export_ta_excel_subadmin_forbidden():
    store = {
        "admin_users": {
            "dc_gaya": {
                "name": "DC Gaya",
                "role": "SUB_ADMIN",
                "allowed_districts": ["Gaya"]
            }
        }
    }
    mock_db = MagicMock()
    mock_db.collection.side_effect = lambda c: MockCollection(c, store)

    token = create_access_token({
        "sub": "dc_gaya",
        "role": "SUB_ADMIN",
        "name": "DC Gaya",
        "allowed_districts": ["Gaya"]
    })
    headers = {"Authorization": f"Bearer {token}"}

    with patch("main.db", mock_db):
        res = client.get("/api/ta-logs/export-excel?month=2026-09&district=Patna", headers=headers)
        assert res.status_code == 403
